from __future__ import annotations

import json
import time

import pytest

# 无阿拉伯数字、不与源文 10 字雷同的仿写段——顺序无关（并行 fill 下两节可互换）
SEC1_REPLY = json.dumps(
    {"heading": "一、项目概述", "paras": ["项目整体规模与建设进展已按要求完成梳理。"]},
    ensure_ascii=False)
SEC2_REPLY = json.dumps(
    {"heading": "二、季度指标", "paras": ["季度各项指标情况已逐项核对完毕。"]},
    ensure_ascii=False)
PARAS_REPLY = json.dumps({"paras": ["该部分内容已按要求完成整体梳理。"]}, ensure_ascii=False)


@pytest.fixture()
def web_app(tmp_path):
    from app.main import create_app

    return create_app(jobs_root=tmp_path / "jobs", com_export=False)


def _wait_status(client, job_id, statuses, timeout=60.0):
    deadline = time.monotonic() + timeout
    d = None
    while time.monotonic() < deadline:
        d = client.get(f"/api/jobs/{job_id}").json()
        if d["status"] in statuses:
            return d
        time.sleep(0.05)
    raise AssertionError(f"等待状态超时：{d['status']}（{d.get('detail')}）")


def test_web_doc_flow_confirm_and_download(basic_docx, web_app, fake_llm_factory):
    from fastapi.testclient import TestClient

    state = {"script": []}  # 规则体裁置信度足够，PLANNED 前零 LLM 调用

    with TestClient(web_app) as client:
        web_app.state.llm_client_factory = lambda: fake_llm_factory(state["script"])[0]

        with open(basic_docx, "rb") as f:
            r = client.post("/api/jobs", files={"file": ("basic.docx", f)},
                            data={"product": "doc"})
        assert r.status_code == 200, r.text
        job_id = r.json()["job_id"]

        d = _wait_status(client, job_id, {"PLANNED"})
        assert d["product"] == "doc"
        assert d["genre"] == "report"
        assert d["skin"] == "business_blue"
        assert [o["type"] for o in d["outline"]] == ["h1", "h1"]
        assert d["outline"][0]["title"] == "一、项目概述"
        assert d["output_docx"] is None

        # 非法体裁 → 409，且任务仍停在 PLANNED
        r = client.post(f"/api/jobs/{job_id}/confirm", data={"genre": "bogus"})
        assert r.status_code == 409
        assert client.get(f"/api/jobs/{job_id}").json()["status"] == "PLANNED"

        state["script"] = [SEC1_REPLY, SEC2_REPLY]
        assert client.post(f"/api/jobs/{job_id}/confirm",
                           data={"genre": "report"}).status_code == 200
        d = _wait_status(client, job_id, {"DONE", "FAILED"})
        assert d["status"] == "DONE", d.get("error")
        assert d["output_docx"].endswith("output.docx")
        assert d["output_pdf"] is None  # com_export=False 未出 PDF
        assert d["pages_png"] == []

        r = client.get(f"/jobs/{job_id}/download?format=docx")
        assert r.status_code == 200
        assert len(r.content) > 1000
        assert "wordprocessingml" in r.headers["content-type"]
        assert 'filename="basic.docx"' in r.headers["content-disposition"]

        assert client.get(f"/jobs/{job_id}/download?format=pdf").status_code == 404
        assert client.get(f"/jobs/{job_id}/download?format=pptx").status_code == 400
        assert client.get(f"/jobs/{job_id}/download").status_code == 400  # doc 默认不许 pptx


def test_web_doc_confirm_changes_genre(basic_docx, web_app, fake_llm_factory):
    from fastapi.testclient import TestClient

    state = {"script": []}

    with TestClient(web_app) as client:
        web_app.state.llm_client_factory = lambda: fake_llm_factory(state["script"])[0]

        with open(basic_docx, "rb") as f:
            r = client.post("/api/jobs", files={"file": ("basic.docx", f)},
                            data={"product": "doc"})
        job_id = r.json()["job_id"]
        _wait_status(client, job_id, {"PLANNED"})

        # report → letter：fill 改走 paras 模式，docir 落 letter
        state["script"] = [PARAS_REPLY, PARAS_REPLY]
        assert client.post(f"/api/jobs/{job_id}/confirm",
                           data={"genre": "letter"}).status_code == 200
        d = _wait_status(client, job_id, {"DONE", "FAILED"})
        assert d["status"] == "DONE", d.get("error")

        import json as _json

        from pathlib import Path

        docir = _json.loads(
            Path(d["output_docx"]).parent.joinpath("docir.json").read_text(encoding="utf-8"))
        assert docir["meta"]["genre"] == "letter"
        kinds = [b["kind"] for b in docir["blocks"]]
        assert "heading" not in kinds  # letter 无标题块


def test_web_doc_upload_restrictions(web_app):
    from fastapi.testclient import TestClient

    with TestClient(web_app) as client:
        r = client.post("/api/jobs", files={"file": ("deck.pptx", b"junk")},
                        data={"product": "doc"})
        assert r.status_code == 400
        assert "docx / doc / pdf" in r.json()["detail"]

        r = client.post("/api/jobs", files={"file": ("x.docx", b"junk")},
                        data={"product": "bogus"})
        assert r.status_code == 400


def test_web_doc_upload_legacy_doc(basic_docx, web_app, fake_llm_factory, monkeypatch):
    """.doc 上传走 Word 转换分支：受理 200，解析后照常 PLANNED。"""
    from pathlib import Path

    import app.services.com_export as ce
    from fastapi.testclient import TestClient

    monkeypatch.setattr(ce, "convert_doc_to_docx",
                        lambda p, out=None: (Path(out).write_bytes(basic_docx.read_bytes()),
                                             Path(out))[1])

    state = {"script": []}
    with TestClient(web_app) as client:
        web_app.state.llm_client_factory = lambda: fake_llm_factory(state["script"])[0]

        with open(basic_docx, "rb") as f:
            r = client.post("/api/jobs", files={"file": ("旧版公文.doc", f)},
                            data={"product": "doc"})
        assert r.status_code == 200, r.text
        job_id = r.json()["job_id"]

        d = _wait_status(client, job_id, {"PLANNED"})
        assert d["product"] == "doc"
        assert d["genre"] == "report"
        assert len(d["outline"]) == 2


def test_web_doc_tablefill_fallback(web_app, fake_llm_factory, tmp_path):
    """长文本格仿写失败 → 退回照搬 + W 级报告行，不阻断 DONE。"""
    from pathlib import Path

    from docx import Document
    from fastapi.testclient import TestClient

    src = tmp_path / "建设汇报.docx"
    doc = Document()
    doc.add_heading("项目建设情况汇报", 0)
    doc.add_heading("一、项目概述", level=1)
    doc.add_paragraph("本项目覆盖 12 个园区，接入设备 8,600 台。")
    doc.add_heading("二、工作进展", level=1)
    doc.add_paragraph("各项工作按计划推进。")
    t = doc.add_table(rows=2, cols=2)
    for j, h in enumerate(["项目", "说明"]):
        t.rows[0].cells[j].text = h
    t.rows[1].cells[0].text = "运维职责"
    t.rows[1].cells[1].text = "统筹产线日常管理与设备运维督导，覆盖 12 条产线"
    doc.save(str(src))

    state = {"script": []}
    with TestClient(web_app) as client:
        web_app.state.llm_client_factory = lambda: fake_llm_factory(state["script"])[0]

        with open(src, "rb") as f:
            r = client.post("/api/jobs", files={"file": (src.name, f)},
                            data={"product": "doc"})
        job_id = r.json()["job_id"]
        _wait_status(client, job_id, {"PLANNED"})

        # tablefill 回复丢了 ⟦1⟧ 占位符 → 该格退回照搬
        bad_tbl = json.dumps(
            {"cells": {"0,1": "统筹产线日常管理与设备运维督导"}},
            ensure_ascii=False)
        # 重试轮仍丢占位符 → 最终退格（tablefill 首轮 + 定向重试各一次调用）
        state["script"] = [SEC1_REPLY, SEC2_REPLY, bad_tbl, bad_tbl]
        assert client.post(f"/api/jobs/{job_id}/confirm",
                           data={"genre": "report"}).status_code == 200
        d = _wait_status(client, job_id, {"DONE", "FAILED"})
        assert d["status"] == "DONE", d.get("error")

        art_dir = Path(d["output_docx"]).parent
        qa = (art_dir / "qa_report.txt").read_text(encoding="utf-8")
        assert "[tablefill] W-CELL-FALLBACK tbl-001 0,1: 占位符回填失败" in qa
        tbl_art = json.loads(
            (art_dir / "tables" / "tbl-001.json").read_text(encoding="utf-8"))
        assert tbl_art["cells"] == {}
        assert tbl_art["fallbacks"] == {"0,1": "占位符回填失败"}

        # 兜底照搬：输出表格该格与源文一致
        out = Document(d["output_docx"])
        assert out.tables[0].rows[1].cells[1].text == \
            "统筹产线日常管理与设备运维督导，覆盖 12 条产线"


def test_web_doc_image_roundtrip(web_app, fake_llm_factory, tmp_path):
    """带图 docx 全链路：解析→规划→fill→组装→渲染，图片原样进输出包。"""
    import base64
    import io

    from docx import Document
    from docx.shared import Cm
    from fastapi.testclient import TestClient

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
        "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    src = tmp_path / "流程手册.docx"
    doc = Document()
    doc.add_heading("一、项目概述", level=1)
    doc.add_paragraph("本项目覆盖 12 个园区，接入设备 8,600 台。")
    doc.add_picture(io.BytesIO(png), width=Cm(5))
    doc.add_heading("二、季度指标", level=1)
    doc.add_paragraph("第三季度各项指标完成情况良好。")
    doc.save(str(src))

    state = {"script": []}
    with TestClient(web_app) as client:
        web_app.state.llm_client_factory = lambda: fake_llm_factory(state["script"])[0]

        with open(src, "rb") as f:
            r = client.post("/api/jobs", files={"file": (src.name, f)},
                            data={"product": "doc"})
        job_id = r.json()["job_id"]
        _wait_status(client, job_id, {"PLANNED"})

        state["script"] = [SEC1_REPLY, SEC2_REPLY]
        assert client.post(f"/api/jobs/{job_id}/confirm",
                           data={"genre": "report"}).status_code == 200
        d = _wait_status(client, job_id, {"DONE", "FAILED"})
        assert d["status"] == "DONE", d.get("error")

        out = Document(d["output_docx"])
        assert len(out.inline_shapes) == 1
        # 图片部件内容与源一致（XML 搬运 + 关系重接）
        img_part = next(r.target_part for r in out.part.rels.values()
                        if r.reltype.endswith("/image"))
        assert img_part.blob == png
