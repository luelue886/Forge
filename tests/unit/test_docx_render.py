from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.render.docx_render import render_docir_to_docx
from app.schema.docir import (
    ClosingBlock,
    DocIR,
    DocIRMeta,
    DocTitleBlock,
    HeadingBlock,
    ImageBlock,
    ParaBlock,
    SalutationBlock,
    SignatureBlock,
    TableBlock,
)
from app.schema.enums import Genre


def _east(run) -> str:
    return run._element.rPr.rFonts.get(qn("w:eastAsia"))


def _fill(cell) -> str | None:
    shd = cell._tc.tcPr.find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def _letter_doc() -> DocIR:
    return DocIR(meta=DocIRMeta(title="调岗申请书", genre=Genre.LETTER), blocks=[
        DocTitleBlock(text="调岗申请书"),
        SalutationBlock(text="尊敬的公司领导："),
        ParaBlock(text="恳请批准调岗。"),
        ClosingBlock(lines=["此致", "敬礼！"]),
        SignatureBlock(signer="申请人：王明", date="2026 年 9 月 18 日"),
    ])


def test_letter_layout(tmp_path: Path):
    out = render_docir_to_docx(_letter_doc(), tmp_path / "letter.docx")
    d = Document(str(out))

    texts = [p.text for p in d.paragraphs]
    assert texts == ["调岗申请书", "尊敬的公司领导：", "恳请批准调岗。",
                     "此致", "敬礼！", "申请人：王明", "2026 年 9 月 18 日"]
    assert len(d.tables) == 0

    title = d.paragraphs[0]
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert title.runs[0].font.size == Pt(22)
    assert title.runs[0].font.bold is True
    assert _east(title.runs[0]) == "黑体"

    salutation = d.paragraphs[1]
    assert salutation.paragraph_format.first_line_indent is None  # 顶格
    assert _east(salutation.runs[0]) == "仿宋"

    body = d.paragraphs[2]
    assert body.paragraph_format.first_line_indent == Pt(28)  # 2 字符
    assert body.paragraph_format.line_spacing == 1.5

    closing1, closing2 = d.paragraphs[3], d.paragraphs[4]
    assert closing1.paragraph_format.first_line_indent == Pt(28)
    assert closing2.paragraph_format.first_line_indent is None

    signer, date = d.paragraphs[5], d.paragraphs[6]
    assert signer.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert date.alignment == WD_ALIGN_PARAGRAPH.RIGHT


def test_report_headings_and_table_verbatim(tmp_path: Path):
    header = ["项目", "金额（万元）"]
    rows = [["营业收入", "1,234.56"], ["回款金额", "986.20"]]
    doc = DocIR(meta=DocIRMeta(title="运营报告", genre=Genre.REPORT), blocks=[
        DocTitleBlock(text="运营报告"),
        HeadingBlock(level=1, text="一、总体情况"),
        ParaBlock(text="营收 1,234.56 万元。"),
        HeadingBlock(level=2, text="（一）明细"),
        ParaBlock(text="明细见下表。"),
        TableBlock(table_id="tbl-001", header=list(header),
                   rows=[list(r) for r in rows]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "report.docx")
    d = Document(str(out))

    h1 = d.paragraphs[1]
    assert h1.runs[0].font.size == Pt(16)
    assert _east(h1.runs[0]) == "黑体"
    h2 = d.paragraphs[3]
    assert h2.runs[0].font.size == Pt(14)

    t = d.tables[0]
    # 表格 verbatim diff：单元格与源数据逐格一致
    assert [t.cell(0, c).text for c in range(2)] == header
    for i, row in enumerate(rows, start=1):
        assert [t.cell(i, c).text for c in range(2)] == row
    assert _fill(t.cell(0, 0)) == "D9D9D9"   # 表头浅灰
    assert _fill(t.cell(1, 0)) is None       # 数据行 1 无底纹
    assert _fill(t.cell(2, 0)) == "F5F5F5"   # 斑马纹
    assert t.cell(0, 0).paragraphs[0].runs[0].font.bold is True
    assert _east(t.cell(1, 1).paragraphs[0].runs[0]) == "宋体"
    assert t.cell(1, 1).paragraphs[0].runs[0].font.size == Pt(10.5)


def test_table_without_header(tmp_path: Path):
    doc = DocIR(meta=DocIRMeta(title="表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="表"),
        TableBlock(table_id="tbl-001", header=[],
                   rows=[["a", "b"], ["c", "d"]]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "form.docx")
    d = Document(str(out))
    t = d.tables[0]
    assert len(t.rows) == 2
    assert [t.cell(0, 0).text, t.cell(0, 1).text] == ["a", "b"]
    assert _fill(t.cell(0, 0)) is None  # 无表头 → 无表头底纹


def test_wide_table_shrinks_font(tmp_path: Path):
    header = [f"列{i}" for i in range(1, 9)]  # 8 列 ≥ 7
    doc = DocIR(meta=DocIRMeta(title="宽表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="宽表"),
        TableBlock(table_id="tbl-001", header=header, rows=[["x"] * 8]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "wide.docx")
    d = Document(str(out))
    assert d.tables[0].cell(1, 0).paragraphs[0].runs[0].font.size == Pt(9)


# ---- C7：合并区/行高渲染（PDF 源表格重建路径）----

def _row_spans(tr):
    """一行 [(gridSpan, vMerge), …]——合并结构指纹。"""
    out = []
    for tc in tr.findall(qn("w:tc")):
        tc_pr = tc.find(qn("w:tcPr"))
        span, vm = 1, None
        if tc_pr is not None:
            gs = tc_pr.find(qn("w:gridSpan"))
            if gs is not None and (gs.get(qn("w:val")) or "").isdigit():
                span = int(gs.get(qn("w:val")))
            v = tc_pr.find(qn("w:vMerge"))
            if v is not None:
                vm = v.get(qn("w:val")) or "continue"
        out.append((span, vm))
    return out


def test_table_merge_and_row_heights_render(tmp_path: Path):
    # 表头行本身含合并区（PDF 重建常见形态：首行整行/局部合并被识别为表头）
    doc = DocIR(meta=DocIRMeta(title="合并表格", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="合并表格"),
        TableBlock(
            table_id="tbl-001",
            header=["合并区甲", "合并区甲", "字段乙", "字段乙"],
            rows=[["合并区甲", "合并区甲", "子项丙", "子项丁"],
                  ["行三一", "行三二", "行三三", "行三四"]],
            col_widths=[0.25, 0.25, 0.25, 0.25],
            merges=[[0, 0, 2, 2], [0, 2, 1, 2]],
            row_heights=[40.0, 50.0, 60.0]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "merged.docx")
    d = Document(str(out))
    assert len(d.tables) == 1
    tbl = d.tables[0]
    rows = tbl._tbl.findall(qn("w:tr"))
    assert len(rows) == 3
    # 首行：2×2 区左上角（restart）+ 横向 1×2 区
    assert _row_spans(rows[0]) == [(2, "restart"), (2, None)]
    # 次行：2×2 区纵向延续 + 两个普通格
    assert _row_spans(rows[1]) == [(2, "continue"), (1, None), (1, None)]
    assert _row_spans(rows[2]) == [(1, None)] * 4
    # 行高 atLeast（pt × 20 twips）
    tr_pr = rows[1].find(qn("w:trPr"))
    tr_h = tr_pr.find(qn("w:trHeight"))
    assert tr_h.get(qn("w:val")) == "1000"
    assert tr_h.get(qn("w:hRule")) == "atLeast"
    # 合并表禁表头/斑马纹底纹
    assert _fill(tbl.cell(0, 0)) is None
    assert _fill(tbl.cell(2, 1)) is None
    # 合并区文本只写一次（展开位去重）
    assert tbl.cell(0, 0).text == "合并区甲"
    assert tbl.cell(1, 0).text == "合并区甲"
    assert tbl.cell(0, 2).text == "字段乙"


def test_table_invalid_merges_skipped(tmp_path: Path):
    # 越界/重叠/退化合并项被丢弃，表格照常渲染不崩
    doc = DocIR(meta=DocIRMeta(title="坏合并", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="坏合并"),
        TableBlock(table_id="tbl-001", header=["甲", "乙"],
                   rows=[["1", "2"], ["3", "4"]],
                   merges=[[0, 0, 9, 9], [0, 0, 1, 1], [1, 1, 2, 2]]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "bad.docx")
    d = Document(str(out))
    rows = d.tables[0]._tbl.findall(qn("w:tr"))
    # [1,1,2,2] 越界丢弃、[0,0,1,1] 退化丢弃 → 无合并生效，
    # 且普通渲染路径（含表头底纹）不受影响
    assert all(span == 1 and vm is None
               for r in rows for span, vm in _row_spans(r))
    assert _fill(d.tables[0].cell(0, 0)) == "D9D9D9"


def test_page_setup_a4_and_margins(tmp_path: Path):
    out = render_docir_to_docx(_letter_doc(), tmp_path / "page.docx")
    sec = Document(str(out)).sections[0]
    # twips 落盘存在 ≤310 EMU 的舍入误差，用容差断言
    assert abs(sec.page_width - Cm(21.0)) < 400
    assert abs(sec.page_height - Cm(29.7)) < 400
    assert abs(sec.top_margin - Cm(2.54)) < 400
    assert abs(sec.left_margin - Cm(3.18)) < 400
    # 页脚有 PAGE 域
    footer_xml = sec.footer.paragraphs[0]._p.xml
    assert "PAGE" in footer_xml


# ---- C3：表格排版保真 ----

def _source_with_merged_table(path: Path) -> Path:
    from docx.shared import Cm as _Cm

    doc = Document()
    doc.add_paragraph("表1：测试")
    t = doc.add_table(rows=2, cols=3)
    for i, w in enumerate((_Cm(6), _Cm(3), _Cm(2))):
        t.columns[i].width = w
    t.cell(0, 0).text = "基本信息"
    t.cell(0, 2).text = "备注"
    t.cell(1, 0).text = "张三"
    t.cell(1, 1).text = "35"
    t.cell(1, 2).text = "该同志负责产线管理工作。"
    t.cell(0, 0).merge(t.cell(0, 1))  # gridSpan=2
    doc.save(str(path))
    return path


def _gridcol_widths(tbl_el) -> list[int]:
    grid = tbl_el.find(qn("w:tblGrid"))
    return [int(gc.get(qn("w:w"))) for gc in grid.findall(qn("w:gridCol"))]


def test_render_copies_source_table_layout(tmp_path: Path):
    src = _source_with_merged_table(tmp_path / "src.docx")
    src_tbl = Document(str(src)).tables[0]._tbl

    doc = DocIR(meta=DocIRMeta(title="信息表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="信息表"),
        TableBlock(table_id="tbl-001", src_index=0,
                   header=["基本信息", "基本信息", "备注"],
                   rows=[["张三", "35", "该同志负责产线管理工作。"]]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx", source=src)
    d = Document(str(out))
    assert len(d.tables) == 1
    got = d.tables[0]
    # 列宽与源逐值一致（deepcopy 保真），内容 verbatim
    assert _gridcol_widths(got._tbl) == _gridcol_widths(src_tbl)
    assert got.cell(1, 1).text == "35"


def test_render_source_index_out_of_range_falls_back(tmp_path: Path):
    src = _source_with_merged_table(tmp_path / "src.docx")
    doc = DocIR(meta=DocIRMeta(title="表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="表"),
        TableBlock(table_id="tbl-001", src_index=99,
                   header=["甲", "乙"], rows=[["1", "2"]]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx", source=src)
    d = Document(str(out))
    t = d.tables[0]
    assert t.cell(0, 0).text == "甲"      # 规范重建路径
    assert _fill(t.cell(0, 0)) == "D9D9D9"  # 规范表头底纹（拷贝路径没有）


def test_render_table_col_widths_ratio(tmp_path: Path):
    """无源可搬（PDF 源）时按 col_widths 比例设列宽。"""
    doc = DocIR(meta=DocIRMeta(title="表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="表"),
        TableBlock(table_id="tbl-001", col_widths=[0.5, 0.3, 0.2],
                   header=["甲", "乙", "丙"], rows=[["1", "2", "3"]]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx")
    t = Document(str(out)).tables[0]
    ws = [tc.tcPr.find(qn("w:tcW")).get(qn("w:w"))
          for tc in t.rows[0]._tr.findall(qn("w:tc"))]
    assert ws and all(w is not None and w.isdigit() for w in ws)
    assert int(ws[0]) > int(ws[1]) > int(ws[2]) > 0


# ---- C4: 图片/流程图复用 ----

import base64
import io

_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def _source_with_image(path: Path) -> Path:
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("一、流程说明")
    doc.add_picture(io.BytesIO(_PNG_1PX), width=Cm(6))
    doc.save(str(path))
    return path


def test_render_copies_source_image(tmp_path: Path):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.shared import Cm

    from app.parsing.docx_parser import parse_docx

    src = _source_with_image(tmp_path / "src.docx")
    img = parse_docx(src).images[0]  # body_index 用解析器同一坐标系

    doc = DocIR(meta=DocIRMeta(title="流程", genre=Genre.REPORT), blocks=[
        DocTitleBlock(text="流程"),
        ImageBlock(image_id=img.image_id, body_index=img.body_index),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx", source=src)
    d = Document(str(out))
    assert len(d.inline_shapes) == 1
    assert d.inline_shapes[0].width == Cm(6)  # 尺寸随源 XML 原样
    rels = [r for r in d.part.rels.values() if r.reltype == RT.IMAGE]
    assert len(rels) == 1 and rels[0].target_part.blob


def test_render_image_body_index_out_of_range(tmp_path: Path):
    src = _source_with_image(tmp_path / "src.docx")
    doc = DocIR(meta=DocIRMeta(title="流程", genre=Genre.REPORT), blocks=[
        DocTitleBlock(text="流程"),
        ImageBlock(image_id="img-001", body_index=99),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx", source=src)
    assert len(Document(str(out)).inline_shapes) == 0  # 降级跳过，不崩


def test_render_pdf_source_image(tmp_path: Path):
    import fitz

    from app.parsing.pdf_parser import parse_pdf
    from app.render import docstyle as st

    pdf_doc = fitz.open()
    page = pdf_doc.new_page(width=595, height=842)
    page.insert_text((72, 72), "年度工作流程说明，整体运转顺畅，各环节衔接有序。",
                     fontname="china-s", fontsize=11)
    page.insert_text((72, 92), "详细流程见下图示意，实际执行中以最新通知为准。",
                     fontname="china-s", fontsize=11)
    page.insert_image(fitz.Rect(150, 150, 450, 300), stream=_PNG_1PX)
    src = tmp_path / "src.pdf"
    pdf_doc.save(str(src))
    pdf_doc.close()

    img = parse_pdf(src).images[0]
    doc = DocIR(meta=DocIRMeta(title="流程", genre=Genre.REPORT), blocks=[
        DocTitleBlock(text="流程"),
        ImageBlock(image_id=img.image_id, page=img.page, bbox=img.bbox),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx", source=src)
    d = Document(str(out))
    assert len(d.inline_shapes) == 1
    # 相对源页宽缩放：300pt / 595pt 占版心的同比例
    content_w = st.PAGE_WIDTH - st.MARGIN_LEFT - st.MARGIN_RIGHT
    expect = int(content_w * 300 / 595)
    assert abs(d.inline_shapes[0].width - expect) < 20000


# ---- C8：宽表缩放（源页边距窄 → dxa 定宽表溢出输出版心）----

def _source_with_dxa_table(path: Path, col_tw: list[int]) -> Path:
    """造 tblW=dxa 定宽表（gridCol/tcW 逐值 + tblInd 152），模拟小边距源文档。"""
    from docx.oxml import OxmlElement
    from docx.shared import Twips

    doc = Document()
    doc.add_paragraph("表1：宽表")
    t = doc.add_table(rows=2, cols=len(col_tw))
    for c, w in enumerate(col_tw):
        t.columns[c].width = Twips(w)
        for r in range(2):
            t.cell(r, c).width = Twips(w)
    tbl_pr = t._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(col_tw)))
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is None:
        ind = OxmlElement("w:tblInd")
        tbl_pr.append(ind)
    ind.set(qn("w:type"), "dxa")
    ind.set(qn("w:w"), "152")
    doc.save(str(path))
    return path


def test_render_wide_source_table_scaled(tmp_path):
    # 10425 dxa（六种合集实测宽度）> 输出版心 ~8300 → 等比缩放 + tblInd 归零
    from docx.shared import Emu

    from app.render import docstyle as st

    cols = [2880, 3742, 3803]  # 和 10425
    src = _source_with_dxa_table(tmp_path / "wide_src.docx", cols)
    doc = DocIR(meta=DocIRMeta(title="宽表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="宽表"),
        TableBlock(table_id="tbl-001", src_index=0,
                   header=["字段", "值", "备注"], rows=[["甲", "乙", "丙"]]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx", source=src)
    tbl = Document(str(out)).tables[0]
    content_tw = Emu(st.PAGE_WIDTH - st.MARGIN_LEFT - st.MARGIN_RIGHT).twips

    tbl_w = tbl._tbl.tblPr.find(qn("w:tblW"))
    assert tbl_w.get(qn("w:type")) == "dxa"
    assert int(tbl_w.get(qn("w:w"))) <= content_tw
    got = _gridcol_widths(tbl._tbl)
    assert sum(got) <= content_tw + 3  # 逐列舍入余量
    for s, g in zip(cols, got):  # 列宽比例保持
        assert abs(g / sum(got) - s / sum(cols)) < 0.02
    ind = tbl._tbl.tblPr.find(qn("w:tblInd"))
    assert ind is not None and int(ind.get(qn("w:w")) or 0) == 0


def test_render_fitting_source_table_untouched(tmp_path):
    # 表宽 ≤ 版心：原样搬运（列宽与 tblInd 均不动）
    cols = [2000, 2500, 3000]  # 和 7500 ≤ ~8300
    src = _source_with_dxa_table(tmp_path / "fit_src.docx", cols)
    doc = DocIR(meta=DocIRMeta(title="表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="表"),
        TableBlock(table_id="tbl-001", src_index=0,
                   header=["字段", "值", "备注"], rows=[["甲", "乙", "丙"]]),
    ])
    out = render_docir_to_docx(doc, tmp_path / "out.docx", source=src)
    tbl = Document(str(out)).tables[0]
    assert _gridcol_widths(tbl._tbl) == cols
    ind = tbl._tbl.tblPr.find(qn("w:tblInd"))
    assert ind is not None and ind.get(qn("w:w")) == "152"


# ---- C9: 表格内嵌证件照/装饰小图删除 ----

_DRAW_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
_V_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"
_V_SHAPE = "{urn:schemas-microsoft-com:vml}shape"


def _src_with_cell_graphic(path, run_xml: str) -> Path:
    from docx.oxml import parse_xml

    doc = Document()
    doc.add_paragraph("表1：含图表格")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "姓名"
    t.rows[0].cells[1].text = "李雷"
    cell = t.cell(1, 0)
    cell.paragraphs[0].add_run("照片：")
    cell.paragraphs[0]._p.append(parse_xml(f'<w:r {_DRAW_NS}>{run_xml}</w:r>'))
    t.rows[1].cells[1].text = "备注"
    doc.save(str(path))
    return path


def _form_docir() -> DocIR:
    return DocIR(meta=DocIRMeta(title="含图表格", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="含图表格"),
        TableBlock(table_id="tbl-001", src_index=0,
                   header=[], rows=[["姓名", "李雷"], ["照片：", "备注"]]),
    ])


def test_render_table_strips_portrait_drawing(tmp_path):
    # 2.8×3.5cm wp:extent → 证件照：表格 XML 搬运时删除，文字保留
    src = _src_with_cell_graphic(
        tmp_path / "photo_src.docx",
        '<w:drawing><wp:inline>'
        '<wp:extent cx="1008000" cy="1260000"/></wp:inline></w:drawing>')
    out = render_docir_to_docx(_form_docir(), tmp_path / "out.docx", source=src)
    tbl = Document(str(out)).tables[0]
    assert tbl._tbl.find(".//" + qn("w:drawing")) is None
    assert tbl.cell(1, 0).text == "照片："


def test_render_table_strips_portrait_vml_pict(tmp_path):
    # VML 图片无 wp:extent，尺寸从 v:shape style 解析（79.4×99.2pt ≈ 2.8×3.5cm）
    src = _src_with_cell_graphic(
        tmp_path / "vml_src.docx",
        '<w:pict><v:shape style="width:79.4pt;height:99.2pt">'
        '<v:imagedata r:id="rId99"/></v:shape></w:pict>')
    out = render_docir_to_docx(_form_docir(), tmp_path / "out.docx", source=src)
    tbl = Document(str(out)).tables[0]
    assert tbl._tbl.find(".//" + _V_IMAGEDATA) is None
    assert tbl.cell(1, 0).text == "照片："


def test_render_table_keeps_content_drawing(tmp_path):
    # 14×3.2cm 横幅图（golden form_personnel 实测尺寸）非证件照 → 保留
    src = _src_with_cell_graphic(
        tmp_path / "banner_src.docx",
        '<w:drawing><wp:inline>'
        '<wp:extent cx="5040000" cy="1152000"/></wp:inline></w:drawing>')
    out = render_docir_to_docx(_form_docir(), tmp_path / "out.docx", source=src)
    tbl = Document(str(out)).tables[0]
    assert tbl._tbl.find(".//" + qn("w:drawing")) is not None


def test_render_table_keeps_textbox_pict(tmp_path):
    # v:pict 内是文本框不是图片 → 不动（无尺寸可判也不猜）
    src = _src_with_cell_graphic(
        tmp_path / "txbx_src.docx",
        '<w:pict><v:shape style="width:79.4pt;height:99.2pt">'
        '<v:textbox><w:txbxContent><w:p><w:r><w:t>备注说明</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict>')
    out = render_docir_to_docx(_form_docir(), tmp_path / "out.docx", source=src)
    tbl = Document(str(out)).tables[0]
    assert tbl._tbl.find(".//" + _V_SHAPE) is not None
