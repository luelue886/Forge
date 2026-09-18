from __future__ import annotations

import json

import pytest

from app.pipeline.fill import FillError, fill_deterministic, fill_page
from app.schema.enums import SlideType, TableStrategy
from app.schema.plan import SlidePlanItem, TableRef


def _item(page_no, slide_type, title="测试页", brief="", src_refs=None, table_ref=None):
    return SlidePlanItem(page_no=page_no, slide_type=slide_type, title=title,
                         brief=brief, src_refs=src_refs or [], table_ref=table_ref)


def _tree(docx_path):
    from app.parsing.docx_parser import parse_docx

    return parse_docx(docx_path)


def test_deterministic_structural_pages(basic_docx):
    tree = _tree(basic_docx)

    cover = fill_deterministic(
        _item(1, SlideType.COVER, title="智慧园区平台建设汇报", brief=""),
        tree)
    assert cover.slide_type is SlideType.COVER
    assert cover.title == "智慧园区平台建设汇报"
    assert not cover.subtitle  # 封面不再带"基于《源文档》仿写"副题

    toc = fill_deterministic(_item(2, SlideType.TOC), tree,
                             toc_entries=["一、项目概述", "二、季度指标"])
    assert toc.title == "目录"
    assert toc.bullets == ["一、项目概述", "二、季度指标"]

    sec = fill_deterministic(_item(3, SlideType.SECTION_HEADER, title="一、项目概述",
                                   brief="章节说明"), tree)
    assert sec.title == "一、项目概述"
    assert sec.subtitle == "章节说明"

    closing = fill_deterministic(_item(10, SlideType.CLOSING), tree,
                                 doc_title="智慧园区平台建设汇报")
    assert closing.title == "谢谢观看"
    assert closing.subtitle == "智慧园区平台建设汇报"

    with pytest.raises(FillError):
        fill_deterministic(_item(2, SlideType.TOC), tree)  # 缺章节条目


def test_structural_pages_never_touch_llm(basic_docx, fake_llm_factory):
    tree = _tree(basic_docx)
    client, fake = fake_llm_factory([])  # 脚本为空：任何真实调用都会 IndexError
    ir = fill_page(_item(1, SlideType.COVER, title="智慧园区平台建设汇报"), tree, client)
    assert ir.slide_type is SlideType.COVER
    assert fake.calls == []


def test_table_page_verbatim(basic_docx):
    tree = _tree(basic_docx)
    item = _item(5, SlideType.TABLE, title="Q3 关键指标",
                 table_ref=TableRef(table_id="tbl-001", strategy=TableStrategy.COPY))
    ir = fill_page(item, tree, None)  # client=None：证明表格页零 LLM
    assert ir.table is not None
    assert ir.table.header == ["指标", "目标", "实际"]
    assert ir.table.rows == [["园区数", "10", "12"], ["设备数", "8000", "8600"]]
    assert ir.note is None


def test_table_page_errors(basic_docx):
    tree = _tree(basic_docx)
    with pytest.raises(FillError):
        fill_deterministic(_item(5, SlideType.TABLE), tree)
    with pytest.raises(FillError):
        fill_deterministic(
            _item(5, SlideType.TABLE,
                  table_ref=TableRef(table_id="tbl-099", strategy=TableStrategy.COPY)),
            tree)


def test_table_trim_with_note(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_heading("大表文档", 0)
    doc.add_heading("一、设备明细", level=1)
    doc.add_paragraph("表1：设备清单")
    t = doc.add_table(rows=8, cols=3)
    for j, h in enumerate(["设备", "数量", "状态"]):
        t.rows[0].cells[j].text = h
    for i in range(7):
        for j, v in enumerate([f"设备{i}", str(100 + i), "在线"]):
            t.rows[i + 1].cells[j].text = v
    p = tmp_path / "big_table.docx"
    doc.save(str(p))

    tree = _tree(p)
    item = _item(1, SlideType.TABLE, title="设备清单",
                 table_ref=TableRef(table_id="tbl-001", strategy=TableStrategy.COPY))
    ir = fill_deterministic(item, tree)
    assert len(ir.table.rows) == 4  # 5 行含表头 → 最多 4 数据行
    assert ir.table.rows[0] == ["设备0", "100", "在线"]  # 从头保留
    assert ir.note is not None and "完整数据见源文档" in ir.note


BAD = json.dumps({
    "title": "这是一个特别特别特别长的标题超过二十个字了呀",
    "bullets": ["覆盖 12 个园区"], "metrics": [], "note": None,
}, ensure_ascii=False)
GOOD = json.dumps({
    "title": "项目覆盖规模",
    "bullets": ["覆盖 12 个园区", "接入设备 8,600 台"], "metrics": [], "note": None,
}, ensure_ascii=False)


def test_fill_llm_validates_and_retries(basic_docx, fake_llm_factory):
    tree = _tree(basic_docx)
    client, fake = fake_llm_factory([BAD, GOOD])
    item = _item(4, SlideType.TEXT_POINTS, title="项目覆盖规模",
                 brief="规模", src_refs=["blk-0001"])

    ir = fill_page(item, tree, client)

    assert ir.title == "项目覆盖规模"
    assert ir.bullets == ["覆盖 12 个园区", "接入设备 8,600 台"]
    assert len(fake.calls) == 2
    assert "V-TITLE-LEN" in json.dumps(fake.calls[1], ensure_ascii=False)  # 纠错明细回传


def test_fill_llm_gives_up_after_two_bad_responses(basic_docx, fake_llm_factory):
    tree = _tree(basic_docx)
    client, _ = fake_llm_factory([BAD, BAD])
    item = _item(4, SlideType.TEXT_POINTS, title="项目覆盖规模",
                 brief="规模", src_refs=["blk-0001"])
    with pytest.raises(FillError):
        fill_page(item, tree, client)
