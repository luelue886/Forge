"""docx_copy：源表格 XML 搬运排版保真 + 单元格文字替换（合并格游标契约）。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.shared import Cm

from app.render import docx_copy


def _make_source(path: Path) -> Path:
    """3 列表：显式列宽 + 表头横合并（gridSpan）+ 长文本格 + 数据行纵合并（vMerge）。"""
    doc = Document()
    doc.add_paragraph("表1：人员信息表")
    t = doc.add_table(rows=4, cols=3)
    for i, w in enumerate((Cm(6), Cm(3), Cm(2))):
        t.columns[i].width = w
    t.cell(0, 0).text = "基本信息"
    t.cell(0, 2).text = "备注"
    t.cell(1, 0).text = "张三"
    t.cell(1, 1).text = "35"
    t.cell(1, 2).text = "该同志长期负责产线管理工作，年度考核为优秀等次。"
    t.cell(2, 0).text = "李四"
    t.cell(2, 1).text = "42"
    t.cell(2, 2).text = "岗位调整待定。"
    t.cell(3, 0).text = "王五"
    t.cell(3, 1).text = "29"
    t.cell(3, 2).text = "新入职员工。"
    t.cell(0, 0).merge(t.cell(0, 1))   # 表头前两格横合并 → gridSpan=2
    t.cell(1, 2).merge(t.cell(2, 2))   # 备注列纵向合并 → vMerge
    doc.save(str(path))
    return path


def _gridcol_widths(tbl_el) -> list[int]:
    grid = tbl_el.find(qn("w:tblGrid"))
    return [int(gc.get(qn("w:w"))) for gc in grid.findall(qn("w:gridCol"))]


def _tr_spans(tbl_el) -> list[list[tuple[int, str]]]:
    """逐 tr 的 (gridSpan, vMerge) 序列——排版结构指纹。"""
    out = []
    for tr in tbl_el.findall(qn("w:tr")):
        desc = []
        for tc in tr.findall(qn("w:tc")):
            tc_pr = tc.find(qn("w:tcPr"))
            span, vm = 1, ""
            if tc_pr is not None:
                gs = tc_pr.find(qn("w:gridSpan"))
                if gs is not None:
                    span = int(gs.get(qn("w:val")))
                v = tc_pr.find(qn("w:vMerge"))
                if v is not None:
                    vm = v.get(qn("w:val")) or "continue"
            desc.append((span, vm))
        out.append(desc)
    return out


def test_copy_table_preserves_layout_and_rewrites_cells(tmp_path: Path):
    src = _make_source(tmp_path / "src.docx")
    src_doc = docx_copy.open_source(src)
    tbl_el = docx_copy.source_table(src_doc, 0)

    out_doc = Document()
    new_el = docx_copy.copy_table(src_doc, out_doc, tbl_el)
    # 目标网格：长文本格改写、数字格 35→36，其余 verbatim
    docx_copy.replace_table_texts(new_el, out_doc, header=["基本信息", "基本信息", "备注"],
                                  rows=[["张三", "36", "该同志多年承担产线管理职责，年度评定为优秀档次。"],
                                        ["李四", "42", "该同志多年承担产线管理职责，年度评定为优秀档次。"],
                                        ["王五", "29", "新入职员工。"]])
    out = tmp_path / "out.docx"
    out_doc.save(str(out))

    d = Document(str(out))
    got = d.tables[0]
    # 排版保真：列宽与合并结构与源逐项一致
    assert _gridcol_widths(got._tbl) == _gridcol_widths(tbl_el)
    assert _tr_spans(got._tbl) == _tr_spans(tbl_el)
    # 内容：改写格就位，verbatim 格原样
    assert got.cell(1, 1).text == "36"
    assert got.cell(1, 2).text == "该同志多年承担产线管理职责，年度评定为优秀档次。"
    assert got.cell(1, 0).text == "张三"
    assert got.cell(0, 2).text == "备注"
    assert got.cell(3, 2).text == "新入职员工。"


def test_replace_cell_texts_vmerge_cursor():
    """vMerge 续格跳过、restart 格写入；gridSpan 后续网格位不误写。"""
    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "标题"
    t.cell(0, 1).text = "值"
    t.cell(1, 0).text = "x"
    t.cell(1, 1).text = "y"
    t.cell(0, 0).merge(t.cell(1, 0))  # 第一列纵向合并
    tbl_el = t._tbl

    docx_copy.replace_cell_texts(tbl_el, {
        (0, 0): "新标题",        # restart 格 → 写入
        (1, 0): "不该写",       # 续格 → 跳过（显示的是 restart 的内容）
        (1, 1): "乙",           # 普通格 → 写入
    })
    # 直接按网格语义读回（wrap 到原文档）
    grid = docx_copy.grid_of(tbl_el, doc)
    assert grid == [["新标题", "值"], ["新标题", "乙"]]


def test_set_tc_text_clears_hyperlink_runs():
    """格内 w:hyperlink（嵌套 run，非段落直属子节点）改写后不得新旧文字叠加。"""
    from docx.oxml import OxmlElement

    doc = Document()
    t = doc.add_table(rows=1, cols=1)
    p = t.cell(0, 0)._tc.find(qn("w:p"))
    hl = OxmlElement("w:hyperlink")
    r = OxmlElement("w:r")
    wt = OxmlElement("w:t")
    wt.text = "旧链接文字"
    r.append(wt)
    hl.append(r)
    p.append(hl)

    docx_copy.replace_cell_texts(t._tbl, {(0, 0): "新文字"})
    assert t.cell(0, 0).text == "新文字"


def test_source_table_ordinal(tmp_path: Path):
    doc2 = Document()
    doc2.add_table(rows=1, cols=1)
    doc2.add_table(rows=1, cols=1)
    doc2.save(str(tmp_path / "two.docx"))
    two = docx_copy.open_source(tmp_path / "two.docx")
    assert docx_copy.source_table(two, 0) is not None
    assert docx_copy.source_table(two, 1) is not None
    assert docx_copy.source_table(two, 2) is None
    assert docx_copy.source_table(two, 99) is None


def test_copy_table_bring_style_chain(tmp_path: Path):
    """tblStyle 引用的样式（含 basedOn 祖先）一并拷进输出包。"""
    src = tmp_path / "styled.docx"
    doc = Document()
    s = doc.styles.add_style("MyTableStyle", WD_STYLE_TYPE.TABLE)
    s.base_style = doc.styles["Table Grid"]
    t = doc.add_table(rows=1, cols=1)
    t.cell(0, 0).text = "内容"
    t.style = s
    doc.save(str(src))

    src_doc = docx_copy.open_source(src)
    tbl_el = docx_copy.source_table(src_doc, 0)
    out_doc = Document()
    docx_copy.copy_table(src_doc, out_doc, tbl_el)

    out_ids = {el.get(qn("w:styleId"))
               for el in out_doc.styles.element.findall(qn("w:style"))}
    assert "MyTableStyle" in out_ids
    # basedOn 祖先 TableGrid 也补入（输出默认模板未必有）
    based = None
    for el in out_doc.styles.element.findall(qn("w:style")):
        if el.get(qn("w:styleId")) == "MyTableStyle":
            bo = el.find(qn("w:basedOn"))
            based = bo.get(qn("w:val")) if bo is not None else None
    assert based in out_ids


# ---- 宽表缩放：auto 型按 gridCol 合计判溢出（用户实测溢出回归）----

def _wide_auto_table(path: Path, col_tw: list[int]) -> Path:
    """tblW=auto 但 gridCol 合计远超版心的表（复刻《安全生产台账》源）。"""
    from docx.oxml import OxmlElement

    doc = Document()
    doc.add_paragraph("宽表")
    t = doc.add_table(rows=2, cols=len(col_tw))
    tbl_pr = t._tbl.tblPr
    for tag in ("w:tblW",):
        el = tbl_pr.find(qn(tag))
        if el is not None:
            tbl_pr.remove(el)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:type"), "auto")
    tbl_w.set(qn("w:w"), "0")
    tbl_pr.append(tbl_w)
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), col_tw):
        gc.set(qn("w:w"), str(w))
    for r in range(2):
        for c in range(len(col_tw)):
            t.cell(r, c).text = "x"
    doc.save(str(path))
    return path


def test_normalize_scales_auto_table_over_content(tmp_path: Path):
    """gridCol 合计 16594 twips（≈2 倍版心）→ 等比缩到版心 + tblW 转固定宽。"""
    src = _make_source(tmp_path / "src.docx")  # 复用公文案版心（窄）
    src_doc = Document(str(src))
    wide = _wide_auto_table(tmp_path / "wide.docx",
                            [907, 2731, 3869, 1709, 1685, 1714, 1704, 2275])
    wide_doc = Document(str(wide))
    tbl_el = wide_doc.tables[0]._tbl
    docx_copy._normalize_table_width(tbl_el, src_doc)
    from docx.shared import Emu

    sec = src_doc.sections[0]
    content_tw = Emu(sec.page_width - sec.left_margin
                     - sec.right_margin).twips
    widths = _gridcol_widths(tbl_el)
    assert sum(widths) <= content_tw + len(widths)  # 舍入容差
    tbl_w = tbl_el.find(qn("w:tblPr")).find(qn("w:tblW"))
    assert tbl_w.get(qn("w:type")) == "dxa"
    assert int(tbl_w.get(qn("w:w"))) == content_tw
    # 等比：最宽列仍最宽
    assert widths[2] == max(widths)


def test_normalize_keeps_fitting_auto_table(tmp_path: Path):
    """gridCol 合计 < 版心的 auto 表不动。"""
    src = _make_source(tmp_path / "src.docx")
    src_doc = Document(str(src))
    small = _wide_auto_table(tmp_path / "small.docx", [1000, 1000, 1000])
    small_doc = Document(str(small))
    tbl_el = small_doc.tables[0]._tbl
    before = _gridcol_widths(tbl_el)
    docx_copy._normalize_table_width(tbl_el, src_doc)
    assert _gridcol_widths(tbl_el) == before


def test_normalize_keeps_pct_table_within_100(tmp_path: Path):
    """tblW pct ≤ 100% 自适应版心（gridCol 只是提示），不缩放。"""
    from docx.oxml import OxmlElement

    src = _make_source(tmp_path / "src.docx")
    src_doc = Document(str(src))
    pct = _wide_auto_table(tmp_path / "pct.docx", [4000, 4000, 4000])
    pct_doc = Document(str(pct))
    tbl_el = pct_doc.tables[0]._tbl
    tbl_w = tbl_el.find(qn("w:tblPr")).find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "pct")
    tbl_w.set(qn("w:w"), "3000")  # 60%
    before = _gridcol_widths(tbl_el)
    docx_copy._normalize_table_width(tbl_el, src_doc)
    assert _gridcol_widths(tbl_el) == before
