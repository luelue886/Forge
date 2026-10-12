"""真模板构建单测：合成模板全路径（复制/替换/删除/清理/后处理），零 COM 零 LLM。"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.render.template_build import BuildError, build_deck
from app.render.tpl_catalog import extract_catalog
from app.schema.enums import IssueCode, SlideType
from app.schema.slotfill import (
    BuildMeta, DeckBuild, PageFill, TableIR, validate_deck_build,
)
from app.schema.tplcatalog import TemplateCatalog
from tests.mk_tpl import GROUP_PIC_RID, TAG_RID, build_synthetic_template

_SLIDE_PART = re.compile(r"^ppt/slides/slide\d+\.xml$")
_SLIDE_RELS = re.compile(r"^ppt/slides/_rels/slide\d+\.xml\.rels$")


@pytest.fixture(scope="module")
def synthetic_tpl(tmp_path_factory) -> Path:
    return build_synthetic_template(tmp_path_factory.mktemp("build") / "synthetic.pptx")


@pytest.fixture(scope="module")
def catalog(synthetic_tpl: Path) -> TemplateCatalog:
    return extract_catalog(synthetic_tpl, "合成测试模板")


def _deck(pages: list[PageFill]) -> DeckBuild:
    return DeckBuild(meta=BuildMeta(title="测试", template="synthetic"), pages=pages)


def _slot(cat: TemplateCatalog, slide_no: int, role: str, idx: int = 0) -> str:
    return cat.slide(slide_no).slots_of_role(role)[idx].slot_id


def _slide_texts(slide) -> str:
    """含组内子形状的递归全文（python-pptx 的 slide.shapes 不下钻组）。"""
    out: list[str] = []

    def rec(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                rec(sh.shapes)
            elif sh.has_text_frame:
                out.append(sh.text_frame.text)

    rec(slide.shapes)
    return "\n".join(out)


class TestBuildDeck:
    def test_basic_deck(self, synthetic_tpl: Path, catalog: TemplateCatalog, tmp_path: Path):
        t1, sub1 = _slot(catalog, 1, "title"), _slot(catalog, 1, "body")
        t3 = _slot(catalog, 3, "title")
        grp = sorted(s.slot_id for s in catalog.slide(3).slots
                     if s.slot_id.startswith("2."))
        assert len(grp) == 4
        t5, sub5 = _slot(catalog, 5, "title"), _slot(catalog, 5, "body")
        pages = [
            PageFill(page_no=1, slide_type=SlideType.COVER, template_slide=1,
                     title="新主题", slot_values={t1: "新主题", sub1: "全新副标题"}),
            PageFill(page_no=2, slide_type=SlideType.TEXT_POINTS, template_slide=3,
                     title="要点一", slot_values={
                         t3: "要点一", grp[0]: "标签甲", grp[1]: "描述甲",
                         grp[2]: "标签乙", grp[3]: "描述乙"},
                     note="第一页讲者备注"),
            PageFill(page_no=3, slide_type=SlideType.TEXT_POINTS, template_slide=3,
                     title="要点二", slot_values={t3: "要点二"}),
            PageFill(page_no=4, slide_type=SlideType.CLOSING, template_slide=5,
                     title="谢谢观看", slot_values={t5: "谢谢观看", sub5: "落款文本"}),
        ]
        out = tmp_path / "out.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)

        prs = Presentation(str(out))
        assert len(prs.slides) == 4
        assert "新主题" in _slide_texts(prs.slides[0])
        assert "标签甲" in _slide_texts(prs.slides[1])
        assert "要点二" in _slide_texts(prs.slides[2])
        assert "谢谢观看" in _slide_texts(prs.slides[3])
        # 原模板页文案全部退场（年度培训主题/目录条目/谢谢观看之外的模板文）
        all_texts = "\n".join(_slide_texts(prs.slides[i]) for i in range(4))
        assert "年度培训主题" not in all_texts
        assert "第1章节" not in all_texts
        # 备注：python-pptx 后处理写入
        assert prs.slides[1].notes_slide.notes_text_frame.text == "第一页讲者备注"
        # 中文 ea 字体兜底：占位符 run 原本无 a:ea
        assert 'typeface="微软雅黑"' in prs.slides[0].shapes[0]._element.xml

        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            assert len([n for n in names if _SLIDE_PART.match(n)]) == 4
            # S3 复制两次：共享 media 存活
            assert any(n.startswith("ppt/media/") for n in names)

    def test_multi_paragraph_fill(self, synthetic_tpl: Path, catalog: TemplateCatalog,
                                  tmp_path: Path):
        t2, b2 = _slot(catalog, 2, "title"), _slot(catalog, 2, "body")
        pages = [PageFill(page_no=1, slide_type=SlideType.TOC, template_slide=2,
                          title="目录", slot_values={t2: "目录", b2: "一\n二\n三"})]
        out = tmp_path / "out.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)
        prs = Presentation(str(out))
        tf = next(sh.text_frame for sh in prs.slides[0].shapes
                  if sh.has_text_frame and sh.text_frame.text.startswith("一"))
        assert [p.text for p in tf.paragraphs] == ["一", "二", "三"]

    def test_group_collapse_and_rels_prune(self, synthetic_tpl: Path,
                                           catalog: TemplateCatalog, tmp_path: Path):
        t3 = _slot(catalog, 3, "title")
        grp = sorted(s.slot_id for s in catalog.slide(3).slots
                     if s.slot_id.startswith("2."))
        pages = [PageFill(page_no=1, slide_type=SlideType.TEXT_POINTS, template_slide=3,
                          title="x", slot_values={t3: "x"}, slot_deletes=grp)]
        out = tmp_path / "out.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)

        prs = Presentation(str(out))
        texts = _slide_texts(prs.slides[0])
        assert "团队协作" not in texts and "数据驱动" not in texts
        # 组整体删除后仅剩标题 + 顶层图片
        assert len(prs.slides[0].shapes) == 2

        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            # 组内图片的 rId 已随组删除而剪除；顶层图片的 media 仍存活
            rels = [n for n in names if _SLIDE_RELS.match(n)]
            assert len(rels) == 1
            content = z.read(rels[0]).decode("utf-8")
            assert GROUP_PIC_RID not in content
            # notesSlide 关系在复制时剥离
            assert "notesSlide" not in content
            assert any(n.startswith("ppt/media/") for n in names)

    def test_orphan_cleanup(self, synthetic_tpl: Path, catalog: TemplateCatalog,
                            tmp_path: Path):
        """只用 S1：S3 的图片/备注部件全部变孤儿清除。"""
        pages = [PageFill(page_no=1, slide_type=SlideType.COVER, template_slide=1,
                          title="t", slot_values={
                              _slot(catalog, 1, "title"): "t",
                              _slot(catalog, 1, "body"): "s"})]
        out = tmp_path / "out.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)
        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            assert len([n for n in names if _SLIDE_PART.match(n)]) == 1
            assert not any(n.startswith("ppt/media/") for n in names)
            assert not any("notesSlides" in n for n in names)
            # notesMaster 是 presentation 关系，保留
            assert any("notesMaster" in n for n in names)

    def test_tag_part_stripped_on_duplicate(self, synthetic_tpl: Path,
                                            catalog: TemplateCatalog, tmp_path: Path):
        """S3 带 tags 部件复制到 2 页：custDataLst / tags rels / tag 部件全剥离。

        多页共享同一 tag 部件会让 PowerPoint 拒绝打开文件——真模板 e2e
        的实弹根因，此处为回归防线。
        """
        title3 = _slot(catalog, 3, "title")
        pages = [
            PageFill(page_no=1, slide_type=SlideType.COVER, template_slide=1,
                     title="测试", slot_values={_slot(catalog, 1, "title"): "测试"}),
            PageFill(page_no=2, slide_type=SlideType.TWO_COLUMN, template_slide=3,
                     title="内容甲", slot_values={title3: "内容甲"}),
            PageFill(page_no=3, slide_type=SlideType.TWO_COLUMN, template_slide=3,
                     title="内容乙", slot_values={title3: "内容乙"}),
        ]
        out = tmp_path / "deck.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)

        with zipfile.ZipFile(out) as z:
            names = z.namelist()
            slide_parts = [n for n in names if _SLIDE_PART.match(n)]
            assert len(slide_parts) == 3  # 3 复制页，原 5 页已删
            for part in slide_parts:
                blob = z.read(part)
                assert b"custDataLst" not in blob, part
                rels = z.read(part.replace("slides/", "slides/_rels/") + ".rels")
                assert b"/tags" not in rels, part
                assert TAG_RID.encode() not in rels, part
            assert not any(n.startswith("ppt/tags/") for n in names)
            ct = z.read("[Content_Types].xml")
            assert b"/ppt/tags/tag1.xml" not in ct  # Override 随孤儿清理同步移除

        prs = Presentation(str(out))  # 剥离后 python-pptx 正常打开
        assert len(prs.slides) == 3

    def test_timing_stripped_on_duplicate(self, synthetic_tpl: Path,
                                           catalog: TemplateCatalog, tmp_path: Path):
        """S3 带 p:timing（spTgt spid 指向组内形状）复制后整树剥离。

        删形状（槽删除/组塌缩）会让 timing 的 spid 悬空，PowerPoint 拒开
        文件——真模板 e2e 的实弹根因（bisect：坏页全带悬空 spid），回归防线。
        """
        with zipfile.ZipFile(synthetic_tpl) as z:
            assert b"p:timing" in z.read("ppt/slides/slide3.xml")  # fixture 守卫

        title3 = _slot(catalog, 3, "title")
        grp = sorted(s.slot_id for s in catalog.slide(3).slots
                     if s.slot_id.startswith("2."))
        pages = [
            PageFill(page_no=1, slide_type=SlideType.TEXT_POINTS, template_slide=3,
                     title="x", slot_values={title3: "x"}, slot_deletes=grp),
            PageFill(page_no=2, slide_type=SlideType.TEXT_POINTS, template_slide=3,
                     title="y", slot_values={title3: "y"}),
        ]
        out = tmp_path / "deck.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)

        with zipfile.ZipFile(out) as z:
            for part in (n for n in z.namelist() if _SLIDE_PART.match(n)):
                assert b"p:timing" not in z.read(part), part
        prs = Presentation(str(out))
        assert len(prs.slides) == 2

    def test_native_table(self, synthetic_tpl: Path, catalog: TemplateCatalog,
                          tmp_path: Path):
        b4 = _slot(catalog, 4, "body")
        pages = [PageFill(page_no=1, slide_type=SlideType.TABLE, template_slide=4,
                          title="表", slot_values={b4: "表格页标题"},
                          table=TableIR(table_id="t0", header=["列甲", "列乙"],
                                        rows=[["1", "2"], ["3", "4"]]))]
        out = tmp_path / "out.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)
        prs = Presentation(str(out))
        slide = prs.slides[0]
        gfs = [sh for sh in slide.shapes if sh.has_table]
        assert len(gfs) == 1
        tbl = gfs[0].table
        assert tbl.cell(0, 0).text == "列甲"
        assert tbl.cell(2, 1).text == "4"
        # 表头底色 = theme accent1（python-pptx 默认主题 4F81BD）
        assert tbl.cell(0, 0).fill.fore_color.rgb == RGBColor(0x4F, 0x81, 0xBD)
        # 表格锚在标题槽下方，不重叠
        title_box = next(sh for sh in slide.shapes if sh.has_text_frame
                         and sh.text_frame.text == "表格页标题")
        assert gfs[0].top >= title_box.top + title_box.height

    def test_noteslide_stripped_and_note_set(self, synthetic_tpl: Path,
                                             catalog: TemplateCatalog, tmp_path: Path):
        t3 = _slot(catalog, 3, "title")
        pages = [PageFill(page_no=1, slide_type=SlideType.TEXT_POINTS, template_slide=3,
                          title="x", slot_values={t3: "x"}, note="讲者备注")]
        out = tmp_path / "out.pptx"
        build_deck(synthetic_tpl, catalog, _deck(pages), out)
        with zipfile.ZipFile(out) as z:
            blob = b"".join(z.read(n) for n in z.namelist())
            assert "模板原始备注".encode("utf-8") not in blob
        prs = Presentation(str(out))
        assert prs.slides[0].notes_slide.notes_text_frame.text == "讲者备注"

    def test_unknown_slot_error(self, synthetic_tpl: Path, catalog: TemplateCatalog,
                                tmp_path: Path):
        pages = [PageFill(page_no=1, slide_type=SlideType.COVER, template_slide=1,
                          title="t", slot_values={"99": "x"})]
        with pytest.raises(BuildError, match="未知槽"):
            build_deck(synthetic_tpl, catalog, _deck(pages), tmp_path / "out.pptx")

    def test_slot_drift_error(self, synthetic_tpl: Path, catalog: TemplateCatalog,
                              tmp_path: Path):
        bad = catalog.model_copy(deep=True)
        s3 = bad.slide(3)
        s3.slots = s3.slots[:-1]
        pages = [PageFill(page_no=1, slide_type=SlideType.TEXT_POINTS, template_slide=3,
                          title="x", slot_values={_slot(catalog, 3, "title"): "x"})]
        with pytest.raises(BuildError, match="不符"):
            build_deck(synthetic_tpl, bad, _deck(pages), tmp_path / "out.pptx")

    def test_template_mismatch_error(self, synthetic_tpl: Path, catalog: TemplateCatalog,
                                     tmp_path: Path):
        deck = DeckBuild(meta=BuildMeta(title="t", template="other"), pages=[
            PageFill(page_no=1, slide_type=SlideType.COVER, template_slide=1, title="t")])
        with pytest.raises(BuildError, match="不一致"):
            build_deck(synthetic_tpl, catalog, deck, tmp_path / "out.pptx")


class TestValidateDeckBuild:
    def test_ok(self, catalog: TemplateCatalog):
        deck = _deck([PageFill(page_no=1, slide_type=SlideType.COVER,
                               template_slide=1, title="t")])
        assert validate_deck_build(deck, catalog) == []

    def test_page_no_gap(self, catalog: TemplateCatalog):
        deck = _deck([
            PageFill(page_no=2, slide_type=SlideType.COVER, template_slide=1, title="a"),
            PageFill(page_no=3, slide_type=SlideType.COVER, template_slide=1, title="b"),
        ])
        issues = validate_deck_build(deck, catalog)
        assert any(i.rule == IssueCode.V_PAGE_NO for i in issues)

    def test_template_slide_out_of_range(self, catalog: TemplateCatalog):
        deck = _deck([PageFill(page_no=1, slide_type=SlideType.COVER,
                               template_slide=99, title="t")])
        issues = validate_deck_build(deck, catalog)
        assert any(i.rule == IssueCode.V_TEMPLATE_SLIDE for i in issues)
