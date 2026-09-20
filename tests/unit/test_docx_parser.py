from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from app.parsing.docx_parser import parse_docx
from app.parsing.base import ParseError


def test_basic_structure(basic_docx):
    tree = parse_docx(basic_docx)
    assert tree.meta.source_format == "docx"
    assert tree.meta.source_name == "basic.docx"
    assert tree.meta.n_chars > 0
    assert tree.meta.parse_warnings == []

    root = tree.sections[0]
    assert root.title == "智慧园区平台建设汇报"
    assert [s.title for s in root.subsections] == ["一、项目概述", "二、季度指标"]
    overview, metrics = root.subsections

    kinds = [b.kind for b in overview.blocks]
    assert kinds == ["para", "list_item", "list_item"]
    assert overview.blocks[0].text == "本项目覆盖 12 个园区，接入设备 8,600 台。"

    # 表格块 + 收尾段落在第二个 section
    assert [b.kind for b in metrics.blocks] == ["para", "table", "para"]
    assert metrics.blocks[1].table_id == "tbl-001"


def test_table_verbatim_and_caption(basic_docx):
    tree = parse_docx(basic_docx)
    t = tree.tables[0]
    assert t.table_id == "tbl-001"
    assert t.header == ["指标", "目标", "实际"]
    assert t.rows == [["园区数", "10", "12"], ["设备数", "8000", "8600"]]
    assert t.n_rows == 3 and t.n_cols == 3
    assert t.caption == "表1：Q3 关键指标"
    assert t.section_id == "sec-0002"
    # full_text 必须包含表格原文（数字溯源基准）
    assert "8600" in tree.full_text
    # 渲染对齐锚点：body 首个表格 + 归一化列宽
    assert t.src_index == 0
    assert t.col_widths and len(t.col_widths) == 3
    assert abs(sum(t.col_widths) - 1.0) < 1e-6


def test_resolve_refs(basic_docx):
    tree = parse_docx(basic_docx)
    refs = tree.known_refs()
    assert {"sec-0000", "sec-0001", "sec-0002", "tbl-001"} <= refs
    assert any(r.startswith("blk-") for r in refs)
    assert "8,600" in tree.resolve(["blk-0001"])
    assert "指标" in tree.resolve(["tbl-001"])
    assert "一、项目概述" in tree.resolve(["sec-0001"])
    assert tree.resolve(["不存在"]) == ""


def test_no_headings(tmp_path):
    doc = Document()
    doc.add_paragraph("第一段，含数字 42。")
    doc.add_paragraph("第二段。")
    p = tmp_path / "flat.docx"
    doc.save(str(p))
    tree = parse_docx(p)
    root = tree.sections[0]
    assert root.subsections == []
    assert [b.kind for b in root.blocks] == ["para", "para"]
    assert "42" in tree.full_text


def test_table_without_header(tmp_path):
    doc = Document()
    doc.add_paragraph("无表头表：")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = ""
    t.rows[0].cells[1].text = "数据列"
    t.rows[1].cells[0].text = "甲"
    t.rows[1].cells[1].text = "乙"
    p = tmp_path / "noheader.docx"
    doc.save(str(p))
    tree = parse_docx(p)
    t = tree.tables[0]
    assert t.header == []
    assert len(t.rows) == 2
    assert any("表头" in w for w in tree.meta.parse_warnings)


def test_merged_cells_warning(tmp_path):
    doc = Document()
    doc.add_paragraph("表2：合并示例")
    t = doc.add_table(rows=2, cols=3)
    t.rows[0].cells[0].text = "维度"
    t.rows[0].cells[1].text = "数值"
    t.rows[0].cells[2].text = "备注"
    t.rows[1].cells[0].text = "整体"
    t.rows[1].cells[1].text = "99"
    t.rows[1].cells[2].text = "-"
    t.cell(0, 0).merge(t.cell(0, 1))  # 横向合并表头前两格
    p = tmp_path / "merged.docx"
    doc.save(str(p))
    tree = parse_docx(p)
    assert any("合并单元格" in w for w in tree.meta.parse_warnings)
    # 合并展开为重复值，值本身 verbatim
    assert tree.tables[0].rows == [["整体", "99", "-"]]


def test_missing_file(tmp_path):
    with pytest.raises(ParseError, match="不存在"):
        parse_docx(tmp_path / "nope.docx")


# ---- 旧版 .doc 入口（Word COM 转换 → 持久 converted.docx）----

def _fake_doc_bytes() -> bytes:
    import io

    doc = Document()
    doc.add_paragraph("一、总体情况")
    doc.add_paragraph("全年营收 1,234.56 万元。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_source_legacy_doc_persists_beside_source(tmp_path, monkeypatch):
    """转换产物落源文件旁（渲染要搬源 XML），幂等：已存在则不再调 COM。"""
    import app.services.com_export as ce
    from app.parse_dispatch import parse_source

    fake_doc = tmp_path / "旧版公文.doc"
    fake_doc.write_bytes(b"not a real doc")
    calls = {"n": 0}

    def fake_convert(p, out=None):
        calls["n"] += 1
        dest = Path(out) if out else Path(p).with_suffix(".docx")
        dest.write_bytes(_fake_doc_bytes())
        return dest

    monkeypatch.setattr(ce, "convert_doc_to_docx", fake_convert)

    tree = parse_source(fake_doc)
    assert calls["n"] == 1
    converted = tmp_path / "旧版公文.converted.docx"
    assert converted.exists()
    assert any(".doc" in w for w in tree.meta.parse_warnings)
    assert tree.meta.source_name == "旧版公文.doc"  # 展示原名，不露 converted
    assert tree.meta.source_format == "docx"
    assert "1,234.56" in tree.full_text

    # 二次解析（断点续跑场景）：converted.docx 已在 → 零 COM 调用
    tree2 = parse_source(fake_doc)
    assert calls["n"] == 1
    assert tree2.full_text == tree.full_text


def test_parse_source_legacy_doc_convert_failure(tmp_path, monkeypatch):
    """转换失败（如无 Word）→ ParseError，而非裸 COM 异常。"""
    import app.services.com_export as ce
    from app.parse_dispatch import parse_source

    fake_doc = tmp_path / "broken.doc"
    fake_doc.write_bytes(b"not a real doc")
    monkeypatch.setattr(ce, "convert_doc_to_docx",
                        lambda p, out=None: (_ for _ in ()).throw(RuntimeError("COM 拒绝")))

    with pytest.raises(ParseError, match="Word"):
        parse_source(fake_doc)


def test_ids_stable_across_parses(basic_docx):
    t1, t2 = parse_docx(basic_docx), parse_docx(basic_docx)
    assert t1.known_refs() == t2.known_refs()
    assert t1.full_text == t2.full_text


# ---- C4: 图片/流程图 ----

import base64
import io

_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")


def _doc_with_image(path) -> Path:
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("一、总体情况")
    doc.add_paragraph("流程如下：")
    doc.add_picture(io.BytesIO(_PNG_1PX), width=Cm(5))
    doc.save(str(path))
    return path


def test_parse_inline_image(tmp_path):
    from docx.oxml.ns import qn

    src = _doc_with_image(tmp_path / "img.docx")
    tree = parse_docx(src)
    assert len(tree.images) == 1
    img = tree.images[0]
    assert img.image_id == "img-001"
    assert img.section_id == "sec-0000"
    assert img.cx_emu == 1800000 and img.cy_emu == 1800000  # Cm(5)，1:1 图
    # body_index 与渲染期同一坐标系：指向图片所在段落
    children = list(Document(str(src)).element.body.iterchildren())
    assert children[img.body_index].find(".//" + qn("w:drawing")) is not None
    # 结构块收录 + 图片无文本进 full_text
    img_blk = next(b for b in tree.sections[0].blocks if b.kind == "image")
    assert img_blk.image_id == "img-001" and img_blk.text == ""
    assert tree.full_text == "一、总体情况\n流程如下："


def test_parse_smartart_skipped(tmp_path):
    from lxml import etree

    src = _doc_with_image(tmp_path / "smart.docx")
    doc = Document(str(src))
    drawing = doc.element.body.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}drawing")
    gd = drawing.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}graphicData")
    etree.SubElement(
        gd, "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds")
    doc.save(str(src))

    tree = parse_docx(src)
    assert tree.images == []
    assert any("SmartArt" in w for w in tree.meta.parse_warnings)


def test_parse_two_drawings_in_one_paragraph(tmp_path):
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("流程：")
    doc.add_picture(io.BytesIO(_PNG_1PX), width=Cm(4))
    doc.add_picture(io.BytesIO(_PNG_1PX), width=Cm(4))
    p1, p2 = doc.paragraphs[-2], doc.paragraphs[-1]
    p1._p.append(p2.runs[0]._r)  # 两张图挤进同一段
    src = tmp_path / "two.docx"
    doc.save(str(src))

    tree = parse_docx(src)
    assert len(tree.images) == 1
    assert any("多个图形" in w for w in tree.meta.parse_warnings)


def test_converted_suffix_stripped_from_title(tmp_path):
    # .doc 经 COM 转换产物名为 *.converted.docx：标题回退 stem 时剥掉尾巴
    doc = Document()
    doc.add_paragraph("个人简历模板第一段说明文字，用于排版展示。")
    p = tmp_path / "简历模板.converted.docx"
    doc.save(str(p))
    tree = parse_docx(p)
    assert tree.sections[0].title == "简历模板"
