from __future__ import annotations

from pathlib import Path

import pytest

from app.parsing.base import ParseError
from app.parsing.pptx_parser import parse_pptx


def _bulletize(p):
    """给段落加 a:buChar，模拟真实 bullet 段。"""
    from pptx.oxml.ns import qn

    pPr = p._p.get_or_add_pPr()
    bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu)


@pytest.fixture(scope="module")
def sample_pptx(tmp_path_factory) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    layout_title = prs.slide_layouts[0]
    layout_content = prs.slide_layouts[1]
    layout_title_only = prs.slide_layouts[5]

    # 封面：标题 + 副标题
    s1 = prs.slides.add_slide(layout_title)
    s1.shapes.title.text_frame.text = "产品发布汇报"
    s1.placeholders[1].text_frame.text = "2026 年秋季"

    # 章节头：只有标题
    s2 = prs.slides.add_slide(layout_title_only)
    s2.shapes.title.text_frame.text = "一、市场回顾"

    # 内容页：标题 + 两段（其一为 bullet）+ 3×3 表格
    s3 = prs.slides.add_slide(layout_content)
    s3.shapes.title.text_frame.text = "市场数据"
    tf = s3.placeholders[1].text_frame
    tf.text = "本季度新增客户 128 家。"
    p2 = tf.add_paragraph()
    p2.text = "续约率达到 91%。"
    _bulletize(p2)
    s3.shapes.add_table(3, 3, Inches(1), Inches(4), Inches(8), Inches(2)).table
    tbl = s3.shapes[-1].table
    for j, h in enumerate(["渠道", "数量", "占比"]):
        tbl.rows[0].cells[j].text = h
    for i, row in enumerate([["直销", "56", "43.8%"], ["渠道", "72", "56.2%"]]):
        for j, v in enumerate(row):
            tbl.rows[i + 1].cells[j].text = v

    # 第二个内容页：归属同一章节
    s4 = prs.slides.add_slide(layout_content)
    s4.shapes.title.text_frame.text = "竞品对比"
    s4.placeholders[1].text_frame.text = "主要竞品三家，市占率合计 62.4%。"

    # 空白页（无任何文本）：应被跳过
    prs.slides.add_slide(prs.slide_layouts[6])

    p = tmp_path_factory.mktemp("pptx") / "sample.pptx"
    prs.save(str(p))
    return p


def test_parse_pptx_structure(sample_pptx):
    tree = parse_pptx(sample_pptx)

    assert tree.meta.source_format == "pptx"
    assert tree.meta.source_name == "sample.pptx"
    assert tree.sections[0].title == "产品发布汇报"  # 封面标题 → 文档题名

    # 章节：1 个显式章节头；无隐式章节（内容页都挂在章节头下）
    l1 = [s for s in tree.sections[0].walk() if s.level == 1]
    assert [s.title for s in l1] == ["一、市场回顾"]

    blocks = [b for s in l1 for b in s.blocks]
    kinds = [(b.kind, b.text) for b in blocks]
    # 内容页标题降级为 para；bullet 段识别为 list_item
    assert ("para", "市场数据") in kinds
    assert ("para", "本季度新增客户 128 家。") in kinds
    assert ("list_item", "续约率达到 91%。") in kinds
    assert ("para", "竞品对比") in kinds
    # 空白页被跳过：块总数 = 2 页标题 + 2 段 + 1 表 + 1 竞品段
    assert len(blocks) == 6
    # 封面副标题落在根节
    root_blocks = [b.text for b in tree.sections[0].blocks]
    assert root_blocks == ["2026 年秋季"]


def test_parse_pptx_table_verbatim(sample_pptx):
    tree = parse_pptx(sample_pptx)

    assert len(tree.tables) == 1
    t = tree.tables[0]
    assert t.n_rows == 3 and t.n_cols == 3
    assert t.header == ["渠道", "数量", "占比"]
    assert t.rows == [["直销", "56", "43.8%"], ["渠道", "72", "56.2%"]]
    assert t.caption == "市场数据"  # 表格所在页标题作 caption
    # full_text 覆盖表格与全部正文（数字溯源基准）
    for needle in ["128 家", "91%", "62.4%", "56.2%", "2026 年秋季"]:
        assert needle in tree.full_text
    # 引用系统可用
    assert {t.table_id} <= tree.known_refs()
    assert "直销" in tree.resolve([t.table_id])


def test_parse_pptx_missing_file(tmp_path):
    with pytest.raises(ParseError):
        parse_pptx(tmp_path / "nope.pptx")


def test_parse_source_legacy_ppt(sample_pptx, tmp_path, monkeypatch):
    """旧版 .ppt：入口先转换再按 pptx 解析，并落 parse_warning。"""
    import app.services.com_export as ce
    from app.parse_dispatch import parse_source

    fake_ppt = tmp_path / "旧版课件.ppt"
    fake_ppt.write_bytes(sample_pptx.read_bytes())

    def fake_convert(p, out=None):
        dest = Path(out) if out else Path(p).with_suffix(".pptx")
        dest.write_bytes(Path(p).read_bytes())
        return dest

    monkeypatch.setattr(ce, "convert_to_pptx", fake_convert)

    tree = parse_source(fake_ppt)
    assert any("ppt" in w for w in tree.meta.parse_warnings)
    assert "产品发布汇报" in tree.full_text


def test_parse_source_legacy_ppt_convert_failure(tmp_path, monkeypatch):
    """转换失败（如无 PowerPoint）→ ParseError，而非裸 COM 异常。"""
    import app.services.com_export as ce
    from app.parse_dispatch import parse_source

    fake_ppt = tmp_path / "broken.ppt"
    fake_ppt.write_bytes(b"not a real ppt")
    monkeypatch.setattr(ce, "convert_to_pptx",
                        lambda p, out=None: (_ for _ in ()).throw(RuntimeError("COM 拒绝")))

    with pytest.raises(ParseError, match="PowerPoint"):
        parse_source(fake_ppt)
