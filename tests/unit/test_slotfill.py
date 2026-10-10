"""M5/N2 槽填充：结构页 LLM 主路径 + 确定性兜底、表格页零 LLM、
内容页 LLM 提案→校验→重试→降级、槽级校验链。"""
from __future__ import annotations

import json

import pytest

from app.pipeline.slotfill import (
    FillError, NOTE_WEIGHT_MAX, slotfill_deterministic, slotfill_page,
    validate_page_fill,
)
from app.render.tpl_catalog import extract_catalog
from app.schema.enums import IssueCode, SlideType
from app.schema.plan import SlidePlanItem
from app.schema.slotfill import PageFill
from app.schema.tplcatalog import SlideInfo, TemplateCatalog
from tests.mk_tpl import build_synthetic_template


@pytest.fixture(scope="module")
def catalog(tmp_path_factory) -> TemplateCatalog:
    tpl = build_synthetic_template(tmp_path_factory.mktemp("tpl") / "synthetic.pptx")
    return extract_catalog(tpl, "合成测试模板")


@pytest.fixture(scope="module")
def tree(basic_docx):
    from app.parsing.docx_parser import parse_docx

    return parse_docx(basic_docx)


def _item(no: int, t: SlideType, title: str = "页", brief: str = "",
          refs: list[str] | None = None) -> SlidePlanItem:
    return SlidePlanItem(page_no=no, slide_type=t, title=title, brief=brief,
                         src_refs=refs or [])


def _s3_ids(catalog: TemplateCatalog) -> dict[str, str]:
    """S3 的槽位：title=0，组内 2.0~2.3。"""
    s3 = catalog.slide(3)
    out = {"title": next(s.slot_id for s in s3.slots if s.role == "title")}
    out["grp"] = sorted(s.slot_id for s in s3.slots if s.slot_id.startswith("2."))
    return out


def _good_out(ids: dict) -> str:    return json.dumps({
        "title": "覆盖规模与设备接入",
        "slots": {
            ids["title"]: "覆盖规模与设备接入",
            ids["grp"][0]: "园区覆盖",
            ids["grp"][1]: "十二个园区全面接入",
            ids["grp"][2]: "设备规模",
            ids["grp"][3]: "八千六百台设备在线",
        },
        "deletes": [],
        "note": "本页展开园区覆盖与设备接入规模。",
    }, ensure_ascii=False)


def _bad_out(title_only: dict) -> str:
    """只填 title（覆盖缺失）→ V_SLOT_COVER。"""
    return json.dumps({
        "title": "覆盖规模",
        "slots": {title_only["title"]: "覆盖规模"},
        "deletes": [], "note": None,
    }, ensure_ascii=False)


def _num_out(ids: dict) -> str:
    """坏输出：描述槽带未溯源数字 1200（源文数字集无 1200，且 >99 不豁免）。"""
    out = json.loads(_good_out(ids))
    out["slots"][ids["grp"][1]] = "设备规模突破 1200 万台"
    return json.dumps(out, ensure_ascii=False)


def _plag_out(ids: dict) -> str:
    """坏输出：描述槽 verbatim 抄源文（数字 12 可溯源，隔离出抄袭违规）。"""
    out = json.loads(_good_out(ids))
    out["slots"][ids["grp"][1]] = "本项目覆盖 12 个园区，接入设备"
    return json.dumps(out, ensure_ascii=False)


class TestDeterministic:
    """结构页/表格页零 LLM（函数签名不收 client）。"""

    def test_cover(self, catalog: TemplateCatalog):
        item = _item(1, SlideType.COVER, "智慧园区汇报", brief="副题文本")
        s1 = catalog.slide(1)
        fill = slotfill_deterministic(item, None, s1)  # type: ignore[arg-type]
        title = next(s.slot_id for s in s1.slots if s.role == "title")
        assert fill.slot_values[title] == "智慧园区汇报"
        # 次大可填槽 ← 副题；其余非 design 槽全删
        assert len(fill.slot_values) == 2
        assert "副题文本" in fill.slot_values.values()
        assert set(fill.slot_deletes) == {
            s.slot_id for s in s1.slots
            if s.role not in ("design", "page_no")} - set(fill.slot_values)
        assert fill.template_slide == 1

    def test_toc_merges_entries(self, catalog: TemplateCatalog):
        item = _item(2, SlideType.TOC, "目录")
        s2 = catalog.slide(2)
        body = next(s.slot_id for s in s2.slots if s.role == "body")
        fill = slotfill_deterministic(item, None, s2, toc_entries=["一、概述", "二、指标"])  # type: ignore[arg-type]
        title = next(s.slot_id for s in s2.slots if s.role == "title")
        assert fill.slot_values[title] == "目录"
        # 单 body 槽：条目按段合并（超出并入末槽）
        assert fill.slot_values[body] == "一、概述\n二、指标"

    def test_section_deletes_rest(self, catalog: TemplateCatalog):
        item = _item(3, SlideType.SECTION_HEADER, "一、项目概述")
        s3 = catalog.slide(3)
        ids = _s3_ids(catalog)
        fill = slotfill_deterministic(item, None, s3)  # type: ignore[arg-type]
        assert fill.slot_values == {ids["title"]: "一、项目概述"}
        assert fill.slot_deletes == ids["grp"]

    def test_closing(self, catalog: TemplateCatalog):
        item = _item(5, SlideType.CLOSING, "谢谢观看")
        s5 = catalog.slide(5)
        fill = slotfill_deterministic(item, None, s5, doc_title="智慧园区汇报")  # type: ignore[arg-type]
        assert "谢谢观看" in fill.slot_values.values()
        assert "智慧园区汇报" in fill.slot_values.values()

    def test_table_verbatim(self, catalog: TemplateCatalog, tree):
        item = _item(4, SlideType.TABLE, "Q3 指标", refs=["tbl-001"])
        from app.schema.plan import TableRef
        from app.schema.enums import TableStrategy
        item.table_ref = TableRef(table_id="tbl-001", strategy=TableStrategy.COPY)
        s4 = catalog.slide(4)
        fill = slotfill_deterministic(item, tree, s4)
        assert fill.table is not None
        assert fill.table.header == ["指标", "目标", "实际"]
        assert fill.table.rows == [["园区数", "10", "12"], ["设备数", "8000", "8600"]]
        assert fill.note is None  # 3×3 未超 5×6，不裁剪
        # S4 无 title 槽：计划题落最大字号可填槽
        assert "Q3 指标" in fill.slot_values.values()

    def test_table_trims_oversize(self, catalog: TemplateCatalog, tree, tmp_path):
        from docx import Document

        doc = Document()
        doc.add_heading("大表文档", 0)
        doc.add_paragraph("前置段落。")
        t = doc.add_table(rows=9, cols=8)
        for j in range(8):
            t.rows[0].cells[j].text = f"列{j}"
        for i in range(1, 9):
            for j in range(8):
                t.rows[i].cells[j].text = f"v{i}-{j}"
        p = tmp_path / "big.docx"
        doc.save(str(p))
        from app.parsing.docx_parser import parse_docx
        from app.schema.enums import TableStrategy
        from app.schema.plan import TableRef

        big_tree = parse_docx(p)
        item = _item(1, SlideType.TABLE, "大表", refs=[big_tree.tables[0].table_id])
        item.table_ref = TableRef(table_id=big_tree.tables[0].table_id,
                                  strategy=TableStrategy.COPY)
        fill = slotfill_deterministic(item, big_tree, catalog.slide(4))
        assert fill.table is not None
        assert len(fill.table.rows) == 4  # 5 行含表头 - 1
        assert len(fill.table.header) == 6
        assert fill.note and "超出单页容量" in fill.note

    def test_no_fillable_slot_raises(self, catalog: TemplateCatalog):
        empty = SlideInfo(slide_no=9, layout="", signature="", slots=[],
                          fillable=0, fillable_cap=0.0, role_hint="content")
        item = _item(1, SlideType.COVER, "t")
        with pytest.raises(FillError, match="完全无可用槽"):
            slotfill_deterministic(item, None, empty)  # type: ignore[arg-type]


def _s1_ids(catalog: TemplateCatalog) -> dict[str, str]:
    """S1（cover）的槽位：title ph + subtitle ph。"""
    s1 = catalog.slide(1)
    out = {"title": next(s.slot_id for s in s1.slots if s.role == "title")}
    rest = [s.slot_id for s in s1.slots
            if s.role not in ("title", "design", "page_no")]
    out["sub"] = rest[0]
    return out


def _s4_taglines(catalog: TemplateCatalog) -> list[str]:
    return [s.slot_id for s in catalog.slide(4).slots_of_role("tagline")]


class TestStructuralLLM:
    """N2：结构页走 LLM 主路径，失败兜底确定性；表格页仍零 LLM。"""

    def test_cover_via_llm(self, catalog: TemplateCatalog, tree, fake_llm_factory):
        ids = _s1_ids(catalog)
        out = json.dumps({
            "title": "智慧园区汇报",
            "slots": {ids["title"]: "智慧园区汇报", ids["sub"]: "平台建设全景"},
            "deletes": [], "note": None,
        }, ensure_ascii=False)
        client, fake = fake_llm_factory([out])
        item = _item(1, SlideType.COVER, "智慧园区汇报", brief="副题")
        fill = slotfill_page(item, tree, catalog.slide(1), client,
                             doc_title="智慧园区平台建设汇报")
        assert len(fake.calls) == 1
        assert fill.slot_values[ids["title"]] == "智慧园区汇报"
        # doc_title 进 cover 提示
        assert "智慧园区平台建设汇报" in fake.calls[0]["messages"][1]["content"]

    def test_toc_via_llm_passes_entries(self, catalog: TemplateCatalog, tree,
                                        fake_llm_factory):
        s2 = catalog.slide(2)
        title = next(s.slot_id for s in s2.slots if s.role == "title")
        body = next(s.slot_id for s in s2.slots if s.role == "body")
        out = json.dumps({
            "title": "目录", "slots": {title: "目录", body: "一、概述\n二、指标"},
            "deletes": [], "note": None,
        }, ensure_ascii=False)
        client, fake = fake_llm_factory([out])
        item = _item(2, SlideType.TOC, "目录")
        fill = slotfill_page(item, tree, s2, client,
                             toc_entries=["一、概述", "二、指标"])
        assert len(fake.calls) == 1
        user = fake.calls[0]["messages"][1]["content"]
        assert "目录条目" in user and "二、指标" in user
        assert fill.slot_values[body] == "一、概述\n二、指标"

    def test_structural_skips_ngram(self, catalog: TemplateCatalog, tree,
                                    fake_llm_factory):
        """结构页语义复检放宽：副题 verbatim 回源文（数字可溯源）不算抄袭。"""
        ids = _s1_ids(catalog)
        src = tree.resolve(["blk-0001"])  # "本项目覆盖 12 个园区，接入设备 8,600 台。"
        out = json.dumps({
            "title": "智慧园区汇报",
            "slots": {ids["title"]: "智慧园区汇报", ids["sub"]: src},
            "deletes": [], "note": None,
        }, ensure_ascii=False)
        client, fake = fake_llm_factory([out])
        item = _item(1, SlideType.COVER, "智慧园区汇报")
        fill = slotfill_page(item, tree, catalog.slide(1), client)
        assert len(fake.calls) == 1  # 未因 ngram 拒收重试
        assert fill.slot_values[ids["sub"]] == src

    def test_structural_fallback_after_two_rejections(self, catalog, tree,
                                                      fake_llm_factory):
        """两轮校验未过 → 确定性兜底（title←计划题，次大槽←brief）。"""
        ids = _s1_ids(catalog)
        bad = json.dumps({"title": "t", "slots": {}, "deletes": [],
                          "note": None}, ensure_ascii=False)
        client, fake = fake_llm_factory([bad, bad])
        item = _item(1, SlideType.COVER, "智慧园区汇报", brief="副题文本")
        fill = slotfill_page(item, tree, catalog.slide(1), client)
        assert len(fake.calls) == 2
        assert fill.slot_values[ids["title"]] == "智慧园区汇报"
        assert "副题文本" in fill.slot_values.values()

    def test_structural_fallback_on_llm_error(self, catalog, tree,
                                              fake_llm_factory):
        """LLM 硬异常不再杀 job：结构页兜底确定性填充。"""
        ids = _s1_ids(catalog)
        client, fake = fake_llm_factory([RuntimeError("endpoint down")])
        item = _item(1, SlideType.COVER, "智慧园区汇报", brief="副题文本")
        fill = slotfill_page(item, tree, catalog.slide(1), client)
        assert fill.slot_values[ids["title"]] == "智慧园区汇报"

    def test_table_still_deterministic(self, catalog: TemplateCatalog, tree,
                                       fake_llm_factory):
        """表格页零 LLM：slotfill_page 不消耗任何调用，表内容 verbatim。"""
        from app.schema.enums import TableStrategy
        from app.schema.plan import TableRef

        client, fake = fake_llm_factory([])
        item = _item(4, SlideType.TABLE, "Q3 指标", refs=["tbl-001"])
        item.table_ref = TableRef(table_id="tbl-001", strategy=TableStrategy.COPY)
        fill = slotfill_page(item, tree, catalog.slide(4), client)
        assert fake.calls == []
        assert fill.table is not None
        assert fill.table.rows == [["园区数", "10", "12"], ["设备数", "8000", "8600"]]

    def test_tagline_visible_to_llm(self, catalog: TemplateCatalog, tree,
                                    fake_llm_factory):
        """N1：tagline 槽进 LLM digest 与覆盖义务——必填或必删。"""
        from app.pipeline.slotfill import _slot_digest

        s4 = catalog.slide(4)
        t_ids = _s4_taglines(catalog)
        assert t_ids, "合成模板 S4 应含 tagline 槽"
        digest_ids = {d["slot_id"] for d in _slot_digest(s4)}
        assert set(t_ids) <= digest_ids
        # 只填正文槽、tagline 未覆盖 → V_SLOT_COVER
        body = next(s.slot_id for s in s4.slots if s.role == "body")
        f = PageFill(page_no=1, slide_type=SlideType.TEXT_POINTS,
                     template_slide=4, title="t",
                     slot_values={body: "内容"}, src_refs=[])
        issues = validate_page_fill(f, s4)
        assert any(i.rule is IssueCode.V_SLOT_COVER and t_ids[0] in i.detail
                   for i in issues)
        # 全覆盖（tagline 全删）→ 干净
        f2 = PageFill(page_no=1, slide_type=SlideType.TEXT_POINTS,
                      template_slide=4, title="t",
                      slot_values={body: "内容"}, slot_deletes=t_ids, src_refs=[])
        assert validate_page_fill(f2, s4) == []


class TestVision:
    """N3：模板页截图随提示走视觉模型；缺图退纯文本。"""

    def test_thumb_attaches_image_and_vision_model(self, catalog, tree,
                                                   fake_llm_factory, tmp_path):
        ids = _s1_ids(catalog)
        out = json.dumps({
            "title": "智慧园区汇报",
            "slots": {ids["title"]: "智慧园区汇报", ids["sub"]: "平台建设全景"},
            "deletes": [], "note": None,
        }, ensure_ascii=False)
        client, fake = fake_llm_factory([out])
        thumbs = tmp_path / "tpl.pages"
        thumbs.mkdir()
        (thumbs / "page_01.jpg").write_bytes(b"\xff\xd8fake-jpeg")
        item = _item(1, SlideType.COVER, "智慧园区汇报")
        fill = slotfill_page(item, tree, catalog.slide(1), client,
                             thumbs_dir=thumbs)
        assert len(fake.calls) == 1
        assert fake.calls[0]["model"] == client.vision_model
        content = fake.calls[0]["messages"][1]["content"]
        assert isinstance(content, list)
        assert content[0]["type"] == "text"
        assert content[1]["type"] == "image_url"
        assert content[1]["image_url"]["url"].startswith("data:image/jpeg;base64,")

    def test_missing_thumb_falls_back_to_text(self, catalog, tree,
                                              fake_llm_factory, tmp_path):
        ids = _s1_ids(catalog)
        out = json.dumps({
            "title": "智慧园区汇报",
            "slots": {ids["title"]: "智慧园区汇报", ids["sub"]: "平台建设全景"},
            "deletes": [], "note": None,
        }, ensure_ascii=False)
        client, fake = fake_llm_factory([out])
        item = _item(1, SlideType.COVER, "智慧园区汇报")
        fill = slotfill_page(item, tree, catalog.slide(1), client,
                             thumbs_dir=tmp_path / "nope.pages")
        assert len(fake.calls) == 1
        assert isinstance(fake.calls[0]["messages"][1]["content"], str)
        assert fill.slot_values[ids["title"]] == "智慧园区汇报"


class TestLLMPath:
    def test_retry_bad_then_good(self, catalog: TemplateCatalog, tree, fake_llm_factory):
        ids = _s3_ids(catalog)
        client, fake = fake_llm_factory([_bad_out(ids), _good_out(ids)])
        item = _item(4, SlideType.TEXT_POINTS, "覆盖规模", refs=["blk-0001"])
        fill = slotfill_page(item, tree, catalog.slide(3), client)
        assert len(fake.calls) == 2
        # 第二次调用带上了拒收原因
        second_user = fake.calls[1]["messages"][1]["content"]
        assert "V-SLOT-COVER" in second_user
        # 修正后的填充全量生效
        assert fill.slot_values[ids["title"]] == "覆盖规模与设备接入"
        assert fill.slot_values[ids["grp"][0]] == "园区覆盖"
        assert fill.note == "本页展开园区覆盖与设备接入规模。"

    def test_good_first_shot(self, catalog: TemplateCatalog, tree, fake_llm_factory):
        ids = _s3_ids(catalog)
        client, fake = fake_llm_factory([_good_out(ids)])
        item = _item(4, SlideType.TEXT_POINTS, "覆盖规模", refs=["blk-0001"])
        fill = slotfill_page(item, tree, catalog.slide(3), client)
        assert len(fake.calls) == 1
        assert validate_page_fill(fill, catalog.slide(3)) == []

    def test_degrade_after_two_failures(self, catalog: TemplateCatalog, tree,
                                        fake_llm_factory):
        ids = _s3_ids(catalog)
        client, fake = fake_llm_factory([_bad_out(ids), _bad_out(ids)])
        item = _item(4, SlideType.TEXT_POINTS, "覆盖规模与设备接入要点", refs=["blk-0001"])
        fill = slotfill_page(item, tree, catalog.slide(3), client)
        assert len(fake.calls) == 2
        # 降级：title-only + 源文 verbatim 备注
        assert list(fill.slot_values) == [ids["title"]]
        assert fill.slot_deletes == ids["grp"]
        src = tree.resolve(["blk-0001"])
        assert fill.note == src  # 源文前 300 当量（此处全量）
        assert fill.slot_values[ids["title"]] == "覆盖规模与设备接入要点"

    def test_retry_untraced_number_then_good(self, catalog: TemplateCatalog, tree,
                                             fake_llm_factory):
        # 铁律复检：槽文本数字须源文可溯源，1200 不在源文数字集 → 定向重试
        ids = _s3_ids(catalog)
        client, fake = fake_llm_factory([_num_out(ids), _good_out(ids)])
        item = _item(4, SlideType.TEXT_POINTS, "覆盖规模", refs=["blk-0001"])
        fill = slotfill_page(item, tree, catalog.slide(3), client)
        assert len(fake.calls) == 2
        assert "E-NUM-UNTRACED" in fake.calls[1]["messages"][1]["content"]
        # 第二轮修正后好填充全量生效
        assert fill.slot_values[ids["grp"][1]] == "十二个园区全面接入"

    def test_degrade_after_plagiarism(self, catalog: TemplateCatalog, tree,
                                      fake_llm_factory):
        # 铁律复检：槽文本与源文 ≥10-gram 雷同，两轮均未改写 → 降级
        ids = _s3_ids(catalog)
        client, fake = fake_llm_factory([_plag_out(ids), _plag_out(ids)])
        item = _item(4, SlideType.TEXT_POINTS, "覆盖规模与设备接入要点", refs=["blk-0001"])
        fill = slotfill_page(item, tree, catalog.slide(3), client)
        assert len(fake.calls) == 2
        assert "E-PLAGIARISM" in fake.calls[1]["messages"][1]["content"]
        assert list(fill.slot_values) == [ids["title"]]
        assert fill.slot_deletes == ids["grp"]
        assert fill.note == tree.resolve(["blk-0001"])

    def test_no_fillable_slot_raises(self, catalog: TemplateCatalog, tree,
                                     fake_llm_factory):
        empty = SlideInfo(slide_no=9, layout="", signature="", slots=[],
                          fillable=0, fillable_cap=0.0, role_hint="content")
        client, _ = fake_llm_factory([])
        item = _item(1, SlideType.TEXT_POINTS, "t", refs=["blk-0001"])
        with pytest.raises(FillError, match="完全无可用槽"):
            slotfill_page(item, tree, empty, client)


class TestValidate:
    def _fill(self, catalog: TemplateCatalog, **kw) -> PageFill:
        s3 = catalog.slide(3)
        ids = _s3_ids(catalog)
        base = dict(page_no=1, slide_type=SlideType.TEXT_POINTS,
                    template_slide=3, title="t",
                    slot_values={ids["title"]: "标题", **{g: "内容" for g in ids["grp"]}},
                    slot_deletes=[], src_refs=[])
        base.update(kw)
        return PageFill(**base)

    def test_ok(self, catalog: TemplateCatalog):
        assert validate_page_fill(self._fill(catalog), catalog.slide(3)) == []

    def test_unknown_slot(self, catalog: TemplateCatalog):
        f = self._fill(catalog, slot_values={**self._fill(catalog).slot_values, "99": "x"})
        rules = [i.rule for i in validate_page_fill(f, catalog.slide(3))]
        assert IssueCode.V_SLOT_UNKNOWN in rules

    def test_capacity(self, catalog: TemplateCatalog):
        ids = _s3_ids(catalog)
        long = "长" * 60
        f = self._fill(catalog, slot_values={
            **self._fill(catalog).slot_values, ids["grp"][0]: long})
        rules = [i.rule for i in validate_page_fill(f, catalog.slide(3))]
        assert IssueCode.V_SLOT_CAP in rules

    def test_coverage_and_title(self, catalog: TemplateCatalog):
        ids = _s3_ids(catalog)
        # 只填组槽、不填 title、漏一个组槽
        f = self._fill(catalog, slot_values={ids["grp"][0]: "a"},
                       slot_deletes=[ids["grp"][1]])
        rules = [i.rule for i in validate_page_fill(f, catalog.slide(3))]
        assert IssueCode.V_SLOT_COVER in rules
        assert IssueCode.V_SLOT_TITLE in rules

    def test_design_untouchable(self, catalog: TemplateCatalog):
        s4 = catalog.slide(4)
        page_no_sid = next(s.slot_id for s in s4.slots if s.role == "page_no")
        f = PageFill(page_no=1, slide_type=SlideType.TEXT_POINTS, template_slide=4,
                     title="t", slot_values={page_no_sid: "3"}, src_refs=[])
        rules = [i.rule for i in validate_page_fill(f, s4)]
        assert IssueCode.V_SLOT_UNKNOWN in rules

    def test_note_length(self, catalog: TemplateCatalog):
        f = self._fill(catalog, note="备" * 400)
        issues = validate_page_fill(f, catalog.slide(3))
        note_issues = [i for i in issues if i.rule is IssueCode.V_NOTE_LEN]
        assert note_issues and "300" in note_issues[0].detail


class TestTransportResilience:
    """N5：传输层耗尽（端点空响应）不杀 job——视觉失败退纯文本重试，再败走兜底。"""

    def test_vision_transport_failure_retries_text(self, catalog, tree,
                                                   fake_llm_factory, tmp_path,
                                                   monkeypatch):
        import time

        monkeypatch.setattr(time, "sleep", lambda s: None)
        ids = _s3_ids(catalog)
        client, fake = fake_llm_factory([""] * 9 + [_good_out(ids)])
        thumbs = tmp_path / "tpl.pages"
        thumbs.mkdir()
        (thumbs / "page_03.jpg").write_bytes(b"\xff\xd8fake-jpeg")
        item = _item(4, SlideType.TEXT_POINTS, "覆盖规模", refs=["blk-0001"])
        fill = slotfill_page(item, tree, catalog.slide(3), client,
                             thumbs_dir=thumbs)
        # 第一轮视觉 3×3 连空 → LLMError；第二轮纯文本默认模型提案成功
        assert len(fake.calls) == 10
        assert fake.calls[0]["model"] == client.vision_model
        assert isinstance(fake.calls[0]["messages"][1]["content"], list)
        assert fake.calls[9]["model"] == client.model
        assert isinstance(fake.calls[9]["messages"][1]["content"], str)
        assert fill.slot_values[ids["title"]] == "覆盖规模与设备接入"

    def test_content_degrades_after_transport_exhausted(self, catalog, tree,
                                                        fake_llm_factory,
                                                        monkeypatch):
        import time

        monkeypatch.setattr(time, "sleep", lambda s: None)
        ids = _s3_ids(catalog)
        client, fake = fake_llm_factory([""] * 18)
        item = _item(4, SlideType.TEXT_POINTS, "覆盖规模与设备接入要点",
                     refs=["blk-0001"])
        fill = slotfill_page(item, tree, catalog.slide(3), client)
        # 两轮（各 3×3 次重试）全空 → 内容页降级 title-only，不抛
        assert len(fake.calls) == 18
        assert list(fill.slot_values) == [ids["title"]]
        assert fill.note == tree.resolve(["blk-0001"])

    def test_structural_transport_falls_back_deterministic(self, catalog, tree,
                                                           fake_llm_factory,
                                                           monkeypatch):
        import time

        monkeypatch.setattr(time, "sleep", lambda s: None)
        ids = _s1_ids(catalog)
        client, fake = fake_llm_factory([""] * 18)
        item = _item(1, SlideType.COVER, "智慧园区汇报", brief="副题文本")
        fill = slotfill_page(item, tree, catalog.slide(1), client)
        assert len(fake.calls) == 18
        assert fill.slot_values[ids["title"]] == "智慧园区汇报"
        assert "副题文本" in fill.slot_values.values()
