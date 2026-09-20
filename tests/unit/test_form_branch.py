from __future__ import annotations

from pathlib import Path

import pytest

from app.llm.prompts import PromptManager
from app.pipeline.doc_runner import run_doc_pipeline
from app.pipeline.form_branch import (
    _apply_overlays,
    _doc_order,
    _docx_ok,
    _skeleton_table_block,
    run_form_branch,
)
from app.pipeline.tblarch import FormBranchFallback
from app.pipeline.tblcontent import _CellsOut
from app.schema.docplan import DocPlan, DocPlanItem
from app.schema.doctree import (
    DocBlock,
    DocMeta,
    DocSection,
    DocTable,
    DocTree,
)
from app.schema.enums import Genre, JobStatus
from app.schema.tblskeleton import (
    ArchitectOut,
    TCell,
    TRow,
    TSkeleton,
    TVisualTable,
    VisualOut,
)

PROSE_ORIG = "本表用于登记车间人员基本信息，作为月度考勤与培训记录的存档依据。"
PROSE_RW = "本表用于登记产线人员的基础信息，作为月度考勤和培训记录的归档凭证。"
LONG_ORIG = "统筹产线日常管理与设备运维督导，覆盖 12 条产线"
LONG_RW = "负责产线管理及设备维护督导工作，涉及 ⟦1⟧ 条生产线"
TITLE = "人员登记表"


# ---- 夹具 ----

class _FakeJob:
    def __init__(self, root: Path, source_name: str = "f.pdf",
                 confirmed: bool = True):
        self.dir = root
        self.source_name = source_name
        self.confirmed = confirmed
        self.error: str | None = None
        self.statuses: list[tuple] = []

    def set_status(self, status, detail: str = "") -> None:
        self.statuses.append((status, detail))

    def save_state(self) -> None:
        pass


class _ReplyClient:
    """按调用序弹出的假 LLM：架构师 → 内容专家 → 视觉总监。"""

    def __init__(self, replies):
        self._replies = list(replies)
        self.calls = 0

    def structured(self, schema, messages, stage=None, **kw):
        self.calls += 1
        if not self._replies:
            raise AssertionError("FakeClient 收到了多余的调用")
        return self._replies.pop(0)


class _NoLLM:
    def structured(self, *a, **kw):
        raise AssertionError("不应调用 LLM")


def _tree() -> DocTree:
    tbl = DocTable(
        table_id="tbl-001", section_id="sec-0000", n_rows=3, n_cols=3,
        header=["维度", "字段", "内容"],
        rows=[["基本情况", "姓名", LONG_ORIG],
              ["基本情况", "电话", "13800001111"]],
        caption=TITLE)
    blocks = [
        DocBlock(block_id="b-001", kind="para", text=PROSE_ORIG),
        DocBlock(block_id="b-002", kind="table", table_id="tbl-001"),
    ]
    sec = DocSection(section_id="sec-0000", level=0, title=TITLE,
                     blocks=blocks)
    full = (f"{TITLE} 维度 字段 内容 基本情况 姓名 {LONG_ORIG} "
            f"电话 13800001111 {PROSE_ORIG}")
    return DocTree(meta=DocMeta(source_format="pdf", source_name="f.pdf",
                                n_chars=len(full)),
                   sections=[sec], tables=[tbl], full_text=full)


def _plan() -> DocPlan:
    return DocPlan(genre=Genre.FORM, title=TITLE, items=[
        DocPlanItem(seq=1, section_id="sec-0000", heading=TITLE,
                    table_ids=["tbl-001"])])


def _arch_skeleton() -> TSkeleton:
    """架构师回声（掩码态）：基本情况 竖向合并跨 2 行。"""
    return TSkeleton(
        table_title=TITLE, total_cols=3, src_tables=["tbl-001"], rows=[
            TRow(cells=[TCell(content="维度", style="header"),
                        TCell(content="字段", style="header"),
                        TCell(content="内容", style="header")]),
            TRow(cells=[TCell(content="基本情况", rowspan=2, style="label"),
                        TCell(content="姓名", style="label"),
                        TCell(content="统筹产线日常管理与设备运维督导，覆盖⟦1⟧条产线",
                              style="input")]),
            TRow(cells=[TCell(content="电话", style="label"),
                        TCell(content="⟦2⟧", style="input")]),
        ])


def _full_client() -> _ReplyClient:
    return _ReplyClient([
        ArchitectOut(tables=[_arch_skeleton()], notes=""),
        _CellsOut(cells={"t0 1,2": LONG_RW, "p0": PROSE_RW}),
        VisualOut(tables=[TVisualTable(table_index=0,
                                       col_widths=[14, 14, 72],
                                       row_heights=[28, 40, 24])]),
    ])


def _write_ok_docx(path: Path, title: str) -> None:
    from docx import Document
    from docx.shared import Cm

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.18)
    d.add_paragraph(title)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "基本情况"
    t.cell(0, 1).text = "内容"
    t.cell(1, 0).text = "姓名"
    t.cell(1, 1).text = "张三"
    t.cell(0, 0).merge(t.cell(1, 0))
    d.save(str(path))


# ---- run_form_branch：全流程（无 COM → python-docx 兜底渲染）----

def test_run_form_branch_fallback_render(tmp_path):
    (tmp_path / "upload").mkdir()
    job = _FakeJob(tmp_path)
    client = _full_client()
    run_form_branch(job, client, PromptManager(), _plan(), _tree(),
                    com_export=False)

    assert [(s, d) for s, d in job.statuses] == [
        (JobStatus.GENERATING, "表格架构重建"),
        (JobStatus.GENERATING, "表格内容仿写"),
        (JobStatus.RENDERED, "render html→docx"),
        (JobStatus.DONE, "1 表 · 1 页"),
    ]
    assert client.calls == 3
    art = tmp_path / "artifacts"
    for name in ("tblarch.json", "tblcontent.json", "tblvisual.json",
                 "qa_report.txt", "output.docx"):
        assert (art / name).exists(), name

    from docx import Document

    d = Document(str(art / "output.docx"))
    assert len(d.tables) == 1
    xml = d.element.xml
    assert "gridSpan" in xml or "vMerge" in xml  # rowspan=2 → vMerge
    full = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
        c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "⟦" not in full
    assert "负责产线管理及设备维护督导工作，涉及 12 条生产线" in full
    assert "13800001111" in full          # 纯数字照搬
    assert PROSE_RW in full               # 散文改写
    assert TITLE in full


def test_run_form_branch_artifact_resume_zero_llm(tmp_path):
    (tmp_path / "upload").mkdir()
    job = _FakeJob(tmp_path)
    run_form_branch(job, _full_client(), PromptManager(), _plan(), _tree(),
                    com_export=False)
    job2 = _FakeJob(tmp_path)
    run_form_branch(job2, _NoLLM(), PromptManager(), _plan(), _tree(),
                    com_export=False)
    assert job2.statuses[-1][0] is JobStatus.DONE
    assert (tmp_path / "artifacts" / "output.docx").exists()


# ---- run_form_branch：COM 主路径（服务打桩）----

def test_run_form_branch_com_path(tmp_path, monkeypatch):
    from app.services import com_export as ce

    seen: dict[str, Path] = {}

    def fake_convert(html_path, out_docx):
        seen["html"] = Path(html_path)
        _write_ok_docx(out_docx, TITLE)
        return out_docx

    def fake_pdf(docx, out_pdf):
        import pymupdf

        doc = pymupdf.open()
        doc.new_page()
        doc.save(str(out_pdf))
        doc.close()
        return out_pdf

    monkeypatch.setattr(ce, "convert_html_to_docx", fake_convert)
    monkeypatch.setattr(ce, "export_docx_pdf", fake_pdf)

    (tmp_path / "upload").mkdir()
    job = _FakeJob(tmp_path)
    run_form_branch(job, _full_client(), PromptManager(), _plan(), _tree(),
                    com_export=True)

    art = tmp_path / "artifacts"
    # HTML 产出：utf-8-sig BOM（Word 嗅探）
    html_raw = (art / "output.html").read_bytes()
    assert html_raw.startswith(b"\xef\xbb\xbf")
    assert seen["html"] == art / "output.html"
    assert "@page" in html_raw.decode("utf-8-sig")
    assert (art / "output.pdf").exists()
    assert len(list((tmp_path / "pages").glob("*.png"))) == 1
    assert job.statuses[-1] == (JobStatus.DONE, "1 表 · 1 页")


def test_run_form_branch_sanity_fail_rerender(tmp_path, monkeypatch):
    """COM 产物 sanity 败（Letter 页面）→ python-docx 兜底重渲。"""
    from docx import Document
    from docx.shared import Cm
    from app.services import com_export as ce

    def bad_convert(html_path, out_docx):
        d = Document()
        d.add_paragraph(TITLE)
        d.add_table(rows=2, cols=2)
        d.save(str(out_docx))
        return out_docx

    def fake_pdf(docx, out_pdf):
        import pymupdf

        doc = pymupdf.open()
        doc.new_page()
        doc.save(str(out_pdf))
        doc.close()
        return out_pdf

    monkeypatch.setattr(ce, "convert_html_to_docx", bad_convert)
    monkeypatch.setattr(ce, "export_docx_pdf", fake_pdf)
    (tmp_path / "upload").mkdir()
    job = _FakeJob(tmp_path)
    run_form_branch(job, _full_client(), PromptManager(), _plan(), _tree(),
                    com_export=True)

    d = Document(str(tmp_path / "artifacts" / "output.docx"))
    sec = d.sections[0]
    assert abs(sec.page_width.cm - 21.0) < 0.2     # 兜底渲染恢复 A4
    assert abs(sec.left_margin.cm - 3.18) < 0.2
    assert d.tables and "vMerge" in d.element.xml
    assert job.statuses[-1][0] is JobStatus.DONE


# ---- _docx_ok 检查项 ----

def test_docx_ok_clean(tmp_path):
    p = tmp_path / "ok.docx"
    _write_ok_docx(p, TITLE)
    assert _docx_ok(p, TITLE, [_arch_skeleton()]) == []


def test_docx_ok_page_and_margins(tmp_path):
    from docx import Document
    from docx.shared import Cm

    p = tmp_path / "letter.docx"
    d = Document()
    d.add_paragraph(TITLE)
    d.add_table(rows=1, cols=1)
    d.save(str(p))
    assert any("A4" in e for e in _docx_ok(p, TITLE, []))

    p2 = tmp_path / "margin.docx"
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(2.0)
    d.add_paragraph(TITLE)
    d.add_table(rows=1, cols=1)
    d.save(str(p2))
    errors = _docx_ok(p2, TITLE, [])
    assert any("左边距" in e for e in errors)
    assert any("右边距" in e for e in errors)


def test_docx_ok_placeholder_and_title(tmp_path):
    from docx import Document

    p = tmp_path / "ph.docx"
    _write_ok_docx(p, TITLE)
    d = Document(str(p))
    d.add_paragraph("残留 ⟦1⟧")
    d.save(str(p))
    errors = _docx_ok(p, TITLE, [])
    assert any("⟦" in e for e in errors)

    d = Document(str(p))
    errors = _docx_ok(p, "另一个标题", [])
    assert any("标题缺失" in e for e in errors)


def test_docx_ok_no_table(tmp_path):
    from docx import Document

    p = tmp_path / "no_tbl.docx"
    d = Document()
    d.add_paragraph(TITLE)
    d.save(str(p))
    assert _docx_ok(p, TITLE, []) == ["输出无表格"]


# ---- 辅助函数 ----

def test_apply_overlays_does_not_mutate():
    sks = [_arch_skeleton()]
    out = _apply_overlays(sks, {0: {"1,2": "替换后的长文本内容"}})
    assert "替换后的长文本内容" in out[0].rows[1].cells[2].content
    assert "⟦1⟧" in sks[0].rows[1].cells[2].content  # 原骨架不动


def test_skeleton_table_block_expansion():
    v = TVisualTable(table_index=0, col_widths=[14, 14, 72],
                     row_heights=[28, 40, 24])
    b = _skeleton_table_block(0, _arch_skeleton(), v)
    assert b.header == ["维度", "字段", "内容"]
    assert b.rows == [["基本情况", "姓名", "统筹产线日常管理与设备运维督导，覆盖⟦1⟧条产线"],
                      ["", "电话", "⟦2⟧"]]
    assert b.merges == [[1, 0, 2, 1]]
    assert b.col_widths == [0.14, 0.14, 0.72]
    assert b.row_heights == [28, 40, 24]


def test_doc_order_absorption_and_tail_append():
    sk = TSkeleton(table_title="", total_cols=1, src_tables=["tbl-001", "tbl-002"],
                   rows=[TRow(cells=[TCell(content="合并格")])])
    sk2 = TSkeleton(table_title="", total_cols=1, src_tables=["tbl-004"],
                    rows=[TRow(cells=[TCell(content="尾部格")])])
    blocks = [
        DocBlock(block_id="b1", kind="para", text=PROSE_ORIG),
        DocBlock(block_id="b2", kind="table", table_id="tbl-001"),
        DocBlock(block_id="b3", kind="table", table_id="tbl-002"),  # 被吸收
        DocBlock(block_id="b4", kind="table", table_id="tbl-003"),  # 无主碎片
        DocBlock(block_id="b5", kind="para", text="短句"),          # <12 照搬
    ]
    sec = DocSection(section_id="s", level=0, title="", blocks=blocks)
    tree = DocTree(meta=DocMeta(source_format="pdf", source_name="f",
                                n_chars=9),
                   sections=[sec], tables=[], full_text="x")
    items, warnings = _doc_order(tree, [sk, sk2], {"p0": PROSE_RW})
    assert items == [("para", PROSE_RW), ("table", 0), ("para", "短句"),
                     ("table", 1)]  # 未放置骨架尾部补挂
    assert len(warnings) == 1 and "tbl-003" in warnings[0]


# ---- doc_runner 分发 ----

def _setup_dispatch_job(root: Path, source_format: str, genre: Genre,
                        with_tables: bool = True) -> _FakeJob:
    art = root / "artifacts"
    art.mkdir(parents=True)
    (root / "upload").mkdir()
    (root / "upload" / "f.pdf").write_bytes(b"%PDF-1.4")
    tree = _tree()
    tree.meta.source_format = source_format  # type: ignore[assignment]
    if not with_tables:
        tree.tables = []
    (art / "doctree.json").write_text(tree.model_dump_json(indent=2),
                                      encoding="utf-8")
    plan = _plan()
    plan.genre = genre
    (art / "docplan.json").write_text(plan.model_dump_json(indent=2),
                                      encoding="utf-8")
    return _FakeJob(root)


class _Sentinel(Exception):
    pass


def test_dispatch_form_pdf_enters_branch(tmp_path, monkeypatch):
    from app.pipeline import doc_runner, form_branch

    job = _setup_dispatch_job(tmp_path, "pdf", Genre.FORM)
    calls: list = []

    def fake_branch(*a, **kw):
        calls.append((a, kw))
        return None

    def no_fill(*a, **kw):
        raise _Sentinel("不应进入旧链路")

    monkeypatch.setattr(form_branch, "run_form_branch", fake_branch)
    monkeypatch.setattr(doc_runner, "fill_all_sections", no_fill)
    run_doc_pipeline(job, client=object(), com_export=False)
    assert len(calls) == 1
    assert calls[0][1]["com_export"] is False


def test_dispatch_form_docx_skips_branch(tmp_path, monkeypatch):
    from app.pipeline import doc_runner

    job = _setup_dispatch_job(tmp_path, "docx", Genre.FORM)
    monkeypatch.setattr(doc_runner, "fill_all_sections",
                        lambda *a, **kw: (_ for _ in ()).throw(_Sentinel()))
    with pytest.raises(_Sentinel):
        run_doc_pipeline(job, client=object(), com_export=False)


def test_dispatch_report_pdf_skips_branch(tmp_path, monkeypatch):
    from app.pipeline import doc_runner

    job = _setup_dispatch_job(tmp_path, "pdf", Genre.REPORT)
    monkeypatch.setattr(doc_runner, "fill_all_sections",
                        lambda *a, **kw: (_ for _ in ()).throw(_Sentinel()))
    with pytest.raises(_Sentinel):
        run_doc_pipeline(job, client=object(), com_export=False)


def test_dispatch_form_pdf_without_tables_skips_branch(tmp_path, monkeypatch):
    from app.pipeline import doc_runner

    job = _setup_dispatch_job(tmp_path, "pdf", Genre.FORM, with_tables=False)
    monkeypatch.setattr(doc_runner, "fill_all_sections",
                        lambda *a, **kw: (_ for _ in ()).throw(_Sentinel()))
    with pytest.raises(_Sentinel):
        run_doc_pipeline(job, client=object(), com_export=False)


def test_dispatch_fallback_continues_legacy(tmp_path, monkeypatch):
    from app.pipeline import doc_runner, form_branch

    job = _setup_dispatch_job(tmp_path, "pdf", Genre.FORM)

    def boom(*a, **kw):
        raise FormBranchFallback("架构师两轮失败")

    monkeypatch.setattr(form_branch, "run_form_branch", boom)
    monkeypatch.setattr(doc_runner, "fill_all_sections",
                        lambda *a, **kw: (_ for _ in ()).throw(_Sentinel()))
    with pytest.raises(_Sentinel):  # 旧链路被继续执行
        run_doc_pipeline(job, client=object(), com_export=False)
