from __future__ import annotations

from pathlib import Path

import pytest

from app.render.template_build import build_deck
from app.render.tpl_catalog import load_catalog, template_pptx_path
from app.schema.enums import IssueCode, SlideType
from app.schema.slotfill import (BuildMeta, DeckBuild, PageFill, TableIR,
                                 TemplateCatalog)
from app.services.com_export import export_pngs, template_checklist

pytestmark = pytest.mark.com

STEM = "办公技巧商务培训PPT模板"  # 33 页中文真模板


@pytest.fixture(scope="module")
def catalog() -> TemplateCatalog:
    return load_catalog(STEM)


def _page(no: int, cat_slide, slide_type: SlideType,
          title_only: bool = False, table: TableIR | None = None) -> PageFill:
    """非 design/page_no 槽：title 必填、可填槽填短文本、小容量槽删除。"""
    values: dict[str, str] = {}
    deletes: list[str] = []
    for s in cat_slide.slots:
        if s.role in ("design", "page_no"):
            continue
        if s.role == "title":
            values[s.slot_id] = "季度成果汇报"
        elif title_only:
            continue  # 不填不删 → 遗留检测样本
        elif s.fillable:
            values[s.slot_id] = "要点说明"
        else:
            deletes.append(s.slot_id)
    return PageFill(page_no=no, slide_type=slide_type,
                    template_slide=cat_slide.slide_no, title="页",
                    slot_values=values, slot_deletes=deletes, table=table)


@pytest.fixture(scope="module")
def deck(catalog: TemplateCatalog) -> DeckBuild:
    # cover(S1) + 表格宿主(S2，纯构图页) + 内容页(S3) + closing(S33)
    return DeckBuild(meta=BuildMeta(title="冒烟", template=STEM), pages=[
        _page(1, catalog.slide(1), SlideType.COVER),
        _page(2, catalog.slide(2), SlideType.TABLE, table=TableIR(
            table_id="tbl-001", header=["指标", "目标", "实际"],
            rows=[["达成率", "100%", "120%"]])),
        _page(3, catalog.slide(3), SlideType.TEXT_POINTS),
        _page(4, catalog.slide(33), SlideType.CLOSING),
    ])


@pytest.fixture(scope="module")
def rendered(tmp_path_factory, catalog, deck) -> Path:
    out = tmp_path_factory.mktemp("com") / "deck.pptx"
    build_deck(template_pptx_path(STEM), catalog, deck, out)
    return out


def test_export_pngs_count(rendered, tmp_path):
    pngs = export_pngs(rendered, tmp_path / "pngs")
    assert len(pngs) == 4
    for p in pngs:
        assert p.exists() and p.stat().st_size > 1000


def test_checklist_clean(rendered, deck, catalog):
    issues = template_checklist(deck, catalog, rendered)
    assert not [i for i in issues if i.code is IssueCode.E_RENDER_MISMATCH], \
        [i.detail for i in issues]


def test_checklist_detects_page_mismatch(rendered, deck, catalog):
    tampered = DeckBuild.model_validate(deck.model_dump())
    tampered.pages = tampered.pages[:-1]  # 模拟 IR 与 pptx 页数不一致
    issues = template_checklist(tampered, catalog, rendered)
    assert any(i.code is IssueCode.E_RENDER_MISMATCH and "页数" in i.detail for i in issues)


def test_checklist_detects_leftover(catalog, tmp_path):
    """真模板遗留检测：S3 只填 title，body 原文残留 → E_LEFTOVER。"""
    page = _page(1, catalog.slide(3), SlideType.TEXT_POINTS, title_only=True)
    deck = DeckBuild(meta=BuildMeta(title="t", template=STEM), pages=[page])
    out = tmp_path / "leftover.pptx"
    build_deck(template_pptx_path(STEM), catalog, deck, out)
    issues = template_checklist(deck, catalog, out)
    assert any(i.code is IssueCode.E_LEFTOVER for i in issues), \
        [i.detail for i in issues]



def test_convert_doc_to_docx_roundtrip(basic_docx, tmp_path):
    """真 Word：docx → .doc（wdFormatDocument=0）→ convert_doc_to_docx 还原可解析。"""
    from app.parse_dispatch import parse_source
    from app.services.com_export import convert_doc_to_docx, get_word_com_service

    def _to_doc(svc):
        app = svc._ensure_app()
        doc = app.Documents.Open(str(Path(basic_docx).resolve()), False, True, False)
        try:
            out = tmp_path / "legacy.doc"
            doc.SaveAs2(str(out.resolve()), FileFormat=0)  # wdFormatDocument
            return out
        finally:
            doc.Close(False)

    legacy = get_word_com_service().run(_to_doc)
    converted = tmp_path / "legacy.converted.docx"
    convert_doc_to_docx(legacy, converted)

    tree = parse_source(legacy)  # converted 已在 → 幂等跳过，直接解析
    assert "智慧园区" in tree.full_text
    assert any(".doc" in w for w in tree.meta.parse_warnings)


def test_word_accepts_copied_table_xml(tmp_path):
    """C3 关键风险：deepcopy 跨包的表格 XML Word 必须能正常打开并导 PDF。"""
    from docx import Document as DocxDocument
    from docx.shared import Cm

    from app.render.docx_render import render_docir_to_docx
    from app.schema.docir import DocIR, DocIRMeta, DocTitleBlock, TableBlock
    from app.schema.enums import Genre
    from app.services.com_export import export_docx_pdf

    src = tmp_path / "src.docx"
    d = DocxDocument()
    d.add_paragraph("表1：人员")
    t = d.add_table(rows=3, cols=3)
    for i, w in enumerate((Cm(6), Cm(3), Cm(2))):
        t.columns[i].width = w
    t.cell(0, 0).text = "基本信息"
    t.cell(0, 2).text = "备注"
    t.cell(1, 0).text = "张三"
    t.cell(1, 1).text = "35"
    t.cell(1, 2).text = "该同志负责产线管理工作。"
    t.cell(2, 0).text = "李四"
    t.cell(2, 1).text = "42"
    t.cell(2, 2).text = "在岗。"
    t.cell(0, 0).merge(t.cell(0, 1))
    d.save(str(src))

    ir = DocIR(meta=DocIRMeta(title="人员信息表", genre=Genre.FORM), blocks=[
        DocTitleBlock(text="人员信息表"),
        TableBlock(table_id="tbl-001", src_index=0,
                   header=["基本信息", "基本信息", "备注"],
                   rows=[["张三", "35", "该同志负责产线管理工作。"],
                         ["李四", "42", "在岗。"]]),
    ])
    docx = render_docir_to_docx(ir, tmp_path / "out.docx", source=src)
    pdf = export_docx_pdf(docx, tmp_path / "out.pdf")
    assert pdf.exists() and pdf.stat().st_size > 2000

    import pymupdf

    with pymupdf.open(str(pdf)) as doc:
        assert len(doc) >= 1
        assert "张三" in doc[0].get_text()


def test_convert_html_to_docx(tmp_path):
    """T4 关键风险：Word HTML 导入 @page/colgroup/rowspan → docx 保真。"""
    from docx import Document as DocxDocument

    from app.render.form_html import FormBlock, build_form_html
    from app.schema.tblskeleton import TCell, TRow, TSkeleton, TVisualTable
    from app.services.com_export import convert_html_to_docx, export_docx_pdf

    sk = TSkeleton(
        table_title="人员登记表", total_cols=3, src_tables=[], rows=[
            TRow(cells=[TCell(content="维度", style="header"),
                        TCell(content="字段", style="header"),
                        TCell(content="内容", style="header")]),
            TRow(cells=[TCell(content="基本情况", rowspan=2, style="label"),
                        TCell(content="姓名", style="label"),
                        TCell(content="负责产线管理，覆盖 12 条产线")]),
            TRow(cells=[TCell(content="电话", style="label"),
                        TCell(content="13800001111")]),
        ])
    v = TVisualTable(table_index=0, col_widths=[14, 14, 72],
                     row_heights=[28, 40, 24])
    html = build_form_html("人员登记表", [FormBlock(kind="table",
                                                   table_index=0)], [sk], [v])
    html_path = tmp_path / "form.html"
    html_path.write_text(html, encoding="utf-8-sig")
    docx = convert_html_to_docx(html_path, tmp_path / "form.docx")

    d = DocxDocument(str(docx))
    assert len(d.tables) == 1
    xml = d.element.xml
    assert "gridSpan" in xml or "vMerge" in xml
    sec = d.sections[0]
    assert abs(sec.page_width.cm - 21.0) < 0.2
    assert abs(sec.page_height.cm - 29.7) < 0.2
    assert abs(sec.top_margin.cm - 2.54) < 0.2
    assert abs(sec.left_margin.cm - 3.18) < 0.2
    full = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
        c.text for t in d.tables for r in t.rows for c in r.cells)
    assert "人员登记表" in full and "基本情况" in full
    assert "13800001111" in full and "12" in full
    assert "⟦" not in full
    # F3：垂直居中后处理——Word HTML 导入常丢 CSS vertical-align，
    # COM 保存前逐格补 VerticalAlignment=1
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                assert c.vertical_alignment == WD_CELL_VERTICAL_ALIGNMENT.CENTER

    pdf = export_docx_pdf(docx, tmp_path / "form.pdf")
    import pymupdf

    with pymupdf.open(str(pdf)) as doc:
        text = "".join(page.get_text() for page in doc)
        assert "基本情况" in text
        assert "13800001111" in text and "12" in text
