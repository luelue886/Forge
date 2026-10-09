"""程序化合成 5 页测试模板（单测零 COM / 零真 LLM / 零真模板）。

页设计覆盖全部提取与构建路径：
  S1 cover——版式占位符（title/subtitle ph 继承字号链）
  S2 toc——标题 + 正文占位符 5 条目（多段单槽）
  S3 content——独立大字号标题 + 图片 + GROUP（非恒等 chOff/chExt，2×(标签+描述)+组内图片）+ notesSlide + 页级 tags
  S4 content——3 个独立文本框，含纯数字 "02 / 14"（page_no）与超小样本（design）
  S5 closing——占位符标题 + 文本框
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
IMG_REL = f"{R}/image"

# GROUP 几何：off=(1",2") ext=(6",3")，子空间 chOff=(100,200) chExt=(300,150)
# 子元素 child 坐标 (cx,cy)：页面绝对 = (1"+(cx-100)*6/300, 2"+(cy-200)*3/150)
GRP_OFF_X, GRP_OFF_Y = 914400, 1828800
GRP_EXT_CX, GRP_EXT_CY = 5486400, 2743200
GRP_CHOFF_X, GRP_CHOFF_Y = 100, 200
GRP_CHEXT_CX, GRP_CHEXT_CY = 300, 150

# 组内图片的关系 id（构建测试断言"组删除后 rels 剪除"用）
GROUP_PIC_RID = "rIdG1"

# 页级 tags 部件引用（复刻真模板 Aspose/OfficePlus 元数据）：构建测试断言
# "复制页共享 tag 部件非法（PowerPoint 拒开文件）"剥离路径用
TAG_RID = "rIdT1"
TAG_PART = "ppt/tags/tag1.xml"
TAG_CT = "application/vnd.openxmlformats-officedocument.presentationml.tags+xml"
CT_PKG = "http://schemas.openxmlformats.org/package/2006/content-types"


def _sp_xml(name: str, x: int, y: int, cx: int, cy: int, text: str,
            pt: int, bold: bool = False) -> str:
    b = ' b="1"' if bold else ""
    esc = (
        text.replace("&", "&amp;").replace("<", "&lt;")
        .replace(">", "&gt;").replace("\n", "&#10;")
    )
    return f"""<p:sp xmlns:p="{P}" xmlns:a="{A}">
  <p:nvSpPr><p:cNvPr id="0" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>
  <p:spPr>
    <a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
  <p:txBody>
    <a:bodyPr/><a:lstStyle/>
    <a:p><a:pPr algn="l"/><a:r><a:rPr lang="zh-CN" sz="{pt * 100}"{b}><a:latin typeface="微软雅黑"/><a:ea typeface="微软雅黑"/></a:rPr><a:t>{esc}</a:t></a:r></a:p>
  </p:txBody>
</p:sp>"""


def _pic_xml(rid: str, x: int, y: int, cx: int, cy: int) -> str:
    return f"""<p:pic xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}">
  <p:nvPicPr><p:cNvPr id="0" name="组内图片"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>
  <p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>
  <p:spPr>
    <a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
  </p:spPr>
</p:pic>"""


def _inject_group(path: Path, slide_xml_name: str) -> None:
    """向指定页 spTree 注入带非恒等子坐标空间的 GROUP（python-pptx 无建组 API）。

    组尾附一张图片（引用顶层图片同一 media），并为该页 rels 注册新 rId——
    覆盖"整组删除时连带图片 + rels 剪除"的构建路径。
    """
    label1 = _sp_xml("标签一", 100, 200, 140, 40, "团队协作", 16, bold=True)
    desc1 = _sp_xml("描述一", 100, 250, 140, 70, "跨部门协同推进项目落地", 12)
    label2 = _sp_xml("标签二", 260, 200, 140, 40, "数据驱动", 16, bold=True)
    desc2 = _sp_xml("描述二", 260, 250, 140, 70, "以指标复盘持续改进", 12)
    pic = _pic_xml(GROUP_PIC_RID, 330, 260, 50, 40)
    group = f"""<p:grpSp xmlns:p="{P}" xmlns:a="{A}" xmlns:r="{R}">
  <p:nvGrpSpPr><p:cNvPr id="0" name="内容组"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
  <p:grpSpPr>
    <a:xfrm>
      <a:off x="{GRP_OFF_X}" y="{GRP_OFF_Y}"/>
      <a:ext cx="{GRP_EXT_CX}" cy="{GRP_EXT_CY}"/>
      <a:chOff x="{GRP_CHOFF_X}" y="{GRP_CHOFF_Y}"/>
      <a:chExt cx="{GRP_CHEXT_CX}" cy="{GRP_CHEXT_CY}"/>
    </a:xfrm>
  </p:grpSpPr>
  {label1}{desc1}{label2}{desc2}{pic}
</p:grpSp>"""
    group_el = etree.fromstring(group.encode())

    rels_name = slide_xml_name.replace("slides/", "slides/_rels/") + ".rels"
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == slide_xml_name:
                root = etree.fromstring(data)
                sp_tree = root.find(f"{{{P}}}cSld/{{{P}}}spTree")
                sp_tree.append(group_el)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item == rels_name:
                root = etree.fromstring(data)
                img_target = next(
                    (rel.get("Target") for rel in root
                     if (rel.get("Type") or "").endswith("/image")), None)
                if img_target:
                    rel = etree.SubElement(root, f"{{{REL_PKG}}}Relationship")
                    rel.set("Id", GROUP_PIC_RID)
                    rel.set("Type", IMG_REL)
                    rel.set("Target", img_target)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    path.write_bytes(buf.getvalue())


def _inject_tags(path: Path, slide_xml_name: str) -> None:
    """向指定页注入页级 tags 部件引用（复刻真模板 Aspose/OfficePlus 元数据）。

    四件套：tag 部件本体 + [Content_Types] Override + 页 rels 关系 +
    cSld 下的 p:custDataLst——多页复制后若共享同一 tag 部件，
    PowerPoint 将拒绝打开文件（构建层必须剥离）。
    """
    tag_xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<p:tagLst xmlns:p="{P}"><p:tag name="forge-test" val="1"/></p:tagLst>'
    ).encode()
    cust = etree.fromstring(
        f'<p:custDataLst xmlns:p="{P}" xmlns:r="{R}">'
        f'<p:tags r:id="{TAG_RID}"/></p:custDataLst>'.encode())

    rels_name = slide_xml_name.replace("slides/", "slides/_rels/") + ".rels"
    ct_name = "[Content_Types].xml"
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == slide_xml_name:
                root = etree.fromstring(data)
                csld = root.find(f"{{{P}}}cSld")
                ext = csld.find(f"{{{P}}}extLst")
                if ext is not None:
                    ext.addprevious(cust)  # 规范顺序：spTree 后、extLst 前
                else:
                    csld.append(cust)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item == rels_name:
                root = etree.fromstring(data)
                rel = etree.SubElement(root, f"{{{REL_PKG}}}Relationship")
                rel.set("Id", TAG_RID)
                rel.set("Type", f"{R}/tags")
                rel.set("Target", "../tags/tag1.xml")
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            elif item == ct_name:
                root = etree.fromstring(data)
                ov = etree.SubElement(root, f"{{{CT_PKG}}}Override")
                ov.set("PartName", f"/{TAG_PART}")
                ov.set("ContentType", TAG_CT)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
        zout.writestr(TAG_PART, tag_xml)
    path.write_bytes(buf.getvalue())


def _inject_timing(path: Path, slide_xml_name: str) -> None:
    """向指定页注入 p:timing 动画树（spTgt spid 指向组内标签形状）。

    复刻真模板实况：删形状（槽删除/组塌缩）后 timing 的 spid 悬空，
    PowerPoint 拒开文件——构建层必须整树剥离，构建测试的回归防线。
    """
    timing = etree.fromstring(f"""<p:timing xmlns:p="{P}" xmlns:a="{A}">
  <p:tnLst>
    <p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
      <p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
        <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>
          <p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
            <p:par><p:cTn id="5" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
              <p:set><p:cBhvr>
                <p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
                <p:tgtEl><p:spTgt spid="0"/></p:tgtEl>
                <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
              </p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
            </p:childTnLst></p:cTn></p:par>
          </p:childTnLst></p:cTn></p:par>
        </p:childTnLst></p:cTn></p:par>
      </p:childTnLst></p:cTn></p:seq>
    </p:childTnLst></p:cTn></p:par>
  </p:tnLst>
</p:timing>""".encode())

    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item == slide_xml_name:
                root = etree.fromstring(data)
                root.append(timing)  # 规范顺序：cSld/clrMapOvr/transition 之后
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            zout.writestr(item, data)
    path.write_bytes(buf.getvalue())


def build_synthetic_template(path: Path) -> Path:
    prs = Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide
    s1.shapes.title.text = "年度培训主题"
    s1.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    s1.placeholders[1].text = "副标题示例文本"

    s2 = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    s2.shapes.title.text = "目录"
    tf = s2.placeholders[1].text_frame
    for i in range(5):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = f"第{i + 1}章节要点概览"

    s3 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    title = s3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
    title.text_frame.text = "内容页大标题"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    img = Image.new("RGB", (8, 8), (120, 160, 200))
    bio = io.BytesIO()
    img.save(bio, format="PNG")
    bio.seek(0)
    s3.shapes.add_picture(bio, Inches(11.5), Inches(0.2), Inches(0.6), Inches(0.6))
    s3.notes_slide.notes_text_frame.text = "模板原始备注"

    s4 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    t1 = s4.shapes.add_textbox(Inches(0.5), Inches(1), Inches(4), Inches(1.2))
    t1.text_frame.text = "独立正文文本框示例"
    t1.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    t2 = s4.shapes.add_textbox(Inches(6), Inches(1), Inches(1.2), Inches(0.4))
    t2.text_frame.text = "02 / 14"
    t2.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    t3 = s4.shapes.add_textbox(Inches(0.5), Inches(3), Inches(0.4), Inches(0.3))
    t3.text_frame.text = "·"
    t3.text_frame.paragraphs[0].runs[0].font.size = Pt(10)

    s5 = prs.slides.add_slide(prs.slide_layouts[0])
    s5.shapes.title.text = "谢谢观看"
    s5.placeholders[1].text = "汇报人落款"

    prs.save(path)
    # python-pptx 页序与文件名一致（默认模板顺序写入）
    _inject_group(path, "ppt/slides/slide3.xml")
    _inject_tags(path, "ppt/slides/slide3.xml")
    _inject_timing(path, "ppt/slides/slide3.xml")
    return path
