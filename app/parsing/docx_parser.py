from __future__ import annotations

import re
from pathlib import Path

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsing.base import (
    ParseError,
    clean_text,
    image_size_class,
    vml_style_size_cm,
)
from app.schema.doctree import (
    DocBlock,
    DocImage,
    DocMeta,
    DocSection,
    DocTable,
    DocTree,
)
from app.schema.textlen import text_weight

_HEADING = re.compile(r"^(?:Heading|标题)\s*(\d)$", re.IGNORECASE)
_CAPTION = re.compile(r"^表\s*\d")
_LIST_STYLE = re.compile(r"^(?:List|列表)", re.IGNORECASE)

_W_PICT = qn("w:pict")
_W_DRAWING = qn("w:drawing")
_W_TXBX_CONTENT = qn("w:txbxContent")
_WP_EXTENT = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
_DGM_RELIDS = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_V_SHAPE = "{urn:schemas-microsoft-com:vml}shape"
_EMU_PER_CM = 360000


def _heading_level(p: Paragraph) -> int | None:
    pPr = p._p.pPr
    if pPr is not None:
        ol = pPr.find(qn("w:outlineLvl"))
        if ol is not None and ol.get(qn("w:val")) is not None:
            v = int(ol.get(qn("w:val")))
            if 0 <= v <= 8:
                return v + 1
    m = _HEADING.match(p.style.name or "")
    return int(m.group(1)) if m else None


def _is_list_item(p: Paragraph) -> bool:
    # numPr 可能在段落直接格式，也可能在样式定义里（如 List Bullet）
    pPr = p._p.pPr
    if pPr is not None and pPr.find(qn("w:numPr")) is not None:
        return True
    style = p.style
    if _LIST_STYLE.match(style.name or ""):
        return True
    spPr = style.element.find(qn("w:pPr"))
    return spPr is not None and spPr.find(qn("w:numPr")) is not None


def _iter_blocks(doc):
    """按文档顺序产出 (body 序号, Paragraph / Table)。

    doc.paragraphs 与 doc.tables 会丢失交错关系；pos 是 body 子元素全序号
    （含非块子元素），渲染期 deepcopy 图片段落用同一序号定位。
    """
    for pos, child in enumerate(doc.element.body.iterchildren()):
        if child.tag == qn("w:p"):
            yield pos, Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield pos, Table(child, doc)


def _txbx_text(p_el) -> list[str]:
    """段落内文本框（w:txbxContent）的文字，按文档顺序；无文本框返回空。

    新式文本框是 mc:AlternateContent 双写（Choice=新格式，Fallback=旧格式
    存同一段文字），只取 Choice 侧防重复；旧式 v:textbox 无包裹天然单份。
    """
    out: list[str] = []
    for box in p_el.findall(f".//{_W_TXBX_CONTENT}"):
        if next(box.iterancestors(_MC_FALLBACK), None) is not None:
            continue
        for para in box.iter(qn("w:p")):
            parts: list[str] = []
            for t in para.iter(qn("w:t")):
                # 更深层嵌套文本框的文字由那个框自己负责，防重复
                if any(a is not box for a in t.iterancestors(_W_TXBX_CONTENT)):
                    continue
                parts.append(t.text or "")
            text = clean_text("".join(parts))
            if text:
                out.append(text)
    return out


def _image_extent(p_el) -> tuple[int | None, int | None]:
    ext = p_el.find(f".//{_WP_EXTENT}")
    if ext is None:
        return None, None
    try:
        return int(ext.get("cx")), int(ext.get("cy"))
    except (TypeError, ValueError):
        return None, None


def _image_size(p_el) -> str | None:
    """段内图形的物理尺寸分类；wp:extent 优先，VML 回退 v:shape style。"""
    cx, cy = _image_extent(p_el)
    if cx is not None and cy is not None:
        return image_size_class(cx / _EMU_PER_CM, cy / _EMU_PER_CM)
    shape = p_el.find(f".//{_V_SHAPE}")
    if shape is not None:
        return image_size_class(*vml_style_size_cm(shape.get("style") or ""))
    return None


def _image_of(p_el, image_id: str, section_id: str, pos: int,
              warnings: list[str]) -> DocImage | None:
    """空文本段落内的 w:drawing / w:pict → DocImage；SmartArt/文本框/小图跳过。"""
    if p_el.find(f".//{_DGM_RELIDS}") is not None:
        warnings.append(f"{image_id}：SmartArt 依赖多个图表部件，暂不支持复用，已跳过")
        return None
    if p_el.find(f".//{_W_TXBX_CONTENT}") is not None:
        return None  # 文本框图形：文字走 _txbx_text，空框不当图片复用
    drawings = p_el.findall(f".//{_W_DRAWING}")
    picts = p_el.findall(f".//{_W_PICT}")
    if not drawings and not picts:
        return None
    if len(drawings) + len(picts) > 1:
        warnings.append(f"{image_id}：段落含多个图形，仅复用首个")
    size = _image_size(p_el)
    if size == "portrait":
        warnings.append(f"{image_id}：疑似证件照，已跳过不搬运")
        return None
    if size == "icon":
        return None
    cx, cy = _image_extent(p_el)
    return DocImage(image_id=image_id, section_id=section_id,
                    body_index=pos, cx_emu=cx, cy_emu=cy)


def _grid_col_widths(tbl: Table) -> list[float] | None:
    """tblGrid/gridCol 的 w:w → 归一化列宽比例；缺失或退化返回 None。"""
    grid = tbl._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return None
    ws: list[int] = []
    for gc in grid.findall(qn("w:gridCol")):
        v = gc.get(qn("w:w"))
        if v is None or not v.strip().isdigit():
            return None
        ws.append(int(v))
    if len(ws) < 2 or sum(ws) <= 0:
        return None
    total = sum(ws)
    return [w / total for w in ws]


def _extract_table(tbl: Table, table_id: str, section_id: str, caption: str | None,
                   warnings: list[str], src_index: int) -> DocTable:
    grid = [[clean_text(c.text) for c in row.cells] for row in tbl.rows]
    if not grid or not grid[0]:
        warnings.append(f"{table_id}：空表格，已跳过内容")
        return DocTable(table_id=table_id, section_id=section_id, n_rows=0, n_cols=0, caption=caption)

    widths = {len(r) for r in grid}
    n_cols = max(widths)
    if len(widths) > 1:
        warnings.append(f"{table_id}：行宽不齐（{sorted(widths)}），已右侧补空对齐")
        grid = [r + [""] * (n_cols - len(r)) for r in grid]

    first = grid[0]
    if len(grid) >= 2 and all(c and text_weight(c) <= 30 for c in first):
        header, rows = first, grid[1:]
    else:
        header, rows = [], grid
        warnings.append(f"{table_id}：未识别出表头，整表按数据行处理")

    xml = tbl._tbl.xml
    if "vMerge" in xml or "gridSpan" in xml:
        warnings.append(f"{table_id}：含合并单元格，按重复值展开（单元格值仍为 verbatim）")

    return DocTable(
        table_id=table_id, section_id=section_id,
        n_rows=len(grid), n_cols=n_cols,
        header=header, rows=rows, caption=caption,
        src_index=src_index, col_widths=_grid_col_widths(tbl),
    )


def parse_docx(path: Path) -> DocTree:
    path = Path(path)
    if not path.exists():
        raise ParseError(f"文件不存在：{path}")
    try:
        doc = DocxDocument(str(path))
    except Exception as e:
        raise ParseError(f"docx 打开失败：{e}") from e

    warnings: list[str] = []
    tables: list[DocTable] = []
    images: list[DocImage] = []
    flat_sections: list[DocSection] = []

    # .doc 经 COM 转换产物名为 *.converted.docx，stem 尾巴不得漏进标题
    stem = path.stem.removesuffix(".converted")
    root = DocSection(section_id="sec-0000", level=0,
                      title=clean_text(doc.core_properties.title or "") or stem or "正文")
    stack: list[DocSection] = [root]
    current = root

    blk_n = sec_n = tbl_n = img_n = 0
    parts: list[str] = []
    last_para = ""
    txbx_warned = False

    for pos, item in _iter_blocks(doc):
        if isinstance(item, Paragraph):
            text = clean_text(item.text)
            style = item.style.name or ""
            if not text:
                txbx = _txbx_text(item._p)
                if txbx:
                    # 文本框设计版（简历等）：文字可仿写，图形排版无法保留
                    if not txbx_warned:
                        warnings.append(
                            "文档含文本框：内容按锚点顺序提取，原设计排版无法保留")
                        txbx_warned = True
                    for line in txbx:
                        blk_n += 1
                        current.blocks.append(DocBlock(
                            block_id=f"blk-{blk_n:04d}", kind="para", text=line))
                        parts.append(line)
                        last_para = line
                    continue
                img = _image_of(item._p, f"img-{img_n + 1:03d}",
                                current.section_id, pos, warnings)
                if img is not None:
                    img_n += 1
                    images.append(img)
                    blk_n += 1
                    current.blocks.append(DocBlock(
                        block_id=f"blk-{blk_n:04d}", kind="image", image_id=img.image_id))
                continue
            if style in ("Title", "标题") and root.title in ("", stem, "正文"):
                root.title = text
                parts.append(text)
                continue
            level = _heading_level(item)
            if level is not None:
                sec_n += 1
                while len(stack) > level:
                    stack.pop()
                sec = DocSection(section_id=f"sec-{sec_n:04d}", level=level, title=text)
                stack[-1].subsections.append(sec)
                flat_sections.append(sec)
                stack.append(sec)
                current = sec
                parts.append(text)
                last_para = text
                continue
            blk_n += 1
            kind = "list_item" if _is_list_item(item) else "para"
            current.blocks.append(DocBlock(block_id=f"blk-{blk_n:04d}", kind=kind, text=text))
            parts.append(text)
            last_para = text
        else:
            tbl_n += 1
            table_id = f"tbl-{tbl_n:03d}"
            caption = last_para if _CAPTION.match(last_para) else None
            t = _extract_table(item, table_id, current.section_id, caption,
                               warnings, src_index=tbl_n - 1)
            tables.append(t)
            blk_n += 1
            current.blocks.append(DocBlock(block_id=f"blk-{blk_n:04d}", kind="table", table_id=table_id))
            parts.append(t.flat_text())
            last_para = ""

    full_text = "\n".join(parts)
    if not full_text:
        warnings.append("文档无文本内容")

    return DocTree(
        meta=DocMeta(source_format="docx", source_name=path.name,
                     n_chars=len(full_text), parse_warnings=warnings),
        sections=[root],
        tables=tables,
        full_text=full_text,
        images=images,
    )
