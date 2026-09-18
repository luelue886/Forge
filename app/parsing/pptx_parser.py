from __future__ import annotations

from pathlib import Path

from pptx import Presentation as PptxPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from app.parsing.base import ParseError, clean_text
from app.schema.doctree import DocBlock, DocMeta, DocSection, DocTable, DocTree
from app.schema.textlen import text_weight

# 章节头判定：标题占位符之外没有实质正文（正文权重 ≤6 视为页脚/日期类装饰）
_HEADER_BODY_MAX = 6


def _is_bullet(p) -> bool:
    pPr = p._p.pPr
    if pPr is None:
        return False
    return (pPr.find(qn("a:buChar")) is not None
            or pPr.find(qn("a:buAutoNum")) is not None)


def _iter_shapes(shapes):
    for shp in shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shp.shapes)
        else:
            yield shp


def _slide_parts(slide) -> tuple[str, list[tuple[str, bool]], list[list[list[str]]]]:
    """(title, [(正文段落, 是否列表项)], 表格网格列表)。正文按 (top, left) 阅读序。"""
    title = ""
    body: list[tuple[tuple[int, int], str, bool]] = []
    grids: list[list[list[str]]] = []

    t = slide.shapes.title
    if t is not None:
        title = clean_text(t.text_frame.text)

    for shp in _iter_shapes(slide.shapes):
        if shp == t:
            continue
        if getattr(shp, "has_table", False):
            grids.append([[clean_text(c.text) for c in row.cells] for row in shp.table.rows])
        elif getattr(shp, "has_text_frame", False):
            for p in shp.text_frame.paragraphs:
                text = clean_text(p.text)
                if text:
                    body.append(((shp.top or 0, shp.left or 0), text, _is_bullet(p)))

    body.sort(key=lambda x: x[0])
    return title, [(b[1], b[2]) for b in body], grids


def _norm_grid(grid: list[list[str]], table_id: str,
               warnings: list[str]) -> tuple[int, list[str], list[list[str]]]:
    widths = {len(r) for r in grid}
    n_cols = max(widths)
    if len(widths) > 1:
        warnings.append(f"{table_id}：行宽不齐（{sorted(widths)}），已右侧补空对齐")
        grid = [r + [""] * (n_cols - len(r)) for r in grid]
    first = grid[0]
    if len(grid) >= 2 and all(c and text_weight(c) <= 30 for c in first):
        header, rows = first, grid[1:]
    else:
        header, rows = [], grid
        warnings.append(f"{table_id}：未识别出表头，整表按数据行处理")
    return n_cols, header, rows


def parse_pptx(path: Path) -> DocTree:
    path = Path(path)
    if not path.exists():
        raise ParseError(f"文件不存在：{path}")
    try:
        prs = PptxPresentation(str(path))
    except Exception as e:
        raise ParseError(f"pptx 打开失败：{e}") from e

    warnings: list[str] = []
    tables: list[DocTable] = []
    flat_sections: list[DocSection] = []

    root = DocSection(section_id="sec-0000", level=0,
                      title=clean_text(prs.core_properties.title or "") or path.stem or "正文")
    current = root
    blk_n = sec_n = tbl_n = 0
    parts: list[str] = []

    def open_section(title: str) -> DocSection:
        nonlocal sec_n, current
        sec_n += 1
        sec = DocSection(section_id=f"sec-{sec_n:04d}", level=1,
                         title=title or f"第 {sec_n} 节")
        root.subsections.append(sec)
        flat_sections.append(sec)
        current = sec
        parts.append(sec.title)
        return sec

    def add_block(kind: str, text: str, table_id: str | None = None) -> None:
        nonlocal blk_n
        blk_n += 1
        current.blocks.append(DocBlock(block_id=f"blk-{blk_n:04d}", kind=kind,
                                       text=text, table_id=table_id))
        parts.append(text)

    for idx, slide in enumerate(prs.slides, 1):
        title, body, grids = _slide_parts(slide)

        if idx == 1:  # 封面：标题作为文档题名，正文（副标题等）挂根节点
            if title and root.title in ("", path.stem, "正文"):
                root.title = title
                parts.append(title)
            for text, is_bullet in body:
                add_block("list_item" if is_bullet else "para", text)
            body = []
            if not grids:
                continue
            title = ""
        if not title and not body and not grids:
            continue

        body_weight = sum(text_weight(b) for b, _ in body)
        if title and not grids and body_weight <= _HEADER_BODY_MAX:
            open_section(title)  # 只有标题的页 → 章节头
            continue

        if current is root and (title or body or grids):
            open_section(title or f"第 {idx} 页")  # 尚无章节时按页开隐式章节
        elif title:
            add_block("para", title)  # 内容页标题降级为普通块，保持 verbatim

        for text, is_bullet in body:
            add_block("list_item" if is_bullet else "para", text)

        for grid in grids:
            if not grid or not grid[0]:
                warnings.append(f"第 {idx} 页：空表格，已跳过")
                continue
            tbl_n += 1
            table_id = f"tbl-{tbl_n:03d}"
            n_cols, header, rows = _norm_grid(grid, table_id, warnings)
            tables.append(DocTable(table_id=table_id, section_id=current.section_id,
                                   n_rows=len(grid), n_cols=n_cols,
                                   header=list(header), rows=[list(r) for r in rows],
                                   caption=title or None))
            add_block("table", "", table_id)
            parts.append(tables[-1].flat_text())

    full_text = "\n".join(parts)
    if not full_text:
        warnings.append("演示文稿无文本内容")

    return DocTree(
        meta=DocMeta(source_format="pptx", source_name=path.name,
                     n_chars=len(full_text), parse_warnings=warnings),
        sections=[root],
        tables=tables,
        full_text=full_text,
    )
