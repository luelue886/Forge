from __future__ import annotations

import atexit
import concurrent.futures
import logging
import queue
import subprocess
import threading
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from app.config import get_settings
from app.render import design
from app.schema.enums import IssueCode, Severity, SlideType
from app.schema.qa import QAIssue
from app.schema.slideir import Deck, SlideIR

log = logging.getLogger(__name__)

SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000
_BOUNDS_TOL_EMU = 25400  # 2pt


class ComError(RuntimeError):
    pass


class ComService:
    """Office COM 单实例：专用 STA 线程 + 队列。绝不附身用户已开的应用。

    progid 如 "PowerPoint.Application" / "Word.Application"；超时强杀对应进程后冷启动重试。
    """

    def __init__(self, progid: str = "PowerPoint.Application",
                 kill_exe: str = "POWERPNT.EXE", timeout_s: float = 120.0):
        self.progid = progid
        self._kill_exe = kill_exe
        self.timeout_s = timeout_s
        self._queue: queue.Queue[tuple | None] = queue.Queue()
        self._thread: threading.Thread | None = None
        self._lock = threading.Lock()
        self._app = None
        atexit.register(self.stop)

    # ---- 生命周期 ----

    def _ensure_thread(self) -> None:
        with self._lock:
            if self._thread is None or not self._thread.is_alive():
                self._thread = threading.Thread(target=self._worker, name="com-sta", daemon=True)
                self._thread.start()

    def _worker(self) -> None:
        import pythoncom

        pythoncom.CoInitialize()
        while True:
            item = self._queue.get()
            if item is None:
                break
            fn, future = item
            try:
                future.set_result(fn(self))
            except BaseException as e:  # noqa: BLE001 — 必须回传到 future
                if not future.done():
                    future.set_exception(e)
        self._quit_app()

    def _ensure_app(self):
        if self._app is None:
            import win32com.client

            self._app = win32com.client.DispatchEx(self.progid)
            try:
                if self.progid == "Word.Application":
                    self._app.Visible = False
                    self._app.DisplayAlerts = 0  # wdAlertsNone
                else:
                    self._app.DisplayAlerts = 1  # ppAlertsNone
            except Exception:
                pass
        return self._app

    def _quit_app(self) -> None:
        app, self._app = self._app, None
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass

    def _hard_kill(self) -> None:
        subprocess.run(
            ["taskkill", "/IM", self._kill_exe, "/F"],
            capture_output=True, creationflags=subprocess.CREATE_NO_WINDOW,
        )
        self._app = None

    # ---- 对外 API ----

    def run(self, fn, timeout_s: float | None = None):
        """在 STA 线程内执行 fn(service)。超时→强杀→冷启动重试一次。"""
        timeout = timeout_s or self.timeout_s
        last_err: Exception | None = None
        for attempt in (1, 2):
            self._ensure_thread()
            fut: concurrent.futures.Future = concurrent.futures.Future()
            self._queue.put((fn, fut))
            try:
                return fut.result(timeout)
            except concurrent.futures.TimeoutError:
                log.warning("COM 超时（第 %d 次），强制重启 PowerPoint", attempt)
                last_err = ComError(f"COM 操作超时（{timeout}s）")
                self._hard_kill()
            except Exception as e:
                if attempt == 1:
                    log.warning("COM 异常：%s；重置后重试一次", e)
                    last_err = e
                    self._hard_kill()
                else:
                    raise
        raise last_err or ComError("COM 操作失败")

    def stop(self) -> None:
        with self._lock:
            thread = self._thread
            self._thread = None
        if thread is not None and thread.is_alive():
            self._queue.put(None)
            thread.join(timeout=5)
        self._quit_app()


_service: ComService | None = None
_word_service: ComService | None = None


def get_com_service() -> ComService:
    global _service
    if _service is None:
        _service = ComService(timeout_s=get_settings().com_timeout_s)
    return _service


def get_word_com_service() -> ComService:
    global _word_service
    if _word_service is None:
        _word_service = ComService(progid="Word.Application", kill_exe="WINWORD.EXE",
                                   timeout_s=get_settings().com_timeout_s)
    return _word_service


# ---- PNG 导出 ----

def export_pngs(pptx_path: Path, out_dir: Path, width_px: int = 1280) -> list[Path]:
    """整份导出 page_NN.png，返回按页序的路径列表。"""

    def _do(svc: ComService):
        app = svc._ensure_app()
        prs = app.Presentations.Open(str(Path(pptx_path).resolve()), True, False, False)
        try:
            out = Path(out_dir).resolve()
            out.mkdir(parents=True, exist_ok=True)
            paths: list[Path] = []
            height_px = int(width_px * 9 / 16)
            for i in range(1, prs.Slides.Count + 1):
                png = out / f"page_{i:02d}.png"
                prs.Slides(i).Export(str(png), "PNG", width_px, height_px)
                paths.append(png)
            return paths
        finally:
            prs.Close()

    return get_com_service().run(_do)


def export_page(pptx_path: Path, page_no: int, out_png: Path, width_px: int = 1280) -> Path:
    """单页导出（单页重生成链路用）。"""

    def _do(svc: ComService):
        app = svc._ensure_app()
        prs = app.Presentations.Open(str(Path(pptx_path).resolve()), True, False, False)
        try:
            if not (1 <= page_no <= prs.Slides.Count):
                raise ComError(f"页码 {page_no} 超出范围 1-{prs.Slides.Count}")
            target = Path(out_png).resolve()
            target.parent.mkdir(parents=True, exist_ok=True)
            prs.Slides(page_no).Export(str(target), "PNG", width_px, int(width_px * 9 / 16))
            return target
        finally:
            prs.Close()

    return get_com_service().run(_do)


# ---- 旧版格式转换 ----

def convert_to_pptx(ppt_path: Path, out_path: Path | None = None) -> Path:
    """.ppt（PowerPoint 97-2003）→ .pptx；out_path 缺省写到源文件旁。走同一 COM 单实例队列。"""

    def _do(svc: ComService):
        app = svc._ensure_app()
        prs = app.Presentations.Open(str(Path(ppt_path).resolve()), True, False, False)
        try:
            out = Path(out_path) if out_path else Path(ppt_path).with_suffix(".pptx")
            prs.SaveAs(str(out.resolve()), 24)  # ppSaveAsOpenXMLPresentation
            return out
        finally:
            prs.Close()

    return get_com_service().run(_do)


# ---- Word：docx → PDF ----

def export_docx_pdf(docx_path: Path, out_pdf: Path | None = None) -> Path:
    """docx → PDF（Word COM，wdFormatPDF=17），版式与 docx 一致。走 Word 专用 STA 队列。"""

    def _do(svc: ComService):
        app = svc._ensure_app()
        doc = app.Documents.Open(str(Path(docx_path).resolve()),
                                 False, True, False)  # ConfirmConversions/ReadOnly/AddToRecentFiles
        try:
            out = Path(out_pdf) if out_pdf else Path(docx_path).with_suffix(".pdf")
            doc.SaveAs2(str(out.resolve()), FileFormat=17)
            return out
        finally:
            doc.Close(False)

    return get_word_com_service().run(_do)


# ---- 渲染清单检查（对已保存 pptx 的确定性核对）----

def _expected_text_shapes(ir: SlideIR) -> set[str]:
    names = set(design.texts_for(ir).keys())
    if ir.slide_type is SlideType.TABLE and ir.table is not None:
        names.add("tbl_frame")
    return names


def checklist(deck: Deck, pptx_path: Path) -> list[QAIssue]:
    """核对渲染产物与 SlideIR：页数、命名形状存在且非空、形状不越画布。"""
    issues: list[QAIssue] = []

    def add(code: IssueCode, page_no: int, detail: str):
        issues.append(QAIssue(
            code=code, severity=Severity.BLOCKING, page_no=page_no, detail=detail,
        ))

    prs = Presentation(str(pptx_path))
    if len(prs.slides) != len(deck.slides):
        add(IssueCode.E_RENDER_MISMATCH, 0,
            f"页数不符：pptx {len(prs.slides)} 页，SlideIR {len(deck.slides)} 页")
        return issues

    for ir, slide in zip(deck.slides, prs.slides):
        expected = _expected_text_shapes(ir)
        found: dict[str, str] = {}
        for shape in slide.shapes:
            if shape.name in expected:
                if shape.has_text_frame:
                    found[shape.name] = shape.text_frame.text.strip()
                else:
                    found[shape.name] = "<table>"
            try:
                l, t = shape.left, shape.top
                w, h = shape.width, shape.height
            except Exception:
                continue
            if l is None or t is None or w is None or h is None:
                continue
            if (l < -_BOUNDS_TOL_EMU or t < -_BOUNDS_TOL_EMU
                    or l + w > SLIDE_W_EMU + _BOUNDS_TOL_EMU
                    or t + h > SLIDE_H_EMU + _BOUNDS_TOL_EMU):
                add(IssueCode.E_RENDER_MISMATCH, ir.page_no,
                    f"形状 {shape.name} 越画布：({Emu(l).inches:.2f},{Emu(t).inches:.2f}) "
                    f"+({Emu(w).inches:.2f}x{Emu(h).inches:.2f})")
        for name in sorted(expected - set(found)):
            add(IssueCode.E_RENDER_MISMATCH, ir.page_no, f"缺少命名形状 {name}")
        for name, text in sorted(found.items()):
            if name != "tbl_frame" and not text:
                add(IssueCode.E_EMPTY_PLACEHOLDER, ir.page_no, f"{name} 内容为空")
    return issues
