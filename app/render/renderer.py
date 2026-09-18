from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from app.render import design
from app.render.design import RectSpec, SlideLayout, TextSpec
from app.render.skins import Skin, load_skin
from app.schema.enums import SlideType
from app.schema.slideir import Deck, SlideIR

SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.upper())


def _color(skin: Skin, role: str) -> RGBColor:
    return _rgb(getattr(skin, role))


def _set_run(run, text: str, pt: float, color: RGBColor, bold: bool, font_name: str) -> None:
    run.text = text
    f = run.font
    f.size = Pt(pt)
    f.bold = bold
    f.name = font_name
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font_name)


def _bulletize(p) -> None:
    """项目符号 + 悬挂缩进。必须在设置 line_spacing/space_before 之后调用。"""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "228600")
    pPr.set("indent", "-228600")
    bu_font = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial", "pitchFamily": "34", "charset": "0"})
    bu_char = pPr.makeelement(qn("a:buChar"), {"char": "•"})
    pPr.append(bu_font)
    pPr.append(bu_char)


def _add_text(slide, name: str, spec: TextSpec, paras: list[str], skin: Skin):
    tb = slide.shapes.add_textbox(Inches(spec.x), Inches(spec.y), Inches(spec.w), Inches(spec.h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE  # 溢出必须显式暴露
    tf.margin_left = Inches(spec.inset)
    tf.margin_right = Inches(spec.inset)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = _ANCHOR[spec.anchor]
    color = _color(skin, spec.color_role)
    for i, text in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[spec.align]
        p.line_spacing = spec.spacing
        if i and spec.space_before:
            p.space_before = Pt(spec.space_before)
        _set_run(p.add_run(), text, spec.font_pt, color, spec.bold, skin.font)
        if spec.bullet:
            _bulletize(p)
    return tb


def _add_rect(slide, spec: RectSpec, skin: Skin):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if spec.rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(spec.x), Inches(spec.y), Inches(spec.w), Inches(spec.h))
    if spec.name:
        shp.name = spec.name
    shp.fill.solid()
    shp.fill.fore_color.rgb = _color(skin, spec.color_role)
    if spec.line_role:
        shp.line.color.rgb = _color(skin, spec.line_role)
        shp.line.width = Pt(1.0)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    if spec.rounded:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    return shp


def _set_bg(slide, skin: Skin, role: str) -> None:
    try:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _color(skin, role)
    except Exception:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(SLIDE_W_EMU), Emu(SLIDE_H_EMU))
        shp.name = "fx_bg"
        shp.fill.solid()
        shp.fill.fore_color.rgb = _color(skin, role)
        shp.line.fill.background()
        shp.shadow.inherit = False


def _add_footer(slide, page_no: int, total: int, deck_title: str, skin: Skin) -> None:
    _add_rect(slide, design.FOOT_LINE, skin)
    _add_text(slide, "fx_foot_l", design.FOOT_L, [deck_title[:30]], skin)
    _add_text(slide, "fx_foot_r", design.FOOT_R, [f"{page_no} / {total}"], skin)


def _add_table(slide, ir: SlideIR, skin: Skin) -> None:
    t = ir.table
    assert t is not None
    n_cols = max([len(t.header)] + [len(r) for r in t.rows]) if (t.header or t.rows) else 0
    n_rows = len(t.rows) + (1 if t.header else 0)
    if n_rows == 0 or n_cols == 0:
        return
    x, y, w, _h = design.TABLE_REGION

    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(min(_h, 0.45 + len(t.rows) * 0.42)))
    gf.name = "tbl_frame"
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    tblPr = tbl._tbl.tblPr
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)

    weights = ([1.25] + [1.0] * (n_cols - 1)) if n_cols >= 3 else [1.0] * n_cols
    total_w = sum(weights)
    for i, col in enumerate(tbl.columns):
        col.width = Emu(int(Inches(w) * weights[i] / total_w))
    for i, row in enumerate(tbl.rows):
        row.height = Inches(design.TABLE_ROW_H_HEADER if (i == 0 and t.header) else design.TABLE_ROW_H_DATA)

    def _fill_cell(cell, text: str, fill: RGBColor, color_role: str, header: bool, first_col: bool) -> None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.margin_top = Inches(0.03)
        cell.margin_bottom = Inches(0.03)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if (header or not first_col) else PP_ALIGN.LEFT
        _set_run(p.add_run(), text, design.TABLE_FONT_PT, _color(skin, color_role), header, skin.font)

    for j, htext in enumerate(t.header):
        _fill_cell(tbl.cell(0, j), htext, _color(skin, "header_bg"), "header_text", header=True, first_col=False)
    offset = 1 if t.header else 0
    for i, row in enumerate(t.rows):
        fill = _color(skin, "zebra") if i % 2 == 1 else _color(skin, "bg")
        for j in range(n_cols):
            text = row[j] if j < len(row) else ""
            _fill_cell(tbl.cell(i + offset, j), text, fill, "text", header=False, first_col=(j == 0))


def render_slide(slide, ir: SlideIR, skin: Skin, deck_title: str, total_pages: int) -> None:
    layout: SlideLayout = design.layout_for(ir)
    _set_bg(slide, skin, layout.bg_role)
    for spec in layout.rects:
        _add_rect(slide, spec, skin)
    tmap = design.texts_for(ir)
    for name, spec in layout.texts.items():
        paras = tmap.get(name)
        if paras:
            _add_text(slide, name, spec, paras, skin)
    if ir.slide_type is SlideType.TABLE and ir.table is not None:
        _add_table(slide, ir, skin)
    if layout.footer:
        _add_footer(slide, ir.page_no, total_pages, deck_title, skin)
    if ir.note:
        slide.notes_slide.notes_text_frame.text = ir.note


def render_deck(deck: Deck, skin: Skin) -> Presentation:
    """pptx 是 SlideIR 的纯函数：同 IR + 同皮肤 → 逐字节等价的排版。"""
    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = prs.slide_layouts[6]
    total = len(deck.slides)
    for ir in deck.slides:
        slide = prs.slides.add_slide(blank)
        render_slide(slide, ir, skin, deck.meta.title, total)
    return prs


def render_to_file(deck: Deck, skin_name: str, out_path: Path) -> Presentation:
    skin = load_skin(skin_name)
    prs = render_deck(deck, skin)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return prs
