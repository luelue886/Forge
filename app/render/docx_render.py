"""DocIR → output.docx 纯函数渲染器（python-docx）。

source（源 docx 路径）存在时，表格走 docx_copy 深拷贝源 XML（排版 100% 保真）
+ 单元格文字替换；否则按规范样式重建，有 col_widths 时设列宽比例。
"""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from app.render import docstyle as st
from app.render import docx_copy
from app.schema.docir import (
    ClosingBlock,
    DocIR,
    DocTitleBlock,
    HeadingBlock,
    ImageBlock,
    ParaBlock,
    SalutationBlock,
    SignatureBlock,
    TableBlock,
)

log = logging.getLogger(__name__)


def _set_font(run, east_asia: str, size: Pt, bold: bool) -> None:
    run.font.name = st.FONT_LATIN
    run.font.size = size
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def _para(doc: Document, east_asia: str, size: Pt, bold: bool,
          align=None, first_line_indent=None, space_before=None,
          space_after=None, line_spacing=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
        pf.element.get_or_add_pPr().get_or_add_ind().set(
            qn("w:firstLineChars"), str(st.FIRST_LINE_CHARS * 100))
    if space_before is not None:
        pf.space_before = space_before
    if space_after is not None:
        pf.space_after = space_after
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    return p, east_asia, size, bold


def _add_run(p, east_asia, size, bold, text: str):
    run = p.add_run(text)
    _set_font(run, east_asia, size, bold)
    return run


def _setup_page(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = st.PAGE_WIDTH, st.PAGE_HEIGHT
    sec.top_margin, sec.bottom_margin = st.MARGIN_TOP, st.MARGIN_BOTTOM
    sec.left_margin, sec.right_margin = st.MARGIN_LEFT, st.MARGIN_RIGHT


def _add_page_number(doc: Document) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.size = st.SIZE_WU
    for tag, attrs, text in (
            ("w:fldChar", {"w:fldCharType": "begin"}, None),
            ("w:instrText", None, "PAGE"),
            ("w:fldChar", {"w:fldCharType": "end"}, None)):
        el = OxmlElement(tag)
        if attrs:
            for k, v in attrs.items():
                el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)


def _shade(cell, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(st.TABLE_BORDER_SIZE))
        el.set(qn("w:color"), st.TABLE_BORDER_COLOR)
        borders.append(el)
    # schema 顺序：tblBorders 须位于 tblLook 之前
    tbl_look = tbl_pr.find(qn("w:tblLook"))
    if tbl_look is not None:
        tbl_look.addprevious(borders)
    else:
        tbl_pr.append(borders)


def _render_table(doc: Document, t: TableBlock) -> None:
    n_cols = len(t.header) or (len(t.rows[0]) if t.rows else 0)
    if n_cols == 0:
        return
    font, size, _ = st.STYLE_TABLE
    if n_cols >= st.WIDE_TABLE_COLS:
        size = st.SIZE_XIAO_WU

    n_rows = len(t.rows) + (1 if t.header else 0)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    if t.col_widths and len(t.col_widths) == n_cols:
        table.autofit = False
        content_w = st.PAGE_WIDTH - st.MARGIN_LEFT - st.MARGIN_RIGHT
        for c, ratio in enumerate(t.col_widths):
            w = Emu(int(content_w * ratio))
            table.columns[c].width = w  # tblGrid/gridCol（固定布局的列宽基准）
            for cell in table.columns[c].cells:
                cell.width = w             # tcW（合并格的宽度基准）
    else:
        table.autofit = True
    _borders(table)

    # 合并区 [[r,c,rs,cs]]（含表头行网格坐标）：维度合法且互不重叠才应用。
    # 有合并时禁表头 shading 与斑马纹——首行常为整行标题合并，展开同值会被
    # 误判表头；斑马纹打断合并区视觉整体感。
    merged = False
    if t.merges:
        occupied: set[tuple[int, int]] = set()
        ok: list[tuple[int, int, int, int]] = []
        for m in t.merges:
            if len(m) != 4:
                continue
            r, c, rs, cs = (int(m[0]), int(m[1]), int(m[2]), int(m[3]))
            if rs < 1 or cs < 1 or r + rs > n_rows or c + cs > n_cols \
                    or (rs == 1 and cs == 1):
                continue
            region = [(i, j) for i in range(r, r + rs) for j in range(c, c + cs)]
            if any(p in occupied for p in region):
                continue
            occupied.update(region)
            ok.append((r, c, rs, cs))
        has_widths = bool(t.col_widths) and len(t.col_widths) == n_cols
        content_w = st.PAGE_WIDTH - st.MARGIN_LEFT - st.MARGIN_RIGHT
        for r, c, rs, cs in ok:
            merged_cell = table.cell(r, c).merge(
                table.cell(r + rs - 1, c + cs - 1))
            if has_widths:
                # merge 后存活 tc 的 tcW 只剩首列宽，显式补成覆盖列之和
                merged_cell.width = Emu(int(
                    content_w * sum(t.col_widths[c:c + cs])))
        merged = bool(ok)

    if t.row_heights and len(t.row_heights) == n_rows:
        for i, h_pt in enumerate(t.row_heights):
            tr_pr = table.rows[i]._tr.get_or_add_trPr()
            tr_h = OxmlElement("w:trHeight")
            tr_h.set(qn("w:val"), str(int(round(h_pt * 20))))  # pt → twips
            tr_h.set(qn("w:hRule"), "atLeast")
            tr_pr.append(tr_h)

    # 合并后同区各坐标返回同一物理格：按 tc 元素去重（表头/数据行统一）。
    # 必须持有元素引用而非 id()——lxml 代理被 GC 后地址会复用造成误判。
    written: set = set()

    def _put(r: int, c: int, text: str, bold: bool, fill: str | None) -> None:
        cell = table.cell(r, c)
        if cell._tc in written:
            return
        written.add(cell._tc)
        if fill and not merged:
            _shade(cell, fill)
        _add_run(cell.paragraphs[0], font, size, bold, text)

    if t.header:
        for c, text in enumerate(t.header):
            _put(0, c, text, True, st.TABLE_HEADER_FILL)
    r0 = 1 if t.header else 0
    for i, row in enumerate(t.rows):
        for c in range(n_cols):
            text = row[c] if c < len(row) else ""
            _put(r0 + i, c, text, False,
                 st.TABLE_ZEBRA_FILL if i % 2 == 1 else None)


def _copy_source_table(src_doc, out_doc: Document, b: TableBlock) -> bool:
    """源表格 XML 搬运 + 目标网格差异格改写；源序号缺失/越界返回 False 走规范重建。"""
    if b.src_index is None:
        return False
    tbl_el = docx_copy.source_table(src_doc, b.src_index)
    if tbl_el is None:
        return False
    new_el = docx_copy.copy_table(src_doc, out_doc, tbl_el)
    docx_copy.replace_table_texts(new_el, out_doc, b.header, b.rows)
    return True


def _insert_pdf_image(doc: Document, source: Path, b: ImageBlock) -> bool:
    """PDF 源：pymupdf 按 bbox 区域渲染位图嵌入（规避 xref=0/smask/跨库序号错位）。"""
    import io

    import pymupdf

    try:
        with pymupdf.open(str(source)) as pdf:
            if b.page is None or b.page < 1 or b.page > len(pdf):
                return False
            page = pdf[b.page - 1]
            clip = pymupdf.Rect(b.bbox) if b.bbox else None
            pix = page.get_pixmap(clip=clip, matrix=pymupdf.Matrix(2, 2))
            png = pix.tobytes("png")
            # 相对源页宽缩放到版心宽，保持原图在页面中的占比
            src_w = clip.width if clip else page.rect.width
            content_w = st.PAGE_WIDTH - st.MARGIN_LEFT - st.MARGIN_RIGHT
            width = Emu(int(content_w * src_w / page.rect.width))
    except Exception as e:  # noqa: BLE001 — 渲染失败降级跳图，不崩整篇
        log.warning("pdf 图片 %s 区域渲染失败：%s", b.image_id, e)
        return False

    doc.add_picture(io.BytesIO(png), width=width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return True


def render_docir_to_docx(doc: DocIR, out: Path, source: Path | None = None) -> Path:
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    _setup_page(d)

    src_doc = None
    if source is not None and Path(source).suffix.lower() == ".docx" \
            and Path(source).exists():
        src_doc = docx_copy.open_source(Path(source))

    for b in doc.blocks:
        if isinstance(b, DocTitleBlock):
            p, font, size, bold = _para(
                d, *st.STYLE_TITLE, align=WD_ALIGN_PARAGRAPH.CENTER,
                space_after=st.TITLE_SPACE_AFTER, line_spacing=1.5)
            _add_run(p, font, size, bold, b.text)
        elif isinstance(b, HeadingBlock):
            style = st.STYLE_H1 if b.level == 1 else st.STYLE_H2
            before = st.H1_SPACE_BEFORE if b.level == 1 else st.H2_SPACE_BEFORE
            after = st.H1_SPACE_AFTER if b.level == 1 else st.H2_SPACE_AFTER
            p, font, size, bold = _para(
                d, *style, space_before=before, space_after=after,
                line_spacing=1.5)
            _add_run(p, font, size, bold, b.text)
        elif isinstance(b, SalutationBlock):
            p, font, size, bold = _para(d, *st.STYLE_BODY, line_spacing=1.5)
            _add_run(p, font, size, bold, b.text)  # 称呼顶格
        elif isinstance(b, ParaBlock):
            p, font, size, bold = _para(
                d, *st.STYLE_BODY, first_line_indent=st.FIRST_LINE_FALLBACK,
                space_after=st.PARA_SPACE_AFTER, line_spacing=st.BODY_LINE_SPACING)
            _add_run(p, font, size, bold, b.text)
        elif isinstance(b, ClosingBlock):
            for j, line in enumerate(b.lines):
                # 惯例：首行（此致）空两格，其余（敬礼）顶格
                indent = st.FIRST_LINE_FALLBACK if j == 0 else None
                p, font, size, bold = _para(
                    d, *st.STYLE_BODY, first_line_indent=indent,
                    line_spacing=1.5)
                _add_run(p, font, size, bold, line)
        elif isinstance(b, SignatureBlock):
            for line in (b.signer, b.date):
                if not line:
                    continue
                p, font, size, bold = _para(
                    d, *st.STYLE_BODY, align=WD_ALIGN_PARAGRAPH.RIGHT,
                    line_spacing=1.5)
                _add_run(p, font, size, bold, line)
        elif isinstance(b, TableBlock):
            if src_doc is None or not _copy_source_table(src_doc, d, b):
                _render_table(d, b)
        elif isinstance(b, ImageBlock):
            if src_doc is not None and b.body_index is not None:
                if not docx_copy.copy_image_paragraph(src_doc, d, b.body_index):
                    log.warning("图片 %s：源段落缺失，跳过", b.image_id)
            elif source is not None and Path(source).suffix.lower() == ".pdf":
                if not _insert_pdf_image(d, Path(source), b):
                    log.warning("图片 %s：PDF 区域渲染失败，跳过", b.image_id)
            else:
                log.warning("图片 %s：无可用源文件，跳过", b.image_id)

    _add_page_number(d)
    d.save(str(out))
    return out
