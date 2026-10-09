"""真模板构建：unzip → 复制映射页 → XML 文本替换/删除 → 删原页 → BFS 清孤儿 → zip → 后处理。

工艺约束（pptx skill 编辑路径）：
- 永不重排 p:presentation 子元素，只增删 sldIdLst 条目与对应 rels；
- 复制页剔除 notesSlide / tags 关系、页级 p:custDataLst 与 p:timing
  （tag 部件被多页共享非法，PowerPoint 拒开文件；演示文稿级 custDataLst
  单引用，保留。timing 的 spid 指向被删形状同样拒开——删形状的页无法
  逐个核对动画引用，整树剥离换取无动画但必开的产物）；
- 文本替换保留供体段落 pPr 与首 run rPr，格式/几何 100% 继承模板；
- 槽寻址与目录共用 shape_children 枚举，重走槽集不符即报错（寻址漂移防护）。
"""
from __future__ import annotations

import copy
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.render.tpl_catalog import (
    NS, _GRP, _HAS_CJK, _SP, _parse, _ph_key, _sp_text, shape_children,
)
from app.schema.slotfill import DeckBuild, PageFill
from app.schema.tplcatalog import TemplateCatalog

_A = f"{{{NS['a']}}}"
_P = f"{{{NS['p']}}}"
_R = f"{{{NS['r']}}}"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

_SLIDE_CT = "application/vnd.openxmlformats-officedocument.presentationml.slide+xml"
_SLIDE_REL = f"{NS['r']}/slide"
_THEME_RE = re.compile(r"^ppt/theme/theme\d+\.xml$")
_NEUTRAL_GRAY = RGBColor(0x66, 0x66, 0x66)


class BuildError(Exception):
    pass


def build_deck(tpl_path: Path, catalog: TemplateCatalog, deck: DeckBuild,
               out_path: Path) -> Path:
    """按 DeckBuild 逐页复制模板页并填充，产出最终 pptx。"""
    if not deck.pages:
        raise BuildError("DeckBuild 页列表为空")
    if deck.meta.template and deck.meta.template != catalog.stem:
        raise BuildError(
            f"DeckBuild 模板 {deck.meta.template} 与目录 {catalog.stem} 不一致")
    tpl_path, out_path = Path(tpl_path), Path(out_path)
    with tempfile.TemporaryDirectory(prefix="forge_tpl_build_") as td:
        work = Path(td) / "unpacked"
        with zipfile.ZipFile(tpl_path) as z:
            z.extractall(work)
        parts = _duplicate_pages(work, catalog, deck.pages)
        for page, part in zip(deck.pages, parts):
            _apply_page(work, catalog, page, part)
        _cleanup_orphans(work)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        _zip_dir(work, out_path)
    _postprocess(out_path, catalog, deck)
    return out_path


# ---------- 复制记账 ----------

def _duplicate_pages(work: Path, catalog: TemplateCatalog,
                     pages: list[PageFill]) -> list[str]:
    pres_path = work / "ppt" / "presentation.xml"
    pres_rels_path = work / "ppt" / "_rels" / "presentation.xml.rels"
    ct_path = work / "[Content_Types].xml"
    pres = _parse(pres_path.read_bytes())
    pres_rels = _parse(pres_rels_path.read_bytes())
    ct = _parse(ct_path.read_bytes())
    sld_id_lst = pres.find(f"{_P}sldIdLst")
    if sld_id_lst is None:
        raise BuildError("模板缺少 sldIdLst")

    ordered = _ordered_slide_parts(pres, pres_rels)
    if len(ordered) != catalog.n_slides:
        raise BuildError(f"模板页数 {len(ordered)} 与目录 {catalog.n_slides} 不符")

    slides_dir = work / "ppt" / "slides"
    existing = sorted(
        int(m.group(1)) for f in slides_dir.glob("slide*.xml")
        if (m := re.match(r"slide(\d+)\.xml$", f.name)))
    next_no = existing[-1] + 1 if existing else 1
    max_rid = _max_rid(pres_rels)
    max_sldid = max((int(s.get("id", "0")) for s in sld_id_lst), default=255)

    originals = list(sld_id_lst)
    new_parts: list[str] = []
    for page in pages:
        if not 1 <= page.template_slide <= len(ordered):
            raise BuildError(
                f"p{page.page_no} template_slide={page.template_slide} 越界（1-{len(ordered)}）")
        src_part = ordered[page.template_slide - 1]
        dst_name = f"slide{next_no}.xml"
        shutil.copyfile(work / src_part, slides_dir / dst_name)
        _copy_slide_rels(work, src_part, dst_name)
        _add_override(ct, f"/ppt/slides/{dst_name}")
        max_rid += 1
        rid = f"rId{max_rid}"
        rel = etree.SubElement(pres_rels, f"{{{_REL_NS}}}Relationship")
        rel.set("Id", rid)
        rel.set("Type", _SLIDE_REL)
        rel.set("Target", f"slides/{dst_name}")
        max_sldid += 1
        sld = etree.SubElement(sld_id_lst, f"{_P}sldId")
        sld.set("id", str(max_sldid))
        sld.set(f"{_R}id", rid)
        new_parts.append(f"ppt/slides/{dst_name}")
        next_no += 1

    # 原模板页全部退场：sldIdLst 条目与 presentation rels 一并移除（孤儿清理由 BFS 收尾）
    for el in originals:
        rid = el.get(f"{_R}id")
        sld_id_lst.remove(el)
        if rid:
            for rel in pres_rels.findall(f"{{{_REL_NS}}}Relationship"):
                if rel.get("Id") == rid:
                    pres_rels.remove(rel)
                    break

    pres_path.write_bytes(_serialize(pres))
    pres_rels_path.write_bytes(_serialize(pres_rels))
    ct_path.write_bytes(_serialize(ct))
    return new_parts


def _ordered_slide_parts(pres, pres_rels) -> list[str]:
    rel_map = {
        rel.get("Id"): _resolve("ppt", rel.get("Target", ""))
        for rel in pres_rels.findall(f"{{{_REL_NS}}}Relationship")
    }
    out: list[str] = []
    for sld in pres.find(f"{_P}sldIdLst"):
        t = rel_map.get(sld.get(f"{_R}id") or "")
        if t and t.startswith("ppt/slides/"):
            out.append(t)
    return out


def _copy_slide_rels(work: Path, src_part: str, dst_name: str) -> None:
    src_name = src_part.rsplit("/", 1)[1]
    src_rels = work / "ppt" / "slides" / "_rels" / f"{src_name}.rels"
    if not src_rels.exists():
        return
    root = _parse(src_rels.read_bytes())
    for rel in list(root.findall(f"{{{_REL_NS}}}Relationship")):
        if (rel.get("Type") or "").endswith("/notesSlide"):
            root.remove(rel)  # 共享 notesSlide 非法；后处理用 python-pptx 另建合法备注
        elif (rel.get("Type") or "").endswith("/tags"):
            root.remove(rel)  # 复制页共享 tag 部件非法（PowerPoint 拒开文件）
    dst = work / "ppt" / "slides" / "_rels" / f"{dst_name}.rels"
    dst.write_bytes(_serialize(root))


def _add_override(ct_root, part_name: str) -> None:
    ov = etree.SubElement(ct_root, f"{{{_CT_NS}}}Override")
    ov.set("PartName", part_name)
    ov.set("ContentType", _SLIDE_CT)


def _max_rid(rels_root) -> int:
    mx = 0
    for rel in rels_root.findall(f"{{{_REL_NS}}}Relationship"):
        m = re.match(r"rId(\d+)$", rel.get("Id", ""))
        if m:
            mx = max(mx, int(m.group(1)))
    return mx


# ---------- 槽填充与删除 ----------

def _apply_page(work: Path, catalog: TemplateCatalog, page: PageFill, part: str) -> None:
    cat_slide = catalog.slide(page.template_slide)
    overlap = set(page.slot_values) & set(page.slot_deletes)
    if overlap:
        raise BuildError(f"p{page.page_no} 槽同时被填充与删除：{sorted(overlap)}")
    slide_path = work / part
    tree = _parse(slide_path.read_bytes())
    for el in tree.findall(f".//{_P}custDataLst"):
        el.getparent().remove(el)  # 多页复制共享同一 tag 部件非法（PowerPoint 拒开文件）
    timing = tree.find(f"{_P}timing")
    if timing is not None:
        tree.remove(timing)  # 动画 spid 指向被删形状 → PowerPoint 拒开；整树剥离
    sp_tree = tree.find(f"{_P}cSld/{_P}spTree")
    if sp_tree is None:
        raise BuildError(f"{part}: 无 spTree")
    walked = _walk_slot_elements(sp_tree)
    expected = {s.slot_id for s in cat_slide.slots}
    if set(walked) != expected:
        raise BuildError(
            f"{part}: 重走槽位 {len(walked)} 个与目录 {len(expected)} 个不符（寻址漂移）")
    unknown = (set(page.slot_values) | set(page.slot_deletes)) - expected
    if unknown:
        raise BuildError(f"p{page.page_no} 未知槽 id：{sorted(unknown)}")

    for sid, text in page.slot_values.items():
        _fill_slot(walked[sid], text)
    if page.slot_deletes:
        _delete_slots(sp_tree, walked, set(page.slot_deletes), cat_slide)
    slide_path.write_bytes(_serialize(tree))
    _prune_slide_rels(
        tree, work / "ppt" / "slides" / "_rels" / f"{Path(part).name}.rels")


def _walk_slot_elements(sp_tree) -> dict:
    """slot_id → sp 元素。与目录提取同一枚举与准入规则（有 txBody 且有文本或是占位符）。"""
    out: dict[str, object] = {}

    def walk(el, path: list[int]):
        for i, child in enumerate(shape_children(el)):
            p = path + [i]
            if child.tag == _GRP:
                walk(child, p)
            elif child.tag == _SP:
                tx = child.find(f"{_P}txBody")
                if tx is None:
                    continue
                text, _has_run = _sp_text(tx)
                if text.strip() or _ph_key(child) != ("", ""):
                    out[".".join(map(str, p))] = child

    walk(sp_tree, [])
    return out


def _fill_slot(sp, text: str) -> None:
    """替换 sp 文本：供体段落（首个含 a:r 的 a:p）的 pPr 与首 run rPr 全保留。"""
    tx = sp.find(f"{_P}txBody")
    if tx is None:
        raise BuildError("填充目标无 txBody")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n") if text else [""]
    paras = tx.findall(f"{_A}p")
    donor = next((p_ for p_ in paras if p_.find(f"{_A}r") is not None), None)
    if donor is None:
        donor = _make_donor(paras)
        tx.append(donor)
    # 修剪供体到 pPr + 单 run
    for extra in donor.findall(f"{_A}r")[1:]:
        donor.remove(extra)
    for tag in ("br", "fld"):
        for el in donor.findall(f"{_A}{tag}"):
            donor.remove(el)
    first_run = donor.find(f"{_A}r")
    _set_run_text(first_run, lines[0])
    for p_ in paras:
        if p_ is not donor:
            tx.remove(p_)
    prev = donor
    for line in lines[1:]:
        np_ = copy.deepcopy(donor)
        _set_run_text(np_.find(f"{_A}r"), line)
        prev.addnext(np_)
        prev = np_
    _ensure_ea_font(first_run, text)


def _make_donor(paras: list) -> object:
    """空占位符的最小供体段落：a:p/a:r/a:rPr/a:t，样式交给占位符继承链。"""
    p_ = etree.Element(f"{_A}p")
    r = etree.SubElement(p_, f"{_A}r")
    rPr = etree.SubElement(r, f"{_A}rPr")
    rPr.set("lang", "zh-CN")
    etree.SubElement(r, f"{_A}t")
    if paras:
        pPr = paras[0].find(f"{_A}pPr")
        if pPr is not None:
            p_.insert(0, copy.deepcopy(pPr))
    return p_


def _set_run_text(run, text: str) -> None:
    t = run.find(f"{_A}t")
    if t is None:
        t = etree.SubElement(run, f"{_A}t")
    t.text = text
    t.set(_XML_SPACE, "preserve")


def _ensure_ea_font(run, text: str) -> None:
    """英文模板 rPr 无 a:ea 时补中文字形，否则中文回退西文字体。"""
    if not text or not _HAS_CJK.search(text):
        return
    rPr = run.find(f"{_A}rPr")
    if rPr is None:
        rPr = etree.Element(f"{_A}rPr")
        run.insert(0, rPr)
    if rPr.find(f"{_A}ea") is not None:
        return
    ea = etree.Element(f"{_A}ea")
    ea.set("typeface", "微软雅黑")
    latin = rPr.find(f"{_A}latin")
    if latin is not None:
        latin.addnext(ea)
        return
    cs = rPr.find(f"{_A}cs")
    if cs is not None:
        cs.addprevious(ea)
    else:
        rPr.append(ea)


def _delete_slots(sp_tree, slot_elements: dict, del_ids: set[str], cat_slide) -> None:
    for sid in sorted(del_ids):
        el = slot_elements.get(sid)
        if el is None or not _in_tree(sp_tree, el):
            continue  # 已随外层组整体删除
        target = el
        path = _element_path(sp_tree, el)
        if path is None:
            continue
        for depth in range(1, len(path)):
            prefix = path[:depth]
            g = _element_at(sp_tree, prefix)
            if g is None or g.tag != _GRP:
                continue
            under = [s.slot_id for s in cat_slide.slots if s.path[:depth] == prefix]
            if under and set(under) <= del_ids:
                target = g  # 组内文本槽全部退场：整组删除（连带图片）
        parent = target.getparent()
        if parent is not None:
            parent.remove(target)


def _in_tree(sp_tree, el) -> bool:
    cur = el
    while cur is not None:
        if cur is sp_tree:
            return True
        cur = cur.getparent()
    return False


def _element_path(sp_tree, el) -> list[int] | None:
    path: list[int] = []
    cur = el
    while cur is not None and cur is not sp_tree:
        parent = cur.getparent()
        if parent is None:
            return None
        kids = shape_children(parent)
        try:
            path.append(kids.index(cur))
        except ValueError:
            return None
        cur = parent
    if cur is not sp_tree:
        return None
    return list(reversed(path))


def _element_at(sp_tree, path: list[int]):
    cur = sp_tree
    for i in path:
        kids = shape_children(cur)
        if i < 0 or i >= len(kids):
            return None
        cur = kids[i]
    return cur


def _prune_slide_rels(slide_tree, rels_path: Path) -> None:
    """删除形状后，剪除页 rels 中不再被 XML 引用的关系（版式关系除外）。"""
    if not rels_path.exists():
        return
    used: set[str] = set()
    for el in slide_tree.iter():
        for k, v in el.attrib.items():
            if k.startswith(_R):
                used.add(v)
    root = _parse(rels_path.read_bytes())
    changed = False
    for rel in list(root.findall(f"{{{_REL_NS}}}Relationship")):
        if rel.get("Id") in used:
            continue
        if (rel.get("Type") or "").endswith("/slideLayout"):
            continue  # 页 XML 不携带版式 rId，但关系必须保留
        root.remove(rel)
        changed = True
    if changed:
        rels_path.write_bytes(_serialize(root))


# ---------- 孤儿清理 ----------

def _cleanup_orphans(work: Path) -> None:
    """从 ppt/presentation.xml 做 BFS 可达性；docProps 与包根部件无条件保留。"""
    rels_cache: dict[str, dict[str, str]] = {}

    def rels_of(part: str) -> dict[str, str]:
        if part in rels_cache:
            return rels_cache[part]
        d, name = part.rsplit("/", 1)
        rp = work / d / "_rels" / f"{name}.rels"
        m: dict[str, str] = {}
        if rp.exists():
            for rel in _parse(rp.read_bytes()).findall(f"{{{_REL_NS}}}Relationship"):
                if rel.get("TargetMode") == "External":
                    continue
                m[rel.get("Id", "")] = _resolve(d, rel.get("Target", ""))
        rels_cache[part] = m
        return m

    reach: set[str] = set()
    queue = ["ppt/presentation.xml"]
    while queue:
        part = queue.pop()
        if part in reach:
            continue
        reach.add(part)
        for t in rels_of(part).values():
            if t and t not in reach:
                queue.append(t)

    keep: set[str] = {"[Content_Types].xml", "_rels/.rels"}
    for part in reach:
        keep.add(part)
        d, name = part.rsplit("/", 1)
        keep.add(f"{d}/_rels/{name}.rels")
    doc_props = work / "docProps"
    if doc_props.exists():
        for p in doc_props.rglob("*"):
            if p.is_file():
                keep.add(p.relative_to(work).as_posix())

    for p in sorted(work.rglob("*"), reverse=True):
        if p.is_file() and p.relative_to(work).as_posix() not in keep:
            p.unlink()
    for p in sorted(work.rglob("*"), reverse=True):
        if p.is_dir() and not any(p.iterdir()):
            p.rmdir()

    ct_path = work / "[Content_Types].xml"
    root = _parse(ct_path.read_bytes())
    removed = False
    for ov in list(root.findall(f"{{{_CT_NS}}}Override")):
        part = ov.get("PartName", "").lstrip("/")
        if part and part not in keep:
            root.remove(ov)
            removed = True
    if removed:
        ct_path.write_bytes(_serialize(root))


# ---------- zip 与后处理 ----------

def _zip_dir(work: Path, out_path: Path) -> None:
    if out_path.exists():
        out_path.unlink()
    ct = "[Content_Types].xml"
    with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as z:
        z.write(work / ct, ct)
        for p in sorted(work.rglob("*")):
            if not p.is_file():
                continue
            arc = p.relative_to(work).as_posix()
            if arc != ct:
                z.write(p, arc)


def _postprocess(out_path: Path, catalog: TemplateCatalog, deck: DeckBuild) -> None:
    prs = Presentation(str(out_path))
    slides = list(prs.slides)
    if len(slides) != len(deck.pages):
        raise BuildError(f"产物页数 {len(slides)} 与规划 {len(deck.pages)} 不符")
    accent = _accent1(out_path) or _NEUTRAL_GRAY
    slide_w_in = prs.slide_width / 914400
    slide_h_in = prs.slide_height / 914400
    for slide, page in zip(slides, deck.pages):
        if page.note and page.note.strip():
            slide.notes_slide.notes_text_frame.text = page.note
        if page.table is not None:
            _add_native_table(slide, page.table, catalog.slide(page.template_slide),
                              slide_w_in, slide_h_in, accent)
    prs.save(str(out_path))


def _accent1(pptx_path: Path) -> RGBColor | None:
    with zipfile.ZipFile(pptx_path) as z:
        theme = next((n for n in z.namelist() if _THEME_RE.match(n)), None)
        if theme is None:
            return None
        el = _parse(z.read(theme)).find(f"{_A}themeElements/{_A}clrScheme/{_A}accent1")
        if el is None:
            return None
        for tag in ("srgbClr", "sysClr"):
            c = el.find(f"{_A}{tag}")
            v = (c.get("val") if tag == "srgbClr" else c.get("lastClr")) if c is not None else None
            if v and len(v) == 6:
                try:
                    return RGBColor(int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))
                except ValueError:
                    return None
    return None


def _add_native_table(slide, table, cat_slide, slide_w_in: float,
                      slide_h_in: float, accent: RGBColor) -> None:
    header = list(table.header)
    rows = [list(r) for r in table.rows]
    widths = [len(header)] + [len(r) for r in rows]
    n_cols = max(widths) if widths else 1
    n_rows = (1 if header else 0) + max(len(rows), 1)
    x, y, w, h = _table_geometry(cat_slide, slide_w_in, slide_h_in)
    gf = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w), Inches(h))
    tbl = gf.table
    if header:
        for j, v in enumerate(header):
            _style_cell(tbl.cell(0, j), v, bold=True, size=14,
                        color=RGBColor(0xFF, 0xFF, 0xFF), fill=accent)
    for i, r in enumerate(rows, start=1 if header else 0):
        for j in range(n_cols):
            _style_cell(tbl.cell(i, j), r[j] if j < len(r) else "", bold=False, size=12)
    tbl.first_row = bool(header)
    tbl.horz_banding = True


def _style_cell(cell, text: str, bold: bool, size: int,
                color: RGBColor | None = None, fill: RGBColor | None = None) -> None:
    cell.text = text
    font = cell.text_frame.paragraphs[0].font
    font.size = Pt(size)
    font.bold = bold
    if color is not None:
        font.color.rgb = color
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def _table_geometry(cat_slide, slide_w_in: float, slide_h_in: float):
    """表格锚在映射页标题槽下方；无标题槽时用最大字号可填槽，再退到页面默认。"""
    titles = cat_slide.slots_of_role("title")
    anchor = titles[0] if titles else max(
        (s for s in cat_slide.slots if s.fillable),
        key=lambda s: s.font_pt, default=None)
    if anchor is not None:
        x, w = anchor.x, anchor.w
        y = anchor.y + anchor.h + 0.2
    else:
        x, y = 0.5, 1.2
        w = slide_w_in - 1.0
    w = max(min(w, slide_w_in - x - 0.3), 3.0)
    h = max(slide_h_in - y - 0.4, 1.0)
    return x, y, w, h


# ---------- 通用 ----------

def _serialize(root) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _resolve(base: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    parts: list[str] = []
    for seg in f"{base}/{target}".split("/"):
        if seg in ("", "."):
            continue
        if seg == "..":
            if parts:
                parts.pop()
            continue
        parts.append(seg)
    return "/".join(parts)
