"""源 docx → 输出 docx 的 XML 节点搬运：表格（及 C4 的图片）排版 100% 保真。

deepcopy 源 w:tbl——合并单元格、列宽、行高、表格样式天然原样——再做单元格
文字替换。跨包坑集中在两处：
- 样式定义：输出包缺 tblStyle 会回退默认样式，copy_style_chain 连 basedOn
  祖先一并补入；
- 图片关系：deepcopy 只拷了 r:id 字符串，指向的是源包的 image part，
  rewire_image_rids 用 relate_to 把图片挂进输出包自己的关系表。
"""

from __future__ import annotations

import logging
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Emu

from app.parsing.base import clean_text, image_size_class, vml_style_size_cm

log = logging.getLogger(__name__)

_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_WP_EXTENT = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
_V_SHAPE = "{urn:schemas-microsoft-com:vml}shape"
_V_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"
_EMU_PER_CM = 360000

_SRC_CACHE: dict[str, Document] = {}
_SRC_CACHE_MAX = 8  # 常驻进程防泄漏：渲染期复用，跨任务不无界增长


def open_source(path: Path) -> Document:
    """打开源 docx（每路径缓存，超出上限淘汰最旧）。"""
    key = str(Path(path).resolve())
    if key not in _SRC_CACHE:
        _SRC_CACHE[key] = Document(key)
        if len(_SRC_CACHE) > _SRC_CACHE_MAX:
            _SRC_CACHE.pop(next(iter(_SRC_CACHE)))
    return _SRC_CACHE[key]


def source_table(src_doc: Document, src_index: int):
    """body 顶层第 src_index 个（0 基）w:tbl；越界返回 None。"""
    n = -1
    for child in src_doc.element.body.iterchildren():
        if child.tag == qn("w:tbl"):
            n += 1
            if n == src_index:
                return child
    return None


def copy_style_chain(src_doc: Document, out_doc: Document, style_id: str) -> None:
    """把源样式定义连同 basedOn 祖先链拷进输出包；已存在的跳过，防环。"""
    seen: set[str] = set()
    sid = style_id
    while sid and sid not in seen:
        seen.add(sid)
        st = None
        for el in src_doc.styles.element.findall(qn("w:style")):
            if el.get(qn("w:styleId")) == sid:
                st = el
                break
        if st is None:
            return
        out_ids = {el.get(qn("w:styleId"))
                   for el in out_doc.styles.element.findall(qn("w:style"))}
        if sid not in out_ids:
            out_doc.styles.element.append(deepcopy(st))
        based = st.find(qn("w:basedOn"))
        sid = based.get(qn("w:val")) if based is not None else None


def rewire_image_rids(src_doc: Document, out_doc: Document, el) -> None:
    """把 el 内图片的 r:embed / v:imagedata r:id 重接到输出包的关系表。"""
    rids: set[str] = set()
    for node in el.iter():
        for attr in (_R_EMBED, _R_ID):
            v = node.get(attr)
            if v:
                rids.add(v)
    for rid in sorted(rids):
        rel = src_doc.part.rels.get(rid)
        if rel is None or rel.reltype != RT.IMAGE:
            continue
        new_rid = out_doc.part.relate_to(rel.target_part, RT.IMAGE)
        if new_rid != rid:
            for node in el.iter():
                for attr in (_R_EMBED, _R_ID):
                    if node.get(attr) == rid:
                        node.set(attr, new_rid)


def _insert_into_body(out_doc: Document, el) -> None:
    """插到 sectPr 之前（python-docx 的 add_paragraph/add_table 同语义）。"""
    body = out_doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is not None:
        sect_pr.addprevious(el)
    else:
        body.append(el)


def _normalize_table_width(tbl_el, out_doc: Document) -> None:
    """源表 dxa 宽度超出输出版心 → 等比缩到版心宽（tblW/gridCol/tcW），tblInd 归零。

    源文档页边距常比输出公文版心窄（709 vs 1803 twips），定宽表按源宽搬运
    会溢出版心：居中的越界、带 tblInd 的靠左。pct/auto 型由 Word 自适应，不动。
    """
    try:
        sec = out_doc.sections[0]
        content_tw = Emu(sec.page_width - sec.left_margin - sec.right_margin).twips
    except Exception:  # noqa: BLE001 — 取不到版心信息则放弃缩放，保持源样
        return
    if content_tw <= 0:
        return
    tbl_pr = tbl_el.find(qn("w:tblPr"))
    if tbl_pr is None:
        return
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None or (tbl_w.get(qn("w:type")) or "") != "dxa":
        return
    val = tbl_w.get(qn("w:w")) or ""
    if not val.isdigit() or int(val) <= content_tw:
        return
    scale = content_tw / int(val)
    tbl_w.set(qn("w:w"), str(content_tw))
    grid = tbl_el.find(qn("w:tblGrid"))
    if grid is not None:
        for gc in grid.findall(qn("w:gridCol")):
            gw = gc.get(qn("w:w")) or ""
            if gw.isdigit():
                gc.set(qn("w:w"), str(int(round(int(gw) * scale))))
    # 只缩本表直接行的 tc（嵌套表有自己的 tblW，不越级代管）
    for tr in tbl_el.findall(qn("w:tr")):
        for tc in tr.findall(qn("w:tc")):
            tc_pr = tc.find(qn("w:tcPr"))
            if tc_pr is None:
                continue
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is not None and (tc_w.get(qn("w:type")) or "") == "dxa":
                tw = tc_w.get(qn("w:w")) or ""
                if tw.isdigit():
                    tc_w.set(qn("w:w"), str(int(round(int(tw) * scale))))
    ind = tbl_pr.find(qn("w:tblInd"))
    if ind is not None:
        ind.set(qn("w:w"), "0")
        ind.set(qn("w:type"), "dxa")


def _strip_portrait_images(tbl_el) -> None:
    """删表格内嵌的证件照/装饰小图（用户要求取消照片填充）。

    只删 w:drawing / w:pict 节点本身，run 里的文字保留。VML 判尺寸走
    v:shape 的 style；无尺寸信息不猜、不删。解析期已挡住表格外证件照，
    这里兜表格 XML 深拷贝带进来的网格内照片。
    """
    doomed = []
    for run in tbl_el.iter(qn("w:r")):
        for child in run:
            if child.tag == qn("w:drawing"):
                ext = child.find(f".//{_WP_EXTENT}")
                if ext is None:
                    continue
                try:
                    cx, cy = int(ext.get("cx")), int(ext.get("cy"))
                except (TypeError, ValueError):
                    continue
                size = image_size_class(cx / _EMU_PER_CM, cy / _EMU_PER_CM)
            elif child.tag == qn("w:pict"):
                if child.find(f".//{_V_IMAGEDATA}") is None:
                    continue  # 文本框等非图片形状不动
                shape = child.find(f".//{_V_SHAPE}")
                style = shape.get("style") if shape is not None else ""
                size = image_size_class(*vml_style_size_cm(style))
            else:
                continue
            if size is not None:  # icon 或 portrait
                doomed.append(child)
    for el in doomed:
        el.getparent().remove(el)


def copy_table(src_doc: Document, out_doc: Document, tbl_el):
    """deepcopy 源 w:tbl 到输出文档末尾（含样式链与内嵌图片重接），返回新元素。"""
    new_el = deepcopy(tbl_el)
    _strip_portrait_images(new_el)
    rewire_image_rids(src_doc, out_doc, new_el)
    _normalize_table_width(new_el, out_doc)
    tbl_pr = new_el.find(qn("w:tblPr"))
    if tbl_pr is not None:
        style = tbl_pr.find(qn("w:tblStyle"))
        if style is not None and style.get(qn("w:val")):
            copy_style_chain(src_doc, out_doc, style.get(qn("w:val")))
    _insert_into_body(out_doc, new_el)
    # 表格后补一个空段：相邻 w:tbl 会被 Word 视作一张表
    new_el.addnext(OxmlElement("w:p"))
    return new_el


def copy_image_paragraph(src_doc: Document, out_doc: Document, body_index: int) -> bool:
    """deepcopy body 第 body_index 个子元素（图片所在段落）到输出文档末尾。

    body_index 与解析期一致：body.iterchildren() 的全序号。返回 False 表示
    越界或该位置不是段落（源文件被人工改动过），调用方降级跳过。
    """
    children = list(src_doc.element.body.iterchildren())
    if body_index < 0 or body_index >= len(children):
        return False
    el = deepcopy(children[body_index])
    if el.tag != qn("w:p"):
        return False
    rewire_image_rids(src_doc, out_doc, el)
    p_pr = el.find(qn("w:pPr"))
    if p_pr is not None:
        ps = p_pr.find(qn("w:pStyle"))
        if ps is not None and ps.get(qn("w:val")):
            copy_style_chain(src_doc, out_doc, ps.get(qn("w:val")))
    _insert_into_body(out_doc, el)
    return True


def _tc_text(tc) -> str:
    """与 python-docx _Cell.text 同语义（段落间 \n），供『值未变则不动』比对。"""
    parts = ["".join(t.text or "" for t in p.iter(qn("w:t")))
             for p in tc.findall(qn("w:p"))]
    return clean_text("\n".join(parts))


def _set_tc_text(tc, text: str) -> None:
    """替换单元格文字：留首段首 run 的 rPr（字体字号随源），其余段落清除。"""
    paras = tc.findall(qn("w:p"))
    if not paras:
        return
    first = paras[0]
    p_pr = first.find(qn("w:pPr"))
    runs = first.findall(qn("w:r"))
    r_pr = deepcopy(runs[0].find(qn("w:rPr"))) if runs and runs[0].find(qn("w:rPr")) is not None else None
    for child in list(first):  # 清到只剩 pPr：w:hyperlink 等嵌套文字一并清，防新旧叠加
        if child is not p_pr:
            first.remove(child)
    r = OxmlElement("w:r")
    if r_pr is not None:
        r.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    first.append(r)
    for p in paras[1:]:
        tc.remove(p)


def replace_cell_texts(tbl_el, texts: dict[tuple[int, int], str]) -> None:
    """按 (行, 列) 替换单元格文字，坐标系与 python-docx row.cells 展开一致。

    - gridSpan：一个 tc 占多列，列游标按 span 前进（键取首列）
    - vMerge 续格：跳过写入（视觉上显示的是 restart 格的内容）
    - 键未命中或值与原文一致：不动原 XML（verbatim 格保 run 级格式）
    - 行列越界 / 格式不良：跳过该项，不抛错
    """
    for r, tr in enumerate(tbl_el.findall(qn("w:tr"))):
        c = 0
        for tc in tr.findall(qn("w:tc")):
            tc_pr = tc.find(qn("w:tcPr"))
            span, vmerge_continue = 1, False
            if tc_pr is not None:
                gs = tc_pr.find(qn("w:gridSpan"))
                if gs is not None and (gs.get(qn("w:val")) or "1").isdigit():
                    span = max(1, int(gs.get(qn("w:val"))))
                vm = tc_pr.find(qn("w:vMerge"))
                if vm is not None and vm.get(qn("w:val")) != "restart":
                    vmerge_continue = True
            if not vmerge_continue:
                new = texts.get((r, c))
                if new is not None and new != _tc_text(tc):
                    _set_tc_text(tc, new)
            c += span


def grid_of(tbl_el, parent_doc) -> list[list[str]]:
    """XML 网格 → python-docx row.cells 展开语义的文本网格（与 docx_parser 同式）。"""
    from docx.table import Table as DocxTable

    t = DocxTable(tbl_el, parent_doc)
    return [[clean_text(c.text) for c in row.cells] for row in t.rows]


def replace_table_texts(tbl_el, parent_doc, header: list[str],
                        rows: list[list[str]]) -> None:
    """目标网格（header+rows，展开语义）与 XML 现网格比对，只改写差异格。

    目标值即 TableBlock 最终内容：C3 阶段全部 verbatim → 零改动纯拷贝；
    C5 起长文本格被改写 → 差异格就位，其余格保 run 级格式原样。
    """
    grid = grid_of(tbl_el, parent_doc)
    target = ([list(header)] if header else []) + [list(r) for r in rows]
    if len(target) != len(grid):
        return  # 行数对不上（不该发生）：放弃改写，纯拷贝
    texts: dict[tuple[int, int], str] = {}
    for r, (g_row, t_row) in enumerate(zip(grid, target)):
        for c in range(min(len(g_row), len(t_row))):
            if t_row[c] != g_row[c]:
                texts[(r, c)] = t_row[c]
    if texts:
        replace_cell_texts(tbl_el, texts)
