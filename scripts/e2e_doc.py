"""8 golden 样例 e2e：run --product doc --auto-confirm → DONE + 仿写硬断言。

断言（独立复检，不信任流水线自报）：
  1) 状态 DONE，qa_report 无 E- 级残留
  2) heading+para 数字 100% 溯源（check_numbers）
  3) heading+para 与源文 10-gram 零命中（ngram_hits）
  4) 表格：header 逐字一致；改写格（在 artifacts/tables 的 cells 里）ngram 零命中
     + 数字溯源；其余格 == 源 grid 照搬；断点产物齐套
  5) 产物齐套：output.docx / output.pdf / pages/*.png，PNG 数 == PDF 页数
  6) 排版保真（docx 源）：逐表 tblGrid 列宽与逐行 (gridSpan, vMerge) 序列
     与源一致；解析期保留的图片（tree.images）必须出现在输出 media
  7) 样例特定：resume 值格已换 + 证件照已删

form+PDF 样例（real_performance_review / real_safety_cost / form_personnel.pdf）
走表格 LLM 重建分支，独立复检：qa 无 E-；骨架合法 + 池覆盖 + 碎片全吸收；
改写单元 ngram/数字零违规；html(BOM)/docx/pdf/PNG 齐套且无 ⟦⟧；源含合并 →
输出含 gridSpan/vMerge；竖排碎片按序重组为输出子串；碎片收敛（人事 ≤8 表、
安全生产 3→2、form_personnel 1 表）；值格虚构 + 纯数字照搬。

用法：.venv/Scripts/python.exe scripts/e2e_doc.py [样例路径 ...]
"""

from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

PY = ROOT / ".venv" / "Scripts" / "python.exe"
GOLDEN = [
    ROOT / "examples" / "letter_application.docx",
    ROOT / "examples" / "sectioned_report.docx",
    ROOT / "examples" / "text_report.pdf",
    ROOT / "examples" / "form_personnel.docx",
    ROOT / "examples" / "resume_sample.docx",
    ROOT / "examples" / "form_personnel.pdf",
    ROOT / "examples" / "real_performance_review.pdf",
    ROOT / "examples" / "real_safety_cost.pdf",
]


def _tbl_signatures(docx_path: Path):
    """逐表 (gridCol 宽度列表, 逐行 [(gridSpan, vMerge), …])——排版指纹。"""
    from docx import Document
    from docx.oxml.ns import qn

    d = Document(str(docx_path))
    out = []
    for tbl in d.tables:
        grid = tbl._tbl.find(qn("w:tblGrid"))
        widths = [gc.get(qn("w:w")) for gc in grid.findall(qn("w:gridCol"))]
        rows = []
        for tr in tbl._tbl.findall(qn("w:tr")):
            sig = []
            for tc in tr.findall(qn("w:tc")):
                tc_pr = tc.find(qn("w:tcPr"))
                span, vm = 1, None
                if tc_pr is not None:
                    gs = tc_pr.find(qn("w:gridSpan"))
                    if gs is not None and (gs.get(qn("w:val")) or "").isdigit():
                        span = int(gs.get(qn("w:val")))
                    v = tc_pr.find(qn("w:vMerge"))
                    if v is not None:
                        vm = v.get(qn("w:val")) or "continue"
                sig.append((span, vm))
            rows.append(sig)
        out.append((widths, rows))
    return out


def _zip_media(docx_path: Path) -> list[str]:
    import zipfile

    with zipfile.ZipFile(docx_path) as z:
        return [n for n in z.namelist() if n.startswith("word/media/")]


def run_sample(src: Path) -> tuple[bool, list[str]]:
    from app.config import DATA_DIR
    from app.qa.ngram import ngram_hits
    from app.qa.numbers import check_numbers
    from app.schema.doctree import DocTree

    errors: list[str] = []
    r = subprocess.run(
        [str(PY), "-m", "app", "run", str(src), "--product", "doc", "--auto-confirm"],
        cwd=str(ROOT), capture_output=True, text=True, encoding="utf-8",
        errors="replace", timeout=3000)
    out = (r.stdout or "") + (r.stderr or "")
    if r.returncode != 0:
        return False, [f"CLI 失败（rc={r.returncode}）：\n{out[-1500:]}"]

    job_id = next((l.split("=", 1)[1].strip() for l in out.splitlines()
                   if l.startswith("job_id=")), None)
    if not job_id:
        return False, [f"未解析到 job_id：\n{out[-500:]}"]
    job_dir = DATA_DIR / "jobs" / job_id
    art = job_dir / "artifacts"

    state = json.loads((job_dir / "state.json").read_text(encoding="utf-8"))
    if state["status"] != "DONE":
        return False, [f"状态 {state['status']}：{state.get('error')}"]

    tree = DocTree.model_validate_json(
        (art / "doctree.json").read_text(encoding="utf-8"))

    # form+PDF → 表格 LLM 重建分支（与 doc_runner 分发镜像；无 docir.json）
    plan_raw = json.loads((art / "docplan.json").read_text(encoding="utf-8"))
    if tree.meta.source_format == "pdf" and plan_raw.get("genre") == "form":
        return run_form_sample(src, job_dir, art, tree)

    docir = json.loads((art / "docir.json").read_text(encoding="utf-8"))

    # 1) 数字 100% 溯源 + 2) 10-gram 零命中（逐块独立复检：
    #    块间拼接会把"段尾句号+下个标题"凑成伪 10-gram，渲染产物里块本是分行）
    for b in docir["blocks"]:
        if b["kind"] not in ("heading", "para") or not b.get("text", "").strip():
            continue
        for tok, ctx in check_numbers(b["text"], tree.full_text):
            errors.append(f"[E-NUM-UNTRACED] {tok}: {ctx}")
        for g in ngram_hits(b["text"], tree.full_text):
            errors.append(f"[E-PLAGIARISM] 与源文连续雷同：{g}")

    # 3) 表格内容分类断言：header 逐字；改写格（断点产物 cells）零雷同 + 数字
    #    溯源；其余格 == 源 grid 照搬
    src_tables = {t.table_id: t for t in tree.tables}
    for b in docir["blocks"]:
        if b["kind"] != "table":
            continue
        tid = b["table_id"]
        t = src_tables.get(tid)
        if t is None:
            errors.append(f"[TABLE-DIFF] {tid} 不在源文档")
            continue
        if b["header"] != list(t.header):
            errors.append(f"[TABLE-DIFF] {tid} header 与源表不一致")
        src_rows = [list(r) for r in t.rows]
        got_rows = b["rows"]
        if len(got_rows) != len(src_rows) or any(
                len(g) != len(s) for g, s in zip(got_rows, src_rows)):
            errors.append(f"[TABLE-DIFF] {tid} 行列结构与源表不一致")
            continue
        art_path = art / "tables" / f"{tid}.json"
        if not art_path.exists():
            errors.append(f"[TABLE-DIFF] {tid} 缺 tablefill 断点产物")
            continue
        rewrites = json.loads(
            art_path.read_text(encoding="utf-8"))["cells"]
        for r, (got_row, src_row) in enumerate(zip(got_rows, src_rows)):
            for c, (got, want) in enumerate(zip(got_row, src_row)):
                key = f"{r},{c}"
                if key in rewrites:
                    if got != rewrites[key]:
                        errors.append(
                            f"[TABLE-DIFF] {tid} {key} 与断点产物不一致")
                    for tok, ctx in check_numbers(got, tree.full_text):
                        errors.append(
                            f"[E-NUM-UNTRACED] {tid} {key} {tok}: {ctx}")
                    for g in ngram_hits(got, tree.full_text):
                        errors.append(
                            f"[E-PLAGIARISM] {tid} {key} 与源文连续雷同：{g}")
                elif got != want:
                    errors.append(
                        f"[TABLE-DIFF] {tid} {key} 非改写格与源不一致"
                        f"：{got!r} != {want!r}")

    # 4) QA 报告无 E- 级残留（W- 级如长段提示/表格退格不算失败）
    qa = (art / "qa_report.txt").read_text(encoding="utf-8")
    residuals = [l for l in qa.splitlines() if "[docqa]" in l and " E-" in l]
    if residuals:
        errors.append(f"QA 残留 {len(residuals)} 项：\n" + "\n".join(residuals[:5]))

    # 5) 产物齐套
    docx, pdf = art / "output.docx", art / "output.pdf"
    if not docx.exists() or docx.stat().st_size < 4000:
        errors.append("output.docx 缺失或过小")
    if not pdf.exists() or pdf.stat().st_size < 4000:
        errors.append("output.pdf 缺失或过小")
    else:
        import pymupdf

        with pymupdf.open(str(pdf)) as d:
            n_pages = len(d)
        pngs = sorted((job_dir / "pages").glob("page_*.png"))
        if len(pngs) != n_pages:
            errors.append(f"预览 PNG {len(pngs)} 张 ≠ PDF {n_pages} 页")

    from docx import Document

    d = Document(str(docx))
    if len(d.paragraphs) < 2:
        errors.append(f"docx 段落数异常（{len(d.paragraphs)}）")

    # 6) 排版保真（docx 源）：逐表 tblGrid 列宽 + 逐行合并结构与源一致；
    #    解析期保留的图片（tree.images，证件照/图标已在解析期过滤）必须
    #    出现在输出 media
    if src.suffix.lower() == ".docx":
        src_docx = job_dir / "upload" / src.name
        if not src_docx.exists():
            errors.append("upload 源 docx 缺失，无法比对排版")
        else:
            if _tbl_signatures(src_docx) != _tbl_signatures(docx):
                errors.append("[LAYOUT-DIFF] 表格 tblGrid/合并结构与源不一致")
    if len(_zip_media(docx)) < len(tree.images):
        errors.append(
            f"[LAYOUT-DIFF] 输出 media {_zip_media(docx)} 少于解析保留的 "
            f"{len(tree.images)} 张图片")

    # 7) 样例特定断言
    if src.name == "resume_sample.docx":
        art_path = art / "tables" / "tbl-001.json"
        cells = (json.loads(art_path.read_text(encoding="utf-8"))["cells"]
                 if art_path.exists() else {})
        if "0,1" not in cells:
            errors.append("[E-VALUE] 姓名值格未被虚构改写（0,1 不在改写产物）")
        elif cells["0,1"] == "张三":
            errors.append("[E-VALUE] 姓名值格照抄原值")
        if _zip_media(docx):
            errors.append("[E-VALUE] 输出仍含图片（证件照未删除）")

    return not errors, errors


def run_form_sample(src: Path, job_dir: Path, art: Path, tree) -> tuple[bool, list[str]]:
    """form+PDF 表格 LLM 重建分支独立复检（不信任流水线自报）。

    ① qa 无 E- 级残留；② 架构师产物独立重校（合法性 + 源池覆盖 + 碎片
    全吸收）；③ 内容专家逐改写单元 ngram/数字溯源零违规；④ html(BOM)/
    docx/pdf/PNG 齐套且无 ⟦⟧ 残留；⑤ 源含合并区 → 输出 XML 含
    gridSpan/vMerge；⑥ 竖排 1 列碎片按序拼接是输出子串；⑦ 碎片收敛；
    ⑧ 样例特定（值格虚构 + 纯数字照搬）。
    """
    import re

    from docx import Document

    from app.pipeline.tblarch import build_mask_pool
    from app.qa.ngram import ngram_hits
    from app.qa.numbers import check_numbers
    from app.schema.tblskeleton import (
        TSkeleton,
        coverage_missing,
        validate_skeleton,
        walk_grid,
    )

    errors: list[str] = []

    # ① QA 报告无 E- 级残留（分支产物仅 W 级可接受）
    qa = (art / "qa_report.txt").read_text(encoding="utf-8")
    bad = [l for l in qa.splitlines() if " E-" in l]
    if bad:
        errors.append(f"QA 残留 E- {len(bad)} 项：" + "; ".join(bad[:3]))

    # ② 架构师独立复检：骨架合法 + 池覆盖 + 源碎片全被吸收
    if not (art / "tblarch.json").exists():
        errors.append("tblarch.json 缺失——form 分支回落了旧链路"
                      "（架构师失败/超时），golden 要求走新分支")
        return not errors, errors
    arch = json.loads((art / "tblarch.json").read_text(encoding="utf-8"))
    skeletons = [TSkeleton.model_validate(t) for t in arch["tables"]]
    for s in skeletons:
        for e in validate_skeleton(s):
            errors.append(f"[TBLARCH] 骨架非法：{e}")
    pool = build_mask_pool(tree)
    for m in coverage_missing(skeletons, pool):
        errors.append(f"[TBLARCH] {m}")
    absorbed = {tid for s in skeletons for tid in s.src_tables}
    for t in tree.tables:
        if t.table_id not in absorbed:
            errors.append(f"[TBLARCH] 源碎片 {t.table_id} 未被任何骨架吸收")

    # ③ 逐改写单元（骨架格 + 散文）ngram/数字溯源零违规
    cells = json.loads((art / "tblcontent.json").read_text(
        encoding="utf-8"))["cells"]
    for k, text in sorted(cells.items()):
        for tok, ctx in check_numbers(text, tree.full_text):
            errors.append(f"[E-NUM-UNTRACED] {k} {tok}: {ctx}")
        for g in ngram_hits(text, tree.full_text):
            errors.append(f"[E-PLAGIARISM] {k} 与源文连续雷同：{g}")

    # ④ 产物齐套：html(BOM)/docx/pdf/PNG==页数；无 ⟦⟧ 残留
    html_raw = (art / "output.html").read_bytes() \
        if (art / "output.html").exists() else b""
    if not html_raw.startswith(b"\xef\xbb\xbf"):
        errors.append("output.html 缺失或无 BOM（COM 主路径未走）")
    docx, pdf = art / "output.docx", art / "output.pdf"
    if not docx.exists() or docx.stat().st_size < 4000:
        errors.append("output.docx 缺失或过小")
        return not errors, errors
    if not pdf.exists() or pdf.stat().st_size < 4000:
        errors.append("output.pdf 缺失或过小")
    else:
        import pymupdf

        with pymupdf.open(str(pdf)) as d:
            n_pages = len(d)
        pngs = sorted((job_dir / "pages").glob("page_*.png"))
        if len(pngs) != n_pages:
            errors.append(f"预览 PNG {len(pngs)} 张 ≠ PDF {n_pages} 页")

    d = Document(str(docx))
    out_text = "\n".join(p.text for p in d.paragraphs) + "\n" + "\n".join(
        c.text for t in d.tables for r in t.rows for c in r.cells)
    if "⟦" in out_text:
        errors.append("docx 残留 ⟦N⟧ 占位符")
    if html_raw and "⟦" in html_raw.decode("utf-8-sig"):
        errors.append("html 残留 ⟦N⟧ 占位符")

    # ⑤ 源含合并区 → 输出 XML 须有 gridSpan/vMerge
    if any(t.merges for t in tree.tables):
        if "gridSpan" not in d.element.xml and "vMerge" not in d.element.xml:
            errors.append("[LAYOUT-DIFF] 源含合并区但输出无 gridSpan/vMerge")

    # ⑥ 竖排重组：每个 1 列源碎片非空格按序拼接是输出（去空白）子串
    out_despaced = "".join(out_text.split())
    cjk = re.compile(r"[一-鿿]")
    for t in tree.tables:
        if t.n_cols != 1:
            continue
        joined = "".join("".join(c.split()) for row in
                         ([t.header] if t.header else []) + t.rows
                         for c in row)
        if cjk.search(joined) and len(joined) >= 3 and joined not in out_despaced:
            errors.append(f"[LAYOUT-DIFF] 竖排碎片未重组：{t.table_id}“{joined}”")

    # ⑦ 碎片收敛 + ⑧ 样例特定
    n_out = len(d.tables)
    if src.name == "real_performance_review.pdf":
        if n_out > 8:
            errors.append(f"[LAYOUT-DIFF] 碎片未收敛：输出 {n_out} 表 > 8"
                          f"（源 {len(tree.tables)} 碎片）")
    elif src.name == "real_safety_cost.pdf":
        if n_out != 2:
            errors.append(f"[LAYOUT-DIFF] 跨页并表后应 2 表，实际 {n_out}")
    elif src.name == "form_personnel.pdf":
        if n_out != 1:
            errors.append(f"[LAYOUT-DIFF] 应 1 表，实际 {n_out}")
        if not skeletons:
            errors.append("[TBLARCH] 无骨架")
        else:
            anchors, _ = walk_grid(skeletons[0])
            name_key = phone_key = None
            for (r, c), cell in anchors.items():
                text = "".join(cell.content.split())
                if text == "张伟":
                    name_key = f"t0 {r},{c}"
                if "13800001234" in text:
                    phone_key = f"t0 {r},{c}"
            if name_key is None:
                errors.append("[E-VALUE] 骨架未含姓名值格（张伟）")
            elif name_key not in cells:
                errors.append(f"[E-VALUE] 姓名值格未被虚构改写"
                              f"（{name_key} 不在改写产物）")
            elif cells[name_key] == "张伟":
                errors.append("[E-VALUE] 姓名值格照抄原值")
            if phone_key is not None and phone_key in cells:
                errors.append(f"[E-VALUE] 纯数字电话格不应进入改写（{phone_key}）")
            if "13800001234" not in out_despaced:
                errors.append("[E-VALUE] 电话号码未照搬进输出")

    return not errors, errors


def main(argv: list[str]) -> int:
    targets = [Path(a) for a in argv] or GOLDEN
    ok_all = True
    for src in targets:
        if not src.exists():
            print(f"[skip] 样例不存在：{src}")
            ok_all = False
            continue
        print(f"\n===== {src.name} =====")
        ok, errs = run_sample(src)
        print("[PASS]" if ok else "[FAIL]")
        for e in errs:
            print(f"  {e}")
        ok_all = ok_all and ok
    print("\n" + ("全部通过" if ok_all else "存在失败"))
    return 0 if ok_all else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
