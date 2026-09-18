"""生成 form golden 样例 examples/form_personnel.docx：表格式文书。

排版保真三要素：gridSpan / vMerge 合并格、tblGrid 显式列宽、内嵌流程图
PNG（PIL 绘制）。另含 2 个 ≥12 汉字当量、带数字的描述格（tablefill 仿写
目标，位于未合并格）。合并格与表头刻意保持短文本——照搬路径，不进 LLM。

用法：.venv/Scripts/python.exe scripts/make_form_sample.py
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "examples" / "form_personnel.docx"


def _flowchart_png() -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    w, h = 960, 220
    img = Image.new("RGB", (w, h), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/simhei.ttf", 22)
    except OSError:
        font = ImageFont.load_default()
    steps = ["调岗申请提出", "部门负责人审批", "人事部备案", "完成岗位调岗"]
    box_w, box_h, gap = 190, 70, 60
    x, y = 20, (h - box_h) // 2
    for i, step in enumerate(steps):
        draw.rectangle([x, y, x + box_w, y + box_h], outline="black", width=2)
        bbox = draw.textbbox((0, 0), step, font=font)
        draw.text((x + (box_w - bbox[2]) / 2, y + (box_h - bbox[3]) / 2),
                  step, fill="black", font=font)
        if i < len(steps) - 1:
            ax = x + box_w + 8
            mid = y + box_h // 2
            draw.line([ax, mid, ax + gap - 16, mid], fill="black", width=2)
            draw.polygon([(ax + gap - 16, mid - 6), (ax + gap - 16, mid + 6),
                          (ax + gap - 6, mid)], fill="black")
        x += box_w + gap
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def main() -> Path:
    doc = Document()
    doc.add_heading("员工岗位变动审批表", 0)

    t = doc.add_table(rows=6, cols=4)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, w in enumerate([Cm(2.8), Cm(6.6), Cm(3.0), Cm(3.5)]):
        t.columns[i].width = w
        for cell in t.columns[i].cells:
            cell.width = w

    # 合并：行 2 两组横向 gridSpan；行 3-4 首列纵向 vMerge；行 5 整行 gridSpan
    t.cell(2, 0).merge(t.cell(2, 1))
    t.cell(2, 2).merge(t.cell(2, 3))
    t.cell(3, 0).merge(t.cell(4, 0))
    t.cell(5, 0).merge(t.cell(5, 3))

    rows = [
        ["姓名", "部门", "入职日期", "联系电话"],
        ["张伟", "市场部", "2021 年 3 月", "13800001234"],
        ["变动类型：平级调岗", "", "生效日期：2026-10-01", ""],
        ["工作交接情况", "负责华东区域 3 个城市的市场维护，移交客户档案 128 份",
         "待办事项", "12 项"],
        ["", "完成整体交接事项 12 项，相关手续已全部办理完毕", "责任人", "李明"],
        ["审批意见：同意", "", "", ""],
    ]
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            if text:
                t.cell(r, c).text = text

    doc.add_picture(io.BytesIO(_flowchart_png()), width=Cm(14))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"生成 {OUT}")
    return OUT


if __name__ == "__main__":
    main()
